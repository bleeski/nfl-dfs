from __future__ import annotations

import json
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.views import Selection

from .contracts import CertificationManifest, Lineup
from .qa import QAFinding
from .system import workbook_is_closed


NAVY = "17324D"
BLUE = "2F75B5"
PALE_BLUE = "DDEBF7"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "E7E6E6"


class WorkbookLockedError(RuntimeError):
    pass


def _safe_display_value(value: object) -> object:
    """Return inert Excel display text for every provider/operator string.

    OpenPyXL interprets a string beginning with ``=`` as a formula. Spreadsheet
    applications also recognize leading ``+``, ``-``, ``@`` and whitespace-
    prefixed formulas when a file is opened or copied. The presentation copy is
    prefixed with an apostrophe; exact source values remain in the hash-bound
    JSON/CSV artifacts and are never used from this escaped display cell.
    """

    if not isinstance(value, str):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("�", value)
    probe = text.lstrip(" \t\r\n")
    if text[:1] in {"\t", "\r", "\n"} or probe[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _set_display(cell, value: object) -> None:
    cell.value = _safe_display_value(value)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _configure_review_print_layout(
    ws,
    *,
    print_area: str,
    repeat_rows: str,
    orientation: str = "landscape",
) -> None:
    """Give the generated review copy a predictable, readable print surface."""

    ws.print_area = print_area
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = repeat_rows


def _finish_generated_sheet(
    ws,
    *,
    widths: Mapping[int, float],
    print_area: str,
    repeat_rows: str = "1:4",
) -> None:
    ws.freeze_panes = "A5"
    # Portfolio starts with a two-axis freeze pane. Replacing it with an A5
    # freeze does not make openpyxl discard the old right-pane selections,
    # which native Excel reports as a repaired sheet view. Generated review
    # sheets have one valid bottom-left selection regardless of their origin.
    ws.sheet_view.selection = [
        Selection(pane="bottomLeft", activeCell="A5", sqref="A5")
    ]
    _configure_review_print_layout(
        ws,
        print_area=print_area,
        repeat_rows=repeat_rows,
    )
    ws.sheet_view.showGridLines = False
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def _populate_readable_review(workbook, review: Mapping[str, object]) -> None:
    """Add the SD5 human surface without changing the five-sheet input contract."""

    portfolio = workbook["Portfolio"]
    portfolio.tables.clear()
    if portfolio.max_row >= 4:
        portfolio.delete_rows(4, portfolio.max_row - 3)
    _title(
        portfolio,
        "Prior-only exact entry assignments",
        "Every row is reconciled to the exact audited assignment and exported review CSV. Scores are PRIOR_ONLY central estimates.",
        15,
    )
    lineup_headers = (
        "Entry ID", "Contest", "Contest ID", "Slot", "Player", "Exact DK roster ID",
        "Underlying person ID", "Team", "Position", "Salary", "Lineup salary",
        "Salary remaining", "Prior-only slot points", "Prior-only lineup points",
        "Official activity / role concern",
    )
    for column, value in enumerate(lineup_headers, start=1):
        _set_display(portfolio.cell(4, column), value)
        portfolio.cell(4, column).fill = PatternFill("solid", fgColor=BLUE)
        portfolio.cell(4, column).font = Font(bold=True, color=WHITE)
    row_number = 5
    for raw_entry in review.get("entries", []):
        if not isinstance(raw_entry, Mapping):
            continue
        for raw_slot in raw_entry.get("slots", []):
            if not isinstance(raw_slot, Mapping):
                continue
            findings = "; ".join(
                str(item.get("finding", ""))
                for item in raw_slot.get("role_findings", [])
                if isinstance(item, Mapping)
            )
            concern = " | ".join(
                value
                for value in (
                    f"official={raw_slot.get('official_activity')}",
                    f"role={raw_slot.get('role_evidence_state')}",
                    findings,
                )
                if value
            )
            values = (
                raw_entry.get("entry_id"), raw_entry.get("contest_name") or "Contest label unavailable",
                raw_entry.get("contest_id"), raw_slot.get("slot"), raw_slot.get("name"),
                raw_slot.get("dk_roster_id"), raw_slot.get("underlying_person_id"),
                raw_slot.get("team"), raw_slot.get("position"), raw_slot.get("salary"),
                raw_entry.get("salary_total"), raw_entry.get("salary_remaining"),
                raw_slot.get("prior_only_central_estimate_points"),
                raw_entry.get("prior_only_central_estimate_points"), concern,
            )
            for column, value in enumerate(values, start=1):
                _set_display(portfolio.cell(row_number, column), value)
            for column in (10, 11, 12):
                portfolio.cell(row_number, column).number_format = '"$"#,##0'
            for column in (13, 14):
                portfolio.cell(row_number, column).number_format = "0.000"
            portfolio.row_dimensions[row_number].height = 42
            row_number += 1
    portfolio.auto_filter.ref = f"A4:O{max(4, row_number - 1)}"
    _finish_generated_sheet(
        portfolio,
        print_area=f"A1:O{max(4, row_number - 1)}",
        widths={1: 15, 2: 28, 3: 14, 4: 10, 5: 24, 6: 18, 7: 32, 8: 9,
                9: 10, 10: 11, 11: 13, 12: 15, 13: 15, 14: 17, 15: 44},
    )

    exposure = workbook.create_sheet("Exposure")
    _title(
        exposure,
        "Actual portfolio exposure",
        "Counts and percentages are recomputed from the exact audited assignment artifact. Combined-person and Captain exposure are separate.",
        13,
    )
    exposure_headers = (
        "Player", "Underlying person ID", "Team", "Combined count", "Combined %",
        "Combined max count", "Combined max %", "Captain count", "Captain %",
        "Captain max count", "Captain max %", "Excluded", "Exclusion basis",
    )
    for column, value in enumerate(exposure_headers, start=1):
        _set_display(exposure.cell(4, column), value)
        exposure.cell(4, column).fill = PatternFill("solid", fgColor=BLUE)
        exposure.cell(4, column).font = Font(bold=True, color=WHITE)
    row_number = 5
    exposure_payload = review.get("exposure", {})
    if not isinstance(exposure_payload, Mapping):
        exposure_payload = {}
    for raw in exposure_payload.get("people", []):
        if not isinstance(raw, Mapping):
            continue
        values = (
            raw.get("name"), raw.get("underlying_person_id"), raw.get("team"),
            raw.get("combined_count"), raw.get("combined_percentage"),
            raw.get("combined_max_count"), raw.get("combined_max_percentage"),
            raw.get("captain_count"), raw.get("captain_percentage"),
            raw.get("captain_max_count"), raw.get("captain_max_percentage"),
            "YES" if raw.get("excluded") else "NO", raw.get("exclusion_source"),
        )
        for column, value in enumerate(values, start=1):
            _set_display(exposure.cell(row_number, column), value)
        for column in (5, 7, 9, 11):
            exposure.cell(row_number, column).number_format = '0.000"%"'
        exposure.row_dimensions[row_number].height = 30
        row_number += 1
    summary_row = row_number + 1
    summary = (
        ("Canonical uniqueness", exposure_payload.get("canonical_uniqueness")),
        ("Unique lineups required", exposure_payload.get("unique_required")),
        ("Configured pairwise person overlap", exposure_payload.get("configured_pairwise_person_overlap")),
        ("Effective pairwise person overlap", exposure_payload.get("effective_pairwise_person_overlap")),
    )
    for offset, (label, value) in enumerate(summary):
        _set_display(exposure.cell(summary_row + offset, 1), label)
        _set_display(exposure.cell(summary_row + offset, 2), value)
        exposure.cell(summary_row + offset, 1).font = Font(bold=True)
    overlap_row = summary_row + len(summary) + 2
    _section_header(exposure, overlap_row, 1, 4, "Configured and actual pairwise underlying-person overlap")
    for column, value in enumerate(("Entry A", "Entry B", "Actual shared people", "Maximum"), start=1):
        _set_display(exposure.cell(overlap_row + 1, column), value)
        exposure.cell(overlap_row + 1, column).font = Font(bold=True)
    for raw in exposure_payload.get("pairwise_overlap", []):
        if not isinstance(raw, Mapping):
            continue
        overlap_row += 1
        values = (raw.get("entry_id_a"), raw.get("entry_id_b"), raw.get("actual_people"), raw.get("maximum_people"))
        for column, value in enumerate(values, start=1):
            _set_display(exposure.cell(overlap_row + 1, column), value)
    exposure.auto_filter.ref = f"A4:M{max(4, row_number - 1)}"
    _finish_generated_sheet(
        exposure,
        print_area=f"A1:M{exposure.max_row}",
        widths={1: 24, 2: 34, 3: 9, 4: 14, 5: 13, 6: 18, 7: 16, 8: 13,
                9: 12, 10: 17, 11: 15, 12: 11, 13: 28},
    )

    evidence = workbook.create_sheet("Review Evidence")
    _title(
        evidence,
        "Evidence, role, and model observations",
        "Official activity is separate from role evidence. Qualitative observations do not create numerical projections or confidence scores.",
        7,
    )
    evidence_headers = ("Category", "State", "Observation", "Observed", "Expires", "Source", "Next evidence action")
    for column, value in enumerate(evidence_headers, start=1):
        _set_display(evidence.cell(4, column), value)
        evidence.cell(4, column).fill = PatternFill("solid", fgColor=BLUE)
        evidence.cell(4, column).font = Font(bold=True, color=WHITE)
    row_number = 5
    for raw in review.get("evidence_observations", []):
        if not isinstance(raw, Mapping):
            continue
        values = tuple(raw.get(key) for key in ("category", "state", "observation", "observed_at", "expires_at", "source", "next_action"))
        for column, value in enumerate(values, start=1):
            if isinstance(value, (dict, list, tuple)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
            _set_display(evidence.cell(row_number, column), value)
        evidence.row_dimensions[row_number].height = 44
        row_number += 1
    evidence.auto_filter.ref = f"A4:G{max(4, row_number - 1)}"
    _finish_generated_sheet(
        evidence,
        print_area=f"A1:G{max(4, row_number - 1)}",
        widths={1: 22, 2: 30, 3: 58, 4: 25, 5: 25, 6: 52, 7: 52},
    )

    artifact_sheet = workbook.create_sheet("Artifacts")
    _title(
        artifact_sheet,
        "Artifact provenance and exact hashes",
        "The readable package passed independent exact-byte reconciliation. Use these paths and hashes to identify the reviewed bytes.",
        3,
    )
    for column, value in enumerate(("Artifact", "Local path", "SHA-256"), start=1):
        _set_display(artifact_sheet.cell(4, column), value)
        artifact_sheet.cell(4, column).fill = PatternFill("solid", fgColor=BLUE)
        artifact_sheet.cell(4, column).font = Font(bold=True, color=WHITE)
    row_number = 5
    for raw in review.get("artifacts", []):
        if not isinstance(raw, Mapping):
            continue
        for column, value in enumerate((raw.get("name"), raw.get("path"), raw.get("sha256")), start=1):
            _set_display(artifact_sheet.cell(row_number, column), value)
        href = raw.get("href")
        if isinstance(href, str):
            artifact_sheet.cell(row_number, 2).hyperlink = href
            artifact_sheet.cell(row_number, 2).style = "Hyperlink"
            artifact_sheet.cell(row_number, 2).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        artifact_sheet.row_dimensions[row_number].height = 54
        row_number += 1
    row_number += 1
    _section_header(artifact_sheet, row_number, 1, 3, "Every bound hash")
    row_number += 1
    for column, value in enumerate(("Hash label", "SHA-256", ""), start=1):
        _set_display(artifact_sheet.cell(row_number, column), value)
        artifact_sheet.cell(row_number, column).font = Font(bold=True)
    for label, digest in sorted(review.get("hashes", {}).items()):
        row_number += 1
        _set_display(artifact_sheet.cell(row_number, 1), label)
        _set_display(artifact_sheet.cell(row_number, 2), digest)
    _finish_generated_sheet(
        artifact_sheet,
        print_area=f"A1:C{artifact_sheet.max_row}",
        widths={1: 34, 2: 76, 3: 70},
    )

    upload = workbook["Upload"]
    truths = review.get("truths", {})
    if not isinstance(truths, Mapping):
        truths = {}
    for row, label in enumerate(
        ("FILE_VALID", "EVIDENCE_STATE", "MODEL_STATUS", "RELEASE_DECISION"),
        start=5,
    ):
        value = truths.get(label)
        _set_display(upload.cell(row, 2), str(value).upper() if isinstance(value, bool) else value)
    _set_display(upload["A17"], "DISPLAY_RECONCILIATION")
    _set_display(upload["B17"], review.get("reconciliation", {}).get("status") if isinstance(review.get("reconciliation"), Mapping) else "UNKNOWN")
    _set_display(upload["A18"], "READABLE_WARNING")
    _set_display(upload["B18"], review.get("warning"))
    for row in (17, 18):
        upload.cell(row, 1).font = Font(bold=True)
        upload.row_dimensions[row].height = 42
    upload.column_dimensions["A"].width = 28
    upload.column_dimensions["B"].width = 88

    # These four sheets begin life as an operator input template. Restrict only
    # the generated review copy to logical printable regions so empty template
    # columns and title merges do not create fragmented or blank PDF pages.
    _configure_review_print_layout(
        workbook["Run Control"],
        print_area=f"A1:D{workbook['Run Control'].max_row}",
        repeat_rows="1:4",
    )
    _configure_review_print_layout(
        workbook["Evidence Paste"],
        print_area="A1:E105,G4:J108,L4:O105,Q4:W37",
        repeat_rows="4:5",
    )
    _configure_review_print_layout(
        workbook["QA"],
        print_area=f"A1:G{workbook['QA'].max_row}",
        repeat_rows="1:4",
    )
    _configure_review_print_layout(
        upload,
        print_area="A1:B18",
        repeat_rows="1:3",
        orientation="portrait",
    )


@dataclass(frozen=True)
class OperatorInput:
    values: dict[str, str | float | int | None]
    official_status_rows: tuple[tuple[object, ...], ...]
    payout_rows: tuple[tuple[object, ...], ...]
    ownership_rows: tuple[tuple[object, ...], ...]


def _title(ws, title: str, subtitle: str, columns: int = 10) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sub = ws.cell(2, 1, subtitle)
    sub.font = Font(name="Aptos", size=10, italic=True, color="666666")
    sub.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    ws.sheet_view.showGridLines = False


def _section_header(ws, row: int, start_col: int, end_col: int, value: str) -> None:
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row, start_col, value)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Aptos", bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")


def _table(ws, reference: str, name: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def create_operator_input_workbook(path: str | Path) -> Path:
    target = Path(path).resolve()
    if target.exists() and not workbook_is_closed(target):
        raise WorkbookLockedError(
            f"Close {target.name} in Excel, then run the command again. No data was changed."
        )
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    run_control = workbook.create_sheet("Run Control")
    evidence = workbook.create_sheet("Evidence Paste")
    portfolio = workbook.create_sheet("Portfolio")
    qa = workbook.create_sheet("QA")
    upload = workbook.create_sheet("Upload")

    _title(
        run_control,
        "NFL DFS Run Control",
        "Edit only pale-yellow cells in this staged input workbook. Close Excel before running the engine.",
        8,
    )
    headers = ("KEY", "VALUE", "REQUIRED", "OPERATOR GUIDANCE")
    run_control.append([])
    run_control.append(headers)
    rows = [
        ("RUN_LABEL", "", "YES", "Short label such as 2026-W02-MAIN"),
        ("SALARY_CSV", "", "YES", "Full path to the DraftKings salary CSV"),
        ("ENTRY_CSV", "", "YES", "Full path to the reserved-entry CSV"),
        ("PAYOUT_CSV", "", "YES", "Strict payout CSV path"),
        ("ASSIGNMENT_CSV", "", "MANUAL", "Hand-created lineup assignment CSV, if used"),
        ("TEAM_PROJECTION_CSV", "", "MODEL", "Required only for model-assisted builds"),
        ("PLAYER_OPPORTUNITY_CSV", "", "MODEL", "Required only for model-assisted builds"),
        ("OFFICIAL_STATUS_CSV", "", "UPLOAD", "Exact-ID official active/inactive evidence"),
        ("FIELD_SIZE", "", "YES", "Total contest entries"),
        ("MAX_ENTRIES", "", "YES", "Contest maximum entries per person"),
        ("ADVERTISED_PRIZE_VALUE", "", "YES", "Cash plus ticket face value"),
        ("TICKET_FACE_VALUE", "", "SATELLITE", "Face value of each awarded ticket, if applicable"),
        ("CONTEST_OBJECTIVE", "LARGE_GPP", "YES", "LARGE_GPP, SMALL_GPP, CASH, WTA, or SATELLITE"),
        ("MANUAL_GUARDRAIL_MODE", "YES", "YES", "YES allows legality-only workflow; it does not certify model quality"),
        (
            "ROLE_EVIDENCE_JSON",
            "",
            "SHOWDOWN",
            "Generated source-bound kicker-role artifact, if current role is ambiguous",
        ),
        (
            "OFFENSIVE_ROLE_EVIDENCE_JSON", "", "SHOWDOWN",
            "Generated captured-source offensive allocation or unresolved-role facts",
        ),
        (
            "PORTFOLIO_POLICY_JSON", "", "SHOWDOWN",
            "Versioned exact-ID portfolio controls; SD3 validates but SD4 must enforce",
        ),
    ]
    for row in rows:
        run_control.append(row)
    _table(run_control, f"A4:D{3 + len(rows)}", "RunControlTable")
    for row in range(5, 5 + len(rows)):
        run_control.cell(row, 2).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    objective_validation = DataValidation(
        type="list", formula1='"LARGE_GPP,SMALL_GPP,CASH,WTA,SATELLITE"', allow_blank=False
    )
    yes_no_validation = DataValidation(type="list", formula1='"YES,NO"', allow_blank=False)
    run_control.add_data_validation(objective_validation)
    run_control.add_data_validation(yes_no_validation)
    objective_validation.add(run_control["B17"])
    yes_no_validation.add(run_control["B18"])
    run_control.freeze_panes = "A4"
    run_control.column_dimensions["A"].width = 28
    run_control.column_dimensions["B"].width = 58
    run_control.column_dimensions["C"].width = 14
    run_control.column_dimensions["D"].width = 68

    _title(
        evidence,
        "Evidence Paste",
        "Paste source-bound values only. Natural-language notes cannot write numerical projections.",
        24,
    )
    _section_header(evidence, 4, 1, 5, "Official activity evidence")
    inactive_headers = ("TEAM", "PLAYER_OR_GSIS_ID", "STATUS", "SOURCE_URL", "OBSERVED_AT")
    for column, value in enumerate(inactive_headers, start=1):
        evidence.cell(5, column, value)
    for row in range(6, 106):
        for column in range(1, 6):
            evidence.cell(row, column).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    _table(evidence, "A5:E105", "OfficialStatusTable")
    status_validation = DataValidation(type="list", formula1='"ACTIVE,INACTIVE"', allow_blank=True)
    evidence.add_data_validation(status_validation)
    status_validation.add("C6:C105")

    _section_header(evidence, 4, 7, 10, "Payout tiers")
    payout_headers = ("rank_start", "rank_end", "prize_type", "value")
    for column, value in enumerate(payout_headers, start=7):
        evidence.cell(5, column, value)
    for row in range(6, 106):
        for column in range(7, 11):
            evidence.cell(row, column).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    _table(evidence, "G5:J105", "PayoutInputTable")
    prize_validation = DataValidation(type="list", formula1='"CASH,TICKET"', allow_blank=True)
    evidence.add_data_validation(prize_validation)
    prize_validation.add("I6:I105")
    evidence["G108"] = "Entered payout value"
    evidence["J108"] = "=SUMPRODUCT((H6:H105-G6:G105+1)*J6:J105)"
    evidence["J108"].number_format = '"$"#,##0.00'

    _section_header(evidence, 4, 12, 15, "Ownership brackets")
    ownership_headers = ("DK_ID", "LOW", "BASE", "HIGH")
    for column, value in enumerate(ownership_headers, start=12):
        evidence.cell(5, column, value)
    for row in range(6, 106):
        for column in range(12, 16):
            evidence.cell(row, column).fill = PatternFill("solid", fgColor=PALE_YELLOW)
            if column >= 13:
                evidence.cell(row, column).number_format = "0.0%"
    _table(evidence, "L5:O105", "OwnershipBracketTable")

    _section_header(evidence, 4, 17, 23, "Manual market and weather evidence")
    market_headers = (
        "TEAM",
        "BOOK",
        "SPREAD",
        "TOTAL",
        "OBSERVED_AT",
        "WEATHER_STATE",
        "SOURCE_URL",
    )
    for column, value in enumerate(market_headers, start=17):
        evidence.cell(5, column, value)
    for row in range(6, 38):
        for column in range(17, 24):
            evidence.cell(row, column).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    _table(evidence, "Q5:W37", "MarketWeatherTable")
    evidence.freeze_panes = "A5"
    for column, width in {
        "A": 10,
        "B": 24,
        "C": 14,
        "D": 48,
        "E": 24,
        "G": 12,
        "H": 12,
        "I": 14,
        "J": 14,
        "L": 14,
        "M": 12,
        "N": 12,
        "O": 12,
        "Q": 10,
        "R": 18,
        "S": 12,
        "T": 12,
        "U": 24,
        "V": 20,
        "W": 48,
    }.items():
        evidence.column_dimensions[column].width = width

    _title(portfolio, "Portfolio", "Generated output only. Do not edit this sheet.", 16)
    portfolio_headers = (
        "Entry ID",
        "Mode",
        "Slot 1",
        "Slot 2",
        "Slot 3",
        "Slot 4",
        "Slot 5",
        "Slot 6",
        "Slot 7",
        "Slot 8",
        "Slot 9",
        "Salary",
        "Canonical Key",
        "Worst State",
        "Robust Net LCB",
        "Elite Probability",
    )
    for column, value in enumerate(portfolio_headers, start=1):
        portfolio.cell(4, column, value)
    for row in range(5, 25):
        for column in range(1, 17):
            portfolio.cell(row, column).fill = PatternFill("solid", fgColor="F2F2F2")
    _table(portfolio, "A4:P24", "PortfolioOutputTable")
    portfolio.freeze_panes = "C5"
    portfolio.column_dimensions["A"].width = 16
    portfolio.column_dimensions["B"].width = 12
    portfolio.column_dimensions["M"].width = 50
    portfolio.column_dimensions["N"].width = 18
    portfolio.column_dimensions["O"].width = 20
    portfolio.column_dimensions["P"].width = 18
    for column in range(3, 12):
        portfolio.column_dimensions[chr(64 + column)].width = 14

    _title(qa, "QA", "Registered quantitative triggers and binding blockers.", 7)
    qa_headers = ("Code", "Severity", "Trigger", "Threshold", "Blocking", "Message", "Disposition")
    for column, value in enumerate(qa_headers, start=1):
        qa.cell(4, column, value)
    for row in range(5, 45):
        for column in range(1, 8):
            qa.cell(row, column).fill = PatternFill("solid", fgColor="F2F2F2")
    _table(qa, "A4:G44", "QAFindingTable")
    qa.column_dimensions["A"].width = 34
    qa.column_dimensions["B"].width = 14
    qa.column_dimensions["C"].width = 24
    qa.column_dimensions["D"].width = 20
    qa.column_dimensions["E"].width = 12
    qa.column_dimensions["F"].width = 90
    qa.column_dimensions["G"].width = 18
    qa.conditional_formatting.add(
        "E5:E44",
        FormulaRule(formula=["E5=TRUE"], fill=PatternFill("solid", fgColor=PALE_RED)),
    )
    qa.freeze_panes = "A5"

    _title(
        upload,
        "Upload",
        "Only RELEASE_DECISION=CERTIFIED_UPLOAD_PACKAGE may be manually uploaded to DraftKings.",
        6,
    )
    upload_rows = [
        ("Status", "DO_NOT_UPLOAD"),
        ("FILE_VALID", False),
        ("EVIDENCE_STATE", "UNKNOWN"),
        ("MODEL_STATUS", "UNVALIDATED"),
        ("RELEASE_DECISION", "DO_NOT_UPLOAD"),
        ("Certification basis", "MANUAL_GUARDRAIL"),
        ("Output CSV", ""),
        ("SHA-256", ""),
        ("Proposed SHA-256", ""),
        ("Manifest", ""),
        ("Created at", ""),
        ("Operator action", "Resolve every blocker, rerun, review, then upload manually."),
    ]
    for row_number, (label, value) in enumerate(upload_rows, start=4):
        upload.cell(row_number, 1, label).font = Font(bold=True)
        upload.cell(row_number, 2, value)
    upload["B4"].fill = PatternFill("solid", fgColor=PALE_RED)
    upload["B4"].font = Font(bold=True, color="9C0006")
    upload.column_dimensions["A"].width = 22
    upload.column_dimensions["B"].width = 100

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                font = copy(cell.font)
                font.name = font.name or "Aptos"
                cell.font = font
                alignment = copy(cell.alignment)
                alignment.vertical = alignment.vertical or "top"
                cell.alignment = alignment
        ws.auto_filter.ref = None
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


def populate_operator_run_control(
    path: str | Path, values: Mapping[str, str | float | int | bool | None]
) -> Path:
    target = Path(path).resolve()
    if not target.exists():
        create_operator_input_workbook(target)
    if not workbook_is_closed(target):
        raise WorkbookLockedError(
            f"Close {target.name} in Excel, then rerun. No workbook values were changed."
        )
    workbook = load_workbook(target)
    if workbook.sheetnames != ["Run Control", "Evidence Paste", "Portfolio", "QA", "Upload"]:
        raise ValueError("operator workbook sheet contract does not match the engine")
    run_control = workbook["Run Control"]
    rows_by_key = {
        str(run_control.cell(row, 1).value): row
        for row in range(5, run_control.max_row + 1)
        if run_control.cell(row, 1).value
    }
    unknown = sorted(set(values).difference(rows_by_key))
    if unknown:
        raise ValueError(f"unknown Run Control keys: {unknown}")
    for key, value in values.items():
        _set_display(run_control.cell(rows_by_key[key], 2), value)
    workbook.save(target)
    return target


def create_cowork_status_workbook(
    *,
    output_path: str | Path,
    run_values: Mapping[str, str | float | int | bool | None],
    blockers: Iterable[str],
    report_path: str | Path,
    truth_values: Mapping[str, bool | str] | None = None,
    readable_review: Mapping[str, object] | None = None,
) -> Path:
    target = create_operator_input_workbook(output_path)
    populate_operator_run_control(target, run_values)
    workbook = load_workbook(target)
    run_control = workbook["Run Control"]
    run_control["A2"] = (
        "Generated by the Cowork workflow. This is a review artifact; "
        "use the machine-readable run request for reruns."
    )
    for row in range(5, run_control.max_row + 1):
        run_control.cell(row, 2).fill = PatternFill("solid", fgColor="F2F2F2")
    qa = workbook["QA"]
    blocker_values = tuple(blockers)
    for row_number, blocker in enumerate(blocker_values, start=5):
        code, _, message = blocker.partition(":")
        values = (
            code,
            "CRITICAL",
            "MISSING",
            "PASS",
            True,
            message.strip() or blocker,
            "BLOCK",
        )
        for column, value in enumerate(values, start=1):
            _set_display(qa.cell(row_number, column), value)
    upload = workbook["Upload"]
    truths = {
        "FILE_VALID": False,
        "EVIDENCE_STATE": "UNKNOWN",
        "MODEL_STATUS": "UNVALIDATED",
        "RELEASE_DECISION": "DO_NOT_UPLOAD",
        "certification_basis": "MANUAL_GUARDRAIL",
        **dict(truth_values or {}),
    }
    upload["B4"] = "DO_NOT_UPLOAD"
    upload["B5"] = truths["FILE_VALID"]
    upload["B6"] = truths["EVIDENCE_STATE"]
    upload["B7"] = truths["MODEL_STATUS"]
    upload["B8"] = truths["RELEASE_DECISION"]
    upload["B9"] = truths["certification_basis"]
    upload["B4"].fill = PatternFill("solid", fgColor=PALE_RED)
    upload["B4"].font = Font(bold=True, color="9C0006")
    _set_display(upload["B13"], str(Path(report_path).resolve()))
    upload["B14"] = datetime.now(timezone.utc).isoformat()
    upload["B15"] = "DO NOT UPLOAD. Resolve the QA blockers and rerun the Cowork request."
    if readable_review is not None:
        _populate_readable_review(workbook, readable_review)
        next_action = readable_review.get("next_action")
        if next_action:
            _set_display(upload["B15"], next_action)
    workbook.save(target)
    return target


def create_review_workbook(
    *,
    staged_input: str | Path,
    output_path: str | Path,
    lineups: Mapping[str, Lineup],
    qa_findings: Iterable[QAFinding],
    manifest: CertificationManifest,
    manifest_path: str | Path | None = None,
    portfolio_metrics: Mapping[str, float | str] | None = None,
) -> Path:
    source = Path(staged_input).resolve()
    target = Path(output_path).resolve()
    if not workbook_is_closed(source):
        raise WorkbookLockedError(
            f"Close {source.name} in Excel, then rerun. No review workbook was written."
        )
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)
    workbook = load_workbook(target)
    portfolio = workbook["Portfolio"]
    metrics = portfolio_metrics or {}
    for row_number, (entry_id, lineup) in enumerate(sorted(lineups.items()), start=5):
        roster = list(lineup.roster) + [""] * (9 - len(lineup.roster))
        values = [
            entry_id,
            lineup.mode.value,
            *roster,
            lineup.salary,
            lineup.canonical_key,
            metrics.get("worst_state", ""),
            metrics.get("robust_net_payout_lcb", ""),
            metrics.get("elite_probability", ""),
        ]
        for column, value in enumerate(values, start=1):
            _set_display(portfolio.cell(row_number, column), value)
        portfolio.cell(row_number, 12).number_format = '"$"#,##0'
        portfolio.cell(row_number, 15).number_format = '"$"#,##0.00'
        portfolio.cell(row_number, 16).number_format = "0.00%"
    qa_sheet = workbook["QA"]
    for row_number, finding in enumerate(qa_findings, start=5):
        values = (
            finding.code,
            finding.severity,
            str(finding.trigger_value),
            str(finding.threshold),
            finding.blocking,
            finding.message,
            "BLOCK" if finding.blocking else "REVIEW",
        )
        for column, value in enumerate(values, start=1):
            _set_display(qa_sheet.cell(row_number, column), value)
    upload = workbook["Upload"]
    upload["B4"] = manifest.status
    upload["B5"] = manifest.file_valid
    upload["B6"] = manifest.evidence_state.value
    upload["B7"] = manifest.model_status.value
    upload["B8"] = manifest.release_decision.value
    upload["B9"] = manifest.certification_basis.value
    upload["B10"] = manifest.output_path or ""
    upload["B11"] = manifest.output_sha256 or ""
    upload["B12"] = manifest.proposed_output_sha256 or ""
    upload["B13"] = str(Path(manifest_path).resolve()) if manifest_path else ""
    upload["B14"] = manifest.created_at.astimezone(timezone.utc).isoformat()
    if manifest.release_decision.value == "CERTIFIED_UPLOAD_PACKAGE":
        upload["B4"].fill = PatternFill("solid", fgColor=PALE_GREEN)
        upload["B4"].font = Font(bold=True, color="006100")
        upload["B15"] = "Review exact Entry IDs and lineups, then upload manually in DraftKings."
    else:
        upload["B4"].fill = PatternFill("solid", fgColor=PALE_RED)
        upload["B15"] = "DO NOT UPLOAD. Resolve blockers shown on QA and rerun."
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(target)
    return target


def timestamped_review_path(output_dir: str | Path, run_id: str) -> Path:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return Path(output_dir).resolve() / f"NFL_DFS_Review_{run_id}_{stamp}.xlsx"


def read_operator_input(path: str | Path) -> OperatorInput:
    source = Path(path).resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    if not workbook_is_closed(source):
        raise WorkbookLockedError(
            f"Close {source.name} in Excel, then run again. The engine will not read a changing workbook."
        )
    workbook = load_workbook(source, read_only=True, data_only=False)
    if workbook.sheetnames != ["Run Control", "Evidence Paste", "Portfolio", "QA", "Upload"]:
        raise ValueError("operator workbook sheet contract does not match the engine")
    run_control = workbook["Run Control"]
    values: dict[str, str | float | int | None] = {}
    for row in range(5, run_control.max_row + 1):
        key = run_control.cell(row, 1).value
        if key:
            values[str(key)] = run_control.cell(row, 2).value
    evidence = workbook["Evidence Paste"]

    def populated_rows(start_col: int, width: int, start_row: int, end_row: int):
        rows: list[tuple[object, ...]] = []
        for row in range(start_row, end_row + 1):
            values_row = tuple(
                evidence.cell(row, column).value
                for column in range(start_col, start_col + width)
            )
            if any(value not in (None, "") for value in values_row):
                rows.append(values_row)
        return tuple(rows)

    return OperatorInput(
        values=values,
        official_status_rows=populated_rows(1, 5, 6, 105),
        payout_rows=populated_rows(7, 4, 6, 105),
        ownership_rows=populated_rows(12, 4, 6, 105),
    )
