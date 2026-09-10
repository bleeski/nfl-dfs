"""Regressions found in the 2026-09-10 pre-slate review of the Cowork rerun path.

Each test here reproduces a defect a reviewer confirmed against the real NE@SEA
run folder and pins the repair:

- a reloaded ``--request`` plus a fresh ``--input-dir`` built from the stale
  snapshot and said nothing;
- ``--exclude`` on a rerun replaced the saved fade list instead of adding to it;
- a reused ``--run-id`` rewrote the earlier run's ``cowork_run.json``;
- ``--lineup-count`` below the reserved-entry count exported duplicate rosters
  and only the display step caught it, leaving the CSV on disk;
- the legacy distinct-captain rule hard-failed once every selectable person had
  captained a lineup;
- a frozen package bound to older salary bytes blocked even with
  ``build_priors`` set;
- the raw label was used in the export filename;
- the review surface hid the sole-kicker assumption and pool coverage.
"""

from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

from nfl_dfs import priors as priors_module
from nfl_dfs.cowork import CoworkRunRequest, resolve_request_inputs
from nfl_dfs.hashing import sha256_file
from nfl_dfs.prior_review import run_prior_review
from nfl_dfs.readable_review import ReadableReviewError, create_readable_review

from .test_participation import _POOL, _slate
from .test_prior_review_profile import AS_OF, _attachments, _cowork_args, _prepared_run
from .test_prior_selection import _entries_bytes


# --------------------------------------------------------------------------- #
# Request rerun semantics
# --------------------------------------------------------------------------- #


def test_fresh_input_dir_supersedes_reloaded_request_snapshots(tmp_path: Path) -> None:
    salary_path, entry_path, _package, _project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    stale = tmp_path / "stale"
    stale.mkdir()
    stale_salary = stale / "salary.csv"
    stale_salary.write_bytes(salary_path.read_bytes().replace(b",OUT\n", b",Q\n", 1))
    stale_entries = stale / "entries.csv"
    stale_entries.write_bytes(entry_path.read_bytes())
    assert sha256_file(stale_salary) != sha256_file(salary_path)
    request = CoworkRunRequest(
        profile="prior_review", salary_csv=str(stale_salary), entry_csv=str(stale_entries)
    )
    fresh = _attachments(tmp_path, salary_path, entry_path)

    # Reloaded request alone: the recorded snapshots stay authoritative.
    unchanged, _ = resolve_request_inputs(request, allowed_roots=(tmp_path,))
    assert Path(unchanged.salary_csv) == stale_salary.resolve()

    # An operator-named attachment directory outranks the reloaded snapshot.
    resolved, _ = resolve_request_inputs(request, input_dir=fresh, allowed_roots=(tmp_path,))
    assert Path(resolved.salary_csv) == (fresh / "salary.csv").resolve()
    assert Path(resolved.entry_csv) == (fresh / "entries.csv").resolve()


def test_exclude_merges_and_run_id_collision_is_refused(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from nfl_dfs import prior_review as prior_review_module
    from nfl_dfs.dk import parse_salaries

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=datetime.now(timezone.utc) + timedelta(hours=6)
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")
    real = prior_review_module.run_prior_review
    monkeypatch.setattr(
        cli, "run_prior_review", lambda **kwargs: real(**kwargs, project=project)
    )
    slate = parse_salaries(salary_path)
    flex_ids = {
        player.name: player.dk_id for player in slate.players if player.role == "FLEX"
    }
    first_fade = flex_ids["Sea Alpha WR"]
    second_fade = flex_ids["Alpha WR"]

    first = cli.command_cowork_run(
        _cowork_args(
            tmp_path, attachments, run_id="rerun-a",
            prior_package_dir=str(package_dir), exclude=[first_fade],
        )
    )
    assert first == 0
    first_report_path = tmp_path / "outputs" / "rerun-a" / "cowork_run.json"
    first_report = json.loads(first_report_path.read_text(encoding="utf-8"))
    first_record_sha = sha256_file(first_report_path)
    request_path = first_report["request"]
    assert json.loads(Path(request_path).read_text())["exclude_dk_ids"] == [first_fade]
    assert first_report["superseded_request_inputs"] == {}

    # A follow-up fade on the reloaded request adds to the saved list.
    # An out-of-tree package directory is authorized on the command line, so a
    # rerun names it again; the request's own copy is not trusted on its own.
    second = cli.command_cowork_run(
        _cowork_args(
            tmp_path, attachments, run_id="rerun-b", request=request_path,
            input_dir=None, exclude=[second_fade], prior_package_dir=str(package_dir),
        )
    )
    assert second == 0
    second_report = json.loads(
        (tmp_path / "outputs" / "rerun-b" / "cowork_run.json").read_text(encoding="utf-8")
    )
    saved = json.loads(Path(second_report["request"]).read_text())["exclude_dk_ids"]
    assert saved == [first_fade, second_fade]
    rosters = {
        dk_id
        for lineup in second_report["prior_review_reports"]["selection"]["lineups"]
        for dk_id in lineup["roster"]
    }
    excluded_people = {"SEA|WR|Sea Alpha WR", "NE|WR|Alpha WR"}
    by_id = {player.dk_id: player for player in slate.players}
    assert not {by_id[dk_id].underlying_id for dk_id in rosters} & excluded_people
    coverage = second_report["prior_review_reports"]["selection"]["pool_coverage"]
    assert coverage["by_reason"]["OPERATOR_EXCLUDED"]["people"] == 2
    assert coverage["by_reason"]["DK_STATUS_UNAVAILABLE"]["people"] == 2

    # Reusing an existing run id is refused before anything is written, and the
    # earlier run's record is byte-identical afterwards.
    with pytest.raises(ValueError, match="RUN_ID_COLLISION"):
        cli.command_cowork_run(
            _cowork_args(
                tmp_path, attachments, run_id="rerun-a", request=request_path,
                input_dir=None, prior_package_dir=str(package_dir),
            )
        )
    assert sha256_file(first_report_path) == first_record_sha
    assert not (tmp_path / "outputs" / "rerun-a" / "cowork_diagnostic.json").exists()


def test_fresh_attachments_on_a_reloaded_request_are_reported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from nfl_dfs import prior_review as prior_review_module

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=datetime.now(timezone.utc) + timedelta(hours=6)
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")
    real = prior_review_module.run_prior_review
    monkeypatch.setattr(
        cli, "run_prior_review", lambda **kwargs: real(**kwargs, project=project)
    )
    assert cli.command_cowork_run(
        _cowork_args(tmp_path, attachments, run_id="first", prior_package_dir=str(package_dir))
    ) == 0
    request_path = json.loads(
        (tmp_path / "outputs" / "first" / "cowork_run.json").read_text(encoding="utf-8")
    )["request"]

    fresh = tmp_path / "fresh"
    fresh.mkdir()
    (fresh / "salary.csv").write_bytes(salary_path.read_bytes().replace(b",OUT\n", b",Q\n", 1))
    (fresh / "entries.csv").write_bytes(entry_path.read_bytes())
    code = cli.command_cowork_run(
        _cowork_args(
            tmp_path, fresh, run_id="second", request=request_path,
            prior_package_dir=str(package_dir),
        )
    )
    report = json.loads(
        (tmp_path / "outputs" / "second" / "cowork_run.json").read_text(encoding="utf-8")
    )
    # The fresh salary bytes were used (the package bound to the old bytes now
    # mismatches, which is the correct named stop without build_priors), and the
    # supersession is recorded rather than silent.
    assert code == 2
    assert report["blockers"][0].startswith("PRIOR_PACKAGE_SALARY_MISMATCH:")
    assert set(report["superseded_request_inputs"]) == {"salary_csv", "entry_csv"}
    assert sha256_file(fresh / "salary.csv") in report["input_hashes"].values()


# --------------------------------------------------------------------------- #
# Selection-side guards
# --------------------------------------------------------------------------- #


def test_lineup_count_below_reserved_entries_blocks_before_any_export(tmp_path: Path) -> None:
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, lineup_count=1, project=project,
    )
    assert outcome.blocked and outcome.stage == "SELECT"
    assert outcome.blockers[0].startswith("LINEUP_COUNT_BELOW_RESERVED_ENTRIES:")
    assert "bulk_entry_csv" not in outcome.artifacts
    assert not list((tmp_path / "out").rglob("*.csv")) if (tmp_path / "out").exists() else True

    surplus = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run-3", output_root=tmp_path / "out-3",
        prior_package_dir=package_dir, lineup_count=3, project=project,
    )
    assert not surplus.blocked
    summary = surplus.reports["selection"]["assignment_summary"]
    assert summary == {
        "lineups_generated": 3,
        "reserved_entries": 2,
        "unassigned_lineups": 1,
        "note": summary["note"],
    }
    assert "not written to any entry" in summary["note"]


def test_captain_exhaustion_falls_back_to_repeats_instead_of_failing(tmp_path: Path) -> None:
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    # 16 people in the pool, 14 selectable: more reserved entries than that
    # used to end in SOLVER_RETURNED_NO_LINEUP once the captains ran out.
    entry_ids = tuple(str(900000101 + index) for index in range(20))
    entry_path.write_bytes(_entries_bytes(entry_ids))
    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, project=project,
    )
    assert not outcome.blocked, outcome.blockers
    lineups = outcome.reports["selection"]["lineups"]
    assert len(lineups) == 20
    assert len({lineup["canonical_key"] for lineup in lineups}) == 20
    differentiation = outcome.reports["selection"]["selection"]["differentiation"]
    assert differentiation["captain"] == "DISTINCT_UNTIL_POOL_EXHAUSTED_THEN_REPEATED"
    repeats_from = differentiation["captain_repeats_from_index"]
    assert 2 <= repeats_from <= 15
    captains = [lineup["captain_dk_id"] for lineup in lineups]
    assert len(set(captains[: repeats_from - 1])) == repeats_from - 1
    assert sum(differentiation["captain_exposure"].values()) == 20


def test_salary_mismatched_package_rebuilds_only_with_permission(tmp_path: Path) -> None:
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    salary_path.write_bytes(salary_path.read_bytes().replace(b",OUT\n", b",Q\n", 1))

    blocked = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run-a", output_root=tmp_path / "out-a",
        prior_package_dir=package_dir, project=project,
    )
    assert blocked.blocked
    assert blocked.blockers[0].startswith("PRIOR_PACKAGE_SALARY_MISMATCH:")

    calls: list[dict] = []

    def propose(**kwargs):
        calls.append(kwargs)
        raise priors_module.PriorsBuildError("SOURCE_EMPTY:games")

    rebuilt = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run-b", output_root=tmp_path / "out-b",
        prior_package_dir=package_dir, build_priors=True, propose=propose, project=project,
    )
    assert len(calls) == 1
    statuses = [(stage["stage"], stage["status"]) for stage in rebuilt.stages]
    assert ("PRIORS", "REBUILDING_SALARY_MISMATCHED_PACKAGE") in statuses
    assert rebuilt.reports["superseded_package"]["reason"].startswith(
        "PRIOR_PACKAGE_SALARY_MISMATCH:"
    )
    assert rebuilt.blockers[0].startswith("PRIORS_PROPOSE_FAILED:PriorsBuildError:")


def test_a_tampered_frozen_artifact_is_never_rebuilt_over(tmp_path: Path) -> None:
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    tampered = package_dir / "player_prior.json"
    tampered.write_text(tampered.read_text(encoding="utf-8") + " ", encoding="utf-8")
    calls: list[dict] = []

    def propose(**kwargs):
        calls.append(kwargs)
        raise AssertionError("propose must not run on a hash mismatch")

    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, build_priors=True, propose=propose, project=project,
    )
    assert outcome.blocked
    assert outcome.blockers[0].startswith("PRIOR_ARTIFACT_HASH_MISMATCH:player_prior.json:")
    assert calls == []


def test_export_filename_uses_a_sanitized_label(tmp_path: Path) -> None:
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne/sea wk1 ..", as_of=AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, project=project,
    )
    assert not outcome.blocked
    exported = Path(outcome.artifacts["bulk_entry_csv"])
    assert exported.parent == (tmp_path / "out" / "review").resolve()
    assert exported.name == "DK_REVIEW_ENTRY_ne-sea-wk1---.csv"


# --------------------------------------------------------------------------- #
# Review surface
# --------------------------------------------------------------------------- #


def test_review_surface_shows_pool_coverage_and_the_kicker_assumption(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from nfl_dfs import prior_review as prior_review_module
    from openpyxl import load_workbook

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=datetime.now(timezone.utc) + timedelta(hours=6)
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")
    real = prior_review_module.run_prior_review
    monkeypatch.setattr(
        cli, "run_prior_review", lambda **kwargs: real(**kwargs, project=project)
    )
    assert cli.command_cowork_run(
        _cowork_args(tmp_path, attachments, prior_package_dir=str(package_dir))
    ) == 0
    report = json.loads(
        (tmp_path / "outputs" / "prior-review-test" / "cowork_run.json").read_text(encoding="utf-8")
    )
    readable = report["prior_review_reports"]["readable_review"]

    coverage = readable["pool_coverage"]
    assert coverage["people_in_pool"] == len(_POOL)
    assert coverage["selectable_people"] == 14
    unavailable = coverage["by_reason"]["DK_STATUS_UNAVAILABLE"]
    assert unavailable["people"] == 2
    assert unavailable["flex_salary"] == 8200 + 1200
    assert unavailable["cpt_salary"] == 12300 + 1800
    reasons = {row["name"]: row["reason"] for row in coverage["excluded_people"]}
    assert reasons == {"Sea Lead RB": "DK_STATUS_UNAVAILABLE:OUT", "Sea Hurt WR": "DK_STATUS_UNAVAILABLE:IR"}
    exposure_rows = {row["name"]: row for row in readable["exposure"]["people"]}
    assert exposure_rows["Sea Lead RB"]["excluded"] is True
    assert exposure_rows["Sea Lead RB"]["exclusion_source"] == "DK_STATUS_UNAVAILABLE:OUT"
    assert exposure_rows["Starter QB"]["excluded"] is False

    kicker_slots = [
        slot for entry in readable["entries"] for slot in entry["slots"] if slot["position"] == "K"
    ]
    assert kicker_slots, "the synthetic pool always selects a kicker"
    for slot in kicker_slots:
        assert slot["role_evidence_state"] == "PRIOR_ONLY_SOLE_LISTED_ASSUMPTION"
        assert any(
            str(item["finding"]).startswith("SOLE_LISTED_KICKER_ASSUMPTION:")
            for item in slot["role_findings"]
        )
    categories = {row["category"] for row in readable["evidence_observations"]}
    assert {"kicker_role", "unallocated_volume"} <= categories

    html = Path(report["prior_review_artifacts"]["readable_review_html"]).read_text(encoding="utf-8")
    assert "Pool coverage" in html and "SOLE_LISTED_KICKER_ASSUMPTION" in html
    workbook = load_workbook(report["review_workbook"])
    assert workbook.sheetnames == [
        "Run Control", "Evidence Paste", "Portfolio", "QA", "Upload",
        "Exposure", "Review Evidence", "Artifacts",
    ]
    exposure_text = {
        cell.value for row in workbook["Exposure"].iter_rows() for cell in row if isinstance(cell.value, str)
    }
    assert any(value.startswith("Pool coverage:") for value in exposure_text)
    assert "DK_STATUS_UNAVAILABLE" in exposure_text


def test_pool_coverage_is_reconciled_against_exact_salary_bytes(tmp_path: Path) -> None:
    from nfl_dfs.dk import parse_entries, parse_salaries

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, project=project,
    )
    assert not outcome.blocked
    slate = parse_salaries(salary_path)
    template = parse_entries(entry_path)
    truths = {
        "FILE_VALID": True, "EVIDENCE_STATE": "UNKNOWN",
        "MODEL_STATUS": "PRIOR_ONLY", "RELEASE_DECISION": "DO_NOT_UPLOAD",
    }

    def attempt(name: str, hashes):
        return create_readable_review(
            slate=slate, template=template, salary_path=salary_path, entry_path=entry_path,
            assignment_path=outcome.artifacts["assignments"],
            exported_path=outcome.artifacts["bulk_entry_csv"],
            artifacts=outcome.artifacts, expected_hashes=hashes, reports=outcome.reports,
            truth_values=truths, blockers=(), next_action="Rerun.",
            output_dir=tmp_path / name, package_root=tmp_path,
        )

    assert attempt("clean", outcome.hashes).data["pool_coverage"]["selectable_people"] == 14

    selection_path = Path(outcome.artifacts["selection_report"])
    original = selection_path.read_bytes()
    record = json.loads(original)
    record["pool_coverage"]["people"][0]["flex_salary"] += 100
    selection_path.write_text(json.dumps(record), encoding="utf-8")
    with pytest.raises(ReadableReviewError, match="POOL_COVERAGE_SALARY_MISMATCH"):
        attempt("salary", {**outcome.hashes, "selection_report": sha256_file(selection_path)})

    record = json.loads(original)
    record["pool_coverage"]["by_reason"]["SELECTABLE"]["people"] -= 1
    selection_path.write_text(json.dumps(record), encoding="utf-8")
    with pytest.raises(ReadableReviewError, match="POOL_COVERAGE_TOTAL_MISMATCH"):
        attempt("total", {**outcome.hashes, "selection_report": sha256_file(selection_path)})
    selection_path.write_bytes(original)


def test_readable_review_failure_withholds_the_artifact_index_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from nfl_dfs import prior_review as prior_review_module

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=datetime.now(timezone.utc) + timedelta(hours=6)
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")
    real = prior_review_module.run_prior_review
    monkeypatch.setattr(
        cli, "run_prior_review", lambda **kwargs: real(**kwargs, project=project)
    )

    def failing(**kwargs):
        raise ReadableReviewError("READABLE_REVIEW_SELECTION_SALARY_MISMATCH:test")

    monkeypatch.setattr(cli, "create_readable_review", failing)
    code = cli.command_cowork_run(
        _cowork_args(tmp_path, attachments, prior_package_dir=str(package_dir))
    )
    assert code == 2
    report = json.loads(
        (tmp_path / "outputs" / "prior-review-test" / "cowork_run.json").read_text(encoding="utf-8")
    )
    assert report["FILE_VALID"] is False
    assert report["bulk_entry_csv"] is None and report["bulk_entry_sha256"] is None
    assert "bulk_entry_csv" not in report["prior_review_artifacts"]
    assert "bulk_entry_csv" not in report["prior_review_hashes"]
    withheld = report["prior_review_reports"]["readable_review_failure"]["withheld_artifacts"]["bulk_entry_csv"]
    assert Path(withheld["path"]).is_file()
    assert sha256_file(withheld["path"]) == withheld["sha256"]
    marker = tmp_path / "outputs" / "prior-review-test" / "review" / "READABLE_REVIEW_FAILED.json"
    assert marker.is_file()
    assert json.loads(marker.read_text())["FILE_VALID"] is False


def test_supplied_official_status_names_uncovered_selected_people_and_workbook_names_the_csv(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from nfl_dfs import prior_review as prior_review_module
    from nfl_dfs.dk import parse_salaries
    from openpyxl import load_workbook

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=datetime.now(timezone.utc) + timedelta(hours=6)
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")
    real = prior_review_module.run_prior_review
    monkeypatch.setattr(
        cli, "run_prior_review", lambda **kwargs: real(**kwargs, project=project)
    )
    slate = parse_salaries(salary_path)
    # One ACTIVE row for a person who will certainly be selected, nothing else.
    alpha = next(p for p in slate.players if p.name == "Sea Alpha WR" and p.role == "FLEX")
    observed = (datetime.now(timezone.utc) - timedelta(minutes=10)).isoformat()
    status_dir = tmp_path / "status"
    status_dir.mkdir()
    status_path = status_dir / "official.csv"
    status_path.write_text(
        "TEAM,PLAYER_OR_GSIS_ID,STATUS,SOURCE_URL,OBSERVED_AT\n"
        f"SEA,{alpha.dk_id},ACTIVE,https://www.nfl.com/injuries/,{observed}\n",
        encoding="utf-8",
    )
    code = cli.command_cowork_run(
        _cowork_args(
            tmp_path, attachments, prior_package_dir=str(package_dir),
            official_status_csv=str(status_path),
        )
    )
    assert code == 0
    report = json.loads(
        (tmp_path / "outputs" / "prior-review-test" / "cowork_run.json").read_text(encoding="utf-8")
    )
    # Presence of a file no longer reads as coverage: the uncovered selected
    # people are named first, and OFFICIAL_STATUS_REQUIRED is not claimed.
    assert report["blockers"][0].startswith("OFFICIAL_STATUS_INCOMPLETE_FOR_SELECTED: ")
    assert not any(value.startswith("OFFICIAL_STATUS_REQUIRED:") for value in report["blockers"])
    coverage = report["prior_review_reports"]["selection"]["official_status_coverage"]
    assert coverage["selected_with_row"] == 1
    assert coverage["selected_people"] == coverage["selected_with_row"] + len(
        coverage["selected_without_row"]
    )
    saved = json.loads(Path(report["request"]).read_text())
    assert saved["official_status_csv"] is not None and Path(saved["official_status_csv"]).is_file()

    # The Upload sheet names the exact reviewed CSV and hash, marked review-only.
    upload = load_workbook(report["review_workbook"])["Upload"]
    assert str(upload["B10"].value).startswith("REVIEW ONLY (not certified): ")
    assert upload["B11"].value == report["bulk_entry_sha256"]
    assert upload["B8"].value == "DO_NOT_UPLOAD"


def test_reused_frozen_package_reports_its_inherited_weather_basis(tmp_path: Path) -> None:
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=AS_OF + timedelta(hours=6)
    )
    team_prior = package_dir / "team_prior.json"
    payload = json.loads(team_prior.read_text(encoding="utf-8"))
    payload["metadata"]["coverage"]["weather_basis"] = (
        "OPERATOR_SUPPLIED:roof=outdoors|OPERATOR_CAPTURE:"
        "https://api.weather.gov/gridpoints/SEW/125,67/forecast:observed_at=2026-09-08T14:00:00+00:00"
    )
    payload["records"] = [{"team": "NE", "weather_state": "RAIN"}, {"team": "SEA", "weather_state": "RAIN"}]
    team_prior.write_text(json.dumps(payload, sort_keys=True), encoding="utf-8")
    package = json.loads((package_dir / "prior_package.json").read_text(encoding="utf-8"))
    package["artifacts"]["team_prior.json"] = sha256_file(team_prior)
    (package_dir / "prior_package.json").write_text(json.dumps(package, indent=2, sort_keys=True), encoding="utf-8")

    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, project=project,
    )
    assert not outcome.blocked, outcome.blockers
    weather = outcome.reports["weather"]
    assert weather["basis"].startswith("INHERITED_FROM_FROZEN_PACKAGE:")
    assert weather["weather_state"] == "RAIN"
    assert weather["source_uri"] == "https://api.weather.gov/gridpoints/SEW/125,67/forecast"
    assert weather["observed_at"] == "2026-09-08T14:00:00+00:00"
