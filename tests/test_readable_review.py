"""SD5 exact-byte readable review, safe rendering, and replay coverage."""

from __future__ import annotations

import copy
import csv
import json
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nfl_dfs.dk import parse_entries, parse_salaries
from nfl_dfs.hashing import sha256_file
from nfl_dfs.lineups import validate_lineup
from nfl_dfs.review_export import export_review_entries, write_assignments_csv, write_run_record
from nfl_dfs.portfolio_policy import (
    portfolio_policy_template,
    validate_portfolio_policy_file,
    write_normalized_portfolio_policy,
)
from nfl_dfs.prior_review import run_prior_review
from nfl_dfs.readable_review import (
    ReadableReviewError,
    create_readable_review,
    verify_readable_review_artifacts,
)
from nfl_dfs.workbook import create_cowork_status_workbook

from .test_prior_review_profile import _prepared_run
from .test_prior_selection import _entries_bytes
from .test_participation import _POOL, _slate


TRUTHS = {
    "FILE_VALID": True,
    "EVIDENCE_STATE": "UNKNOWN",
    "MODEL_STATUS": "PRIOR_ONLY",
    "RELEASE_DECISION": "DO_NOT_UPLOAD",
    "certification_basis": "MODEL_ASSISTED",
}


def _policy_outcome(
    root: Path,
    *,
    entry_count: int = 2,
    hostile_contest: bool = False,
    missing_contest: bool = False,
):
    now = datetime.now(timezone.utc)
    salary_path, entry_path, package_dir, project = _prepared_run(
        root, expires_at=now + timedelta(hours=6)
    )
    entry_ids = tuple(str(900000001 + index) for index in range(entry_count))
    entry_bytes = _entries_bytes(entry_ids)
    if hostile_contest:
        entry_bytes = entry_bytes.replace(
            b"Test Showdown (NE @ SEA)", b"<script>alert(1)</script> =review"
        )
    elif missing_contest:
        entry_bytes = entry_bytes.replace(b"Test Showdown (NE @ SEA)", b"")
    entry_path.write_bytes(entry_bytes)
    slate = parse_salaries(salary_path)
    template = parse_entries(entry_path)
    source_policy = root / "portfolio_policy.json"
    source_policy.write_text(
        json.dumps(
            portfolio_policy_template(
                slate,
                entry_ids,
                controls={"max_pairwise_person_overlap": 5},
            )
        ),
        encoding="utf-8",
    )
    validation = validate_portfolio_policy_file(
        source_policy, slate=slate, entry_ids=entry_ids
    )
    assert validation.valid and validation.policy is not None
    normalized = write_normalized_portfolio_policy(
        root / "portfolio_policy.normalized.json", validation.policy
    )
    outcome = run_prior_review(
        salary_csv=salary_path,
        entry_csv=entry_path,
        label=f"fixture-{entry_count}",
        as_of=now,
        run_root=root / "run",
        output_root=root / "output",
        prior_package_dir=package_dir,
        portfolio_policy=validation.policy,
        portfolio_policy_source_path=source_policy,
        portfolio_policy_source_sha256=validation.source_sha256,
        portfolio_policy_normalized_path=normalized,
        portfolio_policy_normalized_sha256=validation.policy.normalized_sha256,
        project=project,
    )
    assert outcome.file_valid, outcome.as_report()
    return slate, template, salary_path, entry_path, outcome


def _create(
    root: Path,
    *,
    entry_count: int = 2,
    hostile_contest: bool = False,
    missing_contest: bool = False,
):
    slate, template, salary_path, entry_path, outcome = _policy_outcome(
        root,
        entry_count=entry_count,
        hostile_contest=hostile_contest,
        missing_contest=missing_contest,
    )
    blockers = (
        "OFFICIAL_STATUS_REQUIRED: missing exact current activity evidence",
        "ROLE_ASSUMPTION_UNSUPPORTED: multiple named blockers remain",
        "<script>alert('blocker')</script>",
    )
    readable = create_readable_review(
        slate=slate,
        template=template,
        salary_path=salary_path,
        entry_path=entry_path,
        assignment_path=outcome.artifacts["assignments"],
        exported_path=outcome.artifacts["bulk_entry_csv"],
        artifacts=outcome.artifacts,
        expected_hashes=outcome.hashes,
        reports=outcome.reports,
        truth_values=TRUTHS,
        blockers=blockers,
        next_action="Refresh exact official activity evidence and rerun.",
        output_dir=root / "output" / "review",
        package_root=root,
    )
    return slate, template, salary_path, entry_path, outcome, readable, blockers


def test_exact_lineups_exposures_truths_and_html_are_readable_and_inert(tmp_path: Path) -> None:
    _slate, template, _salary, _entries, _outcome, readable, blockers = _create(
        tmp_path, hostile_contest=True
    )
    data = readable.data
    assert data["reconciliation"]["status"] == "PASS"
    assert data["truths"] == {key: TRUTHS[key] for key in TRUTHS if key != "certification_basis"}
    assert [entry["entry_id"] for entry in data["entries"]] == [
        entry.entry_id for entry in template.authorizations
    ]
    for entry in data["entries"]:
        assert [slot["slot"] for slot in entry["slots"]] == [
            "CPT", "FLEX 1", "FLEX 2", "FLEX 3", "FLEX 4", "FLEX 5"
        ]
        assert all(slot["dk_roster_id"].isdigit() for slot in entry["slots"])
        assert len({slot["underlying_person_id"] for slot in entry["slots"]}) == 6
        assert sum(slot["salary"] for slot in entry["slots"]) == entry["salary_total"]
        assert entry["salary_total"] + entry["salary_remaining"] == 50_000
        assert abs(
            sum(slot["prior_only_central_estimate_points"] for slot in entry["slots"])
            - entry["prior_only_central_estimate_points"]
        ) < 0.01
    exposure = data["exposure"]
    assert exposure["entry_count_denominator"] == 2
    assert exposure["canonical_uniqueness"] == "PASS"
    assert exposure["configured_pairwise_person_overlap"] == 5
    assert exposure["effective_pairwise_person_overlap"] == 5
    assert any(row["combined_count"] == 0 for row in exposure["people"])
    assert any(row["combined_count"] == 1 for row in exposure["people"])
    assert any(row["captain_count"] >= 1 for row in exposure["people"])
    assert all(
        row["combined_percentage"] == round(100 * row["combined_count"] / 2, 3)
        and row["captain_percentage"] == round(100 * row["captain_count"] / 2, 3)
        for row in exposure["people"]
    )
    assert blockers == tuple(data["blockers"])
    assert verify_readable_review_artifacts(
        json_path=readable.json_path,
        json_sha256=readable.json_sha256,
        html_path=readable.html_path,
        html_sha256=readable.html_sha256,
    ) == ()
    html_text = Path(readable.html_path).read_text(encoding="utf-8")
    assert "<script>alert" not in html_text
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html_text
    assert "PRIOR_ONLY central estimates" in html_text
    assert "DO_NOT_UPLOAD" in html_text
    assert all(label in html_text for label in data["truths"])
    assert "Refresh exact official activity evidence and rerun." in html_text
    assert readable.json_sha256 in html_text
    assert {
        "salary_csv", "entry_csv", "portfolio_policy_source",
        "portfolio_policy_normalized", "assignments", "bulk_entry_csv",
    }.issubset(data["hashes"])
    assert {
        "team_projections", "player_opportunities", "source_ledger", "team_splits",
    }.issubset({row["name"] for row in data["artifacts"]})
    states = {(row["category"], row["state"]) for row in data["evidence_observations"]}
    assert ("official_activity", "MISSING") in states
    assert any(category == "offensive_role" for category, _state in states)


def test_missing_contest_and_stale_source_state_remain_visible(tmp_path: Path) -> None:
    slate, template, salary_path, entry_path, outcome = _policy_outcome(
        tmp_path, missing_contest=True
    )
    reports = copy.deepcopy(outcome.reports)
    reports["projection"]["freshness_state"] = "STALE"
    readable = create_readable_review(
        slate=slate,
        template=template,
        salary_path=salary_path,
        entry_path=entry_path,
        assignment_path=outcome.artifacts["assignments"],
        exported_path=outcome.artifacts["bulk_entry_csv"],
        artifacts=outcome.artifacts,
        expected_hashes=outcome.hashes,
        reports=reports,
        truth_values=TRUTHS,
        blockers=("SOURCE_STALE: refresh the exact source package",),
        next_action="Refresh the exact source package and rerun.",
        output_dir=tmp_path / "readable",
        package_root=tmp_path,
    )
    assert all(entry["contest_name"] is None for entry in readable.data["entries"])
    assert any(
        row["category"] == "projection" and row["state"] == "STALE"
        for row in readable.data["evidence_observations"]
    )
    html_text = Path(readable.html_path).read_text(encoding="utf-8")
    assert "Contest label unavailable" in html_text
    assert "SOURCE_STALE" in html_text


def test_workbook_prevents_formula_injection_without_changing_exact_ids(tmp_path: Path) -> None:
    _slate, _template, _salary, _entries, _outcome, readable, _blockers = _create(tmp_path)
    review = copy.deepcopy(readable.data)
    hostile = (
        "=1+1",
        "+cmd|' /C calc'!A0",
        "-2+3",
        "@SUM(1,1)",
        "\t=HYPERLINK(\"https://invalid.example\")",
        "\n=1+1",
        "<script>alert(1)</script>",
    )
    slots = [slot for entry in review["entries"] for slot in entry["slots"]]
    for slot, value in zip(slots[: len(hostile)], hostile, strict=True):
        slot["name"] = value
    review["entries"][0]["contest_name"] = "=unsafe contest"
    review["evidence_observations"][0]["source"] = "@unsafe source"
    review["evidence_observations"][0]["observation"] = "<script>unsafe</script>"
    review["artifacts"][0]["path"] = "+unsafe path"
    workbook_path = tmp_path / "review.xlsx"
    create_cowork_status_workbook(
        output_path=workbook_path,
        run_values={"RUN_LABEL": "=unsafe"},
        blockers=hostile,
        report_path=tmp_path / "cowork_run.json",
        truth_values=TRUTHS,
        readable_review=review,
    )
    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook.sheetnames[-3:] == ["Exposure", "Review Evidence", "Artifacts"]
    assert workbook["Run Control"]["B5"].value == "'=unsafe"
    player_cells = [workbook["Portfolio"].cell(row, 5) for row in range(5, 12)]
    for cell, original in zip(player_cells, hostile, strict=True):
        assert cell.data_type == "s"
        if original.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
            assert cell.value == "'" + original
        else:
            assert cell.value == original
    assert workbook["Portfolio"]["F5"].value.isdigit()
    assert not workbook["Portfolio"]["F5"].value.startswith("'")
    assert workbook["Portfolio"]["B5"].value == "'=unsafe contest"
    assert workbook["Review Evidence"]["F5"].value == "'@unsafe source"
    assert workbook["Review Evidence"]["C5"].value == "<script>unsafe</script>"
    assert workbook["Artifacts"]["B5"].value == "'+unsafe path"
    assert workbook["QA"]["A5"].data_type == "s"
    assert workbook["QA"]["A5"].value == "'=1+1"


def test_repeated_nonascii_names_keep_exact_role_and_person_identity(tmp_path: Path) -> None:
    repeated_name = "José O'Brien <script>"
    pool = tuple(
        (team, position, repeated_name if position == "QB" else name, status, salary)
        for team, position, name, status, salary in _POOL
    )
    slate = _slate(tmp_path, pool)
    salary_path = tmp_path / "DKSalaries.csv"
    entry_path = tmp_path / "DKEntries.csv"
    entry_path.write_bytes(_entries_bytes(("900000001", "900000002", "900000003")))
    template = parse_entries(entry_path)
    by_person: dict[str, dict[str, str]] = {}
    for player in slate.players:
        by_person.setdefault(player.underlying_id, {})[str(player.role)] = player.dk_id
    ne_qb = f"NE|QB|{repeated_name}"
    sea_qb = f"SEA|QB|{repeated_name}"
    others = [person for person in sorted(by_person) if person not in {ne_qb, sea_qb}]
    people = (ne_qb, sea_qb, *others[:5])
    first = (by_person[ne_qb]["CPT"], *(by_person[person]["FLEX"] for person in (sea_qb, *others[:4])))
    second = (by_person[ne_qb]["CPT"], *(by_person[person]["FLEX"] for person in (sea_qb, *others[:3], others[4])))
    third = (by_person[sea_qb]["CPT"], *(by_person[person]["FLEX"] for person in (ne_qb, *others[:4])))
    assignments = {
        template.authorizations[0].entry_id: first,
        template.authorizations[1].entry_id: second,
        template.authorizations[2].entry_id: third,
    }
    assignment_path = tmp_path / "run" / "assignments.csv"
    assignment_sha = write_assignments_csv(
        assignment_path,
        assignments,
        entry_order=tuple(assignments),
    )
    export = export_review_entries(
        slate=slate,
        template=template,
        assignments=assignments,
        output_path=tmp_path / "output" / "review" / "DK_REVIEW_ENTRY_names.csv",
    )
    assert export.file_valid
    score_map = {dk_id: float(index + 1) for index, dk_id in enumerate(sorted({*first, *second, *third}))}
    lineups = []
    for roster in assignments.values():
        validation = validate_lineup(slate, roster)
        assert validation.lineup is not None
        lineups.append(
            {
                "roster": list(roster),
                "salary": validation.lineup.salary,
                "canonical_key": validation.lineup.canonical_key,
                "prior_points": round(sum(score_map[dk_id] for dk_id in roster), 3),
            }
        )
    selection_record = {
        "assignments_sha256": assignment_sha,
        "reserved_entries": list(assignments),
        "lineups": lineups,
        "selected_prior_points_by_dk_id": score_map,
        "selection": {
            "differentiation": {"max_person_overlap": 6},
            "threshold_sensitive": [],
            "score_omissions": ["NO_CALIBRATED_CEILING"],
        },
        "participation_detail": {
            "unavailable_people": [],
            "operator_excluded_people": [],
        },
    }
    selection_path = tmp_path / "run" / "selection_report.json"
    selection_sha = write_run_record(selection_path, selection_record)
    unselected = next(player for player in slate.players if player.underlying_id not in people)
    reports = {
        "official_status": {
            "scope": "SUPPLIED_ROWS_ONLY_NOT_FULL_POOL_ACTIVITY_CERTIFICATION",
            "statuses": {unselected.dk_id: "INACTIVE"},
            "source_urls": {unselected.dk_id: "https://example.com/official"},
            "observed_at": {unselected.dk_id: "2026-09-09T20:00:00+00:00"},
        }
    }
    readable = create_readable_review(
        slate=slate,
        template=template,
        salary_path=salary_path,
        entry_path=entry_path,
        assignment_path=assignment_path,
        exported_path=export.output_path,
        artifacts={
            "assignments": str(assignment_path),
            "bulk_entry_csv": export.output_path,
            "selection_report": str(selection_path),
        },
        expected_hashes={
            "salary_csv": sha256_file(salary_path),
            "entry_csv": sha256_file(entry_path),
            "assignments": assignment_sha,
            "bulk_entry_csv": export.output_sha256,
            "selection_report": selection_sha,
        },
        reports=reports,
        truth_values=TRUTHS,
        blockers=("OFFICIAL_STATUS_PARTIAL: one inactive finding is not full-pool activity proof",),
        next_action="Refresh the complete official status evidence.",
        output_dir=tmp_path / "output" / "review",
        package_root=tmp_path,
    )
    selected_qbs = [
        slot
        for entry in readable.data["entries"]
        for slot in entry["slots"]
        if slot["name"] == repeated_name
    ]
    assert {slot["underlying_person_id"] for slot in selected_qbs} == {ne_qb, sea_qb}
    assert len({slot["dk_roster_id"] for slot in selected_qbs}) == 4
    exposures = {
        row["underlying_person_id"]: row for row in readable.data["exposure"]["people"]
    }
    assert exposures[ne_qb]["captain_count"] == 2
    assert exposures[ne_qb]["combined_count"] == 3
    assert exposures[sea_qb]["captain_count"] == 1
    assert exposures[sea_qb]["combined_count"] == 3
    assert any(
        observation["category"] == "official_status"
        and "INACTIVE" in json.dumps(observation)
        for observation in readable.data["evidence_observations"]
    )
    html_text = Path(readable.html_path).read_text(encoding="utf-8")
    assert repeated_name not in html_text
    assert "José O&#x27;Brien &lt;script&gt;" in html_text


def test_semantic_and_exact_byte_mutations_fail_display_reconciliation(tmp_path: Path) -> None:
    slate, template, salary_path, entry_path, outcome, _readable, blockers = _create(tmp_path)

    def attempt(output_name: str, *, artifacts=None, expected_hashes=None):
        return create_readable_review(
            slate=slate,
            template=template,
            salary_path=salary_path,
            entry_path=entry_path,
            assignment_path=outcome.artifacts["assignments"],
            exported_path=outcome.artifacts["bulk_entry_csv"],
            artifacts=artifacts or outcome.artifacts,
            expected_hashes=expected_hashes or outcome.hashes,
            reports=outcome.reports,
            truth_values=TRUTHS,
            blockers=blockers,
            next_action="Rerun.",
            output_dir=tmp_path / output_name,
            package_root=tmp_path,
        )

    exact_mutations = (
        (salary_path, b"\n", "SALARY_CSV_SHA256_MISMATCH"),
        (entry_path, b"\n", "ENTRY_CSV_SHA256_MISMATCH"),
        (Path(outcome.artifacts["assignments"]), b"tampered\n", "ASSIGNMENTS_SHA256_MISMATCH"),
        (Path(outcome.artifacts["portfolio_policy_source"]), b"\n", "SOURCE_POLICY_SHA256_MISMATCH"),
        (Path(outcome.artifacts["portfolio_policy_normalized"]), b"\n", "NORMALIZED_POLICY_SHA256_MISMATCH"),
        (Path(outcome.artifacts["bulk_entry_csv"]), b"\n", "BULK_ENTRY_CSV_SHA256_MISMATCH"),
    )
    for index, (path, suffix, code) in enumerate(exact_mutations):
        original = path.read_bytes()
        path.write_bytes(original + suffix)
        with pytest.raises(ReadableReviewError, match=code):
            attempt(f"mutation-{index}")
        path.write_bytes(original)

    with pytest.raises(ReadableReviewError, match="ARTIFACT_MISSING"):
        attempt(
            "missing-artifact",
            artifacts={
                **outcome.artifacts,
                "declared_but_missing": str(tmp_path / "missing.json"),
            },
        )

    selection_path = Path(outcome.artifacts["selection_report"])
    selection_original = selection_path.read_bytes()
    selection_record = json.loads(selection_original)
    selection_record["lineups"][0]["salary"] += 1
    selection_path.write_text(json.dumps(selection_record), encoding="utf-8")
    semantic_hashes = {**outcome.hashes, "selection_report": sha256_file(selection_path)}
    with pytest.raises(ReadableReviewError, match="SELECTION_SALARY_MISMATCH"):
        attempt("selection-salary", expected_hashes=semantic_hashes)
    selection_path.write_bytes(selection_original)

    entry_original = entry_path.read_bytes()
    entry_path.write_bytes(entry_original.replace(b"Test Showdown", b"Other Showdown"))
    semantic_hashes = {**outcome.hashes, "entry_csv": sha256_file(entry_path)}
    with pytest.raises(ReadableReviewError, match="ENTRY_REPARSE_MISMATCH"):
        attempt("entry-semantic", expected_hashes=semantic_hashes)
    entry_path.write_bytes(entry_original)

    normalized_path = Path(outcome.artifacts["portfolio_policy_normalized"])
    normalized_original = normalized_path.read_bytes()
    normalized_record = json.loads(normalized_original)
    with Path(outcome.artifacts["assignments"]).open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        assignment_rows = list(csv.reader(handle))
    assignment_map = {row[0]: tuple(row[1:]) for row in assignment_rows[1:]}
    by_dk_id = {player.dk_id: player for player in slate.players}
    selected_person = by_dk_id[
        assignment_map[template.authorizations[0].entry_id][0]
    ].underlying_id
    selected_limit = next(
        row for row in normalized_record["effective"]["people"]
        if row["underlying_id"] == selected_person
    )
    selected_limit["combined_max_entries"] = 0
    selected_limit["captain_max_entries"] = 0
    normalized_path.write_text(json.dumps(normalized_record), encoding="utf-8")
    semantic_hashes = {
        **outcome.hashes,
        "portfolio_policy_normalized": sha256_file(normalized_path),
    }
    with pytest.raises(ReadableReviewError, match="COMBINED_EXPOSURE_EXCEEDED"):
        attempt("policy-semantic", expected_hashes=semantic_hashes)
    normalized_path.write_bytes(normalized_original)

    audit_path = Path(outcome.artifacts["portfolio_policy_audit"])
    audit_original = audit_path.read_bytes()
    audit_record = json.loads(audit_original)
    audit_record["captain_counts"] = {}
    audit_path.write_text(json.dumps(audit_record), encoding="utf-8")
    semantic_hashes = {**outcome.hashes, "portfolio_policy_audit": sha256_file(audit_path)}
    with pytest.raises(ReadableReviewError, match="AUDIT_CAPTAIN_EXPOSURE_MISMATCH"):
        attempt("audit-semantic", expected_hashes=semantic_hashes)
    audit_path.write_bytes(audit_original)

    audit_record = json.loads(audit_original)
    audit_record["hashes"]["assignment_artifact_sha256"] = "0" * 64
    audit_path.write_text(json.dumps(audit_record), encoding="utf-8")
    semantic_hashes = {**outcome.hashes, "portfolio_policy_audit": sha256_file(audit_path)}
    with pytest.raises(ReadableReviewError, match="AUDIT_HASH_MISMATCH"):
        attempt("audit-hash", expected_hashes=semantic_hashes)
    audit_path.write_bytes(audit_original)

    html_path = Path(_readable.html_path)
    html_path.write_bytes(html_path.read_bytes() + b"tampered")
    assert verify_readable_review_artifacts(
        json_path=_readable.json_path,
        json_sha256=_readable.json_sha256,
        html_path=html_path,
        html_sha256=_readable.html_sha256,
    )[0].startswith("READABLE_REVIEW_HTML_SHA256_MISMATCH")


def test_five_entry_fixture_and_copied_layout_replay_are_deterministic(tmp_path: Path) -> None:
    first_root = tmp_path / "first" / "package"
    second_root = tmp_path / "second" / "package"
    first_root.mkdir(parents=True)
    slate, template, salary_path, entry_path, outcome, first, blockers = _create(
        first_root, entry_count=5
    )
    shutil.copytree(first_root, second_root)
    copied_artifacts = {
        name: str(second_root / Path(path).resolve().relative_to(first_root.resolve()))
        for name, path in outcome.artifacts.items()
    }
    second = create_readable_review(
        slate=parse_salaries(second_root / salary_path.relative_to(first_root)),
        template=parse_entries(second_root / entry_path.relative_to(first_root)),
        salary_path=second_root / salary_path.relative_to(first_root),
        entry_path=second_root / entry_path.relative_to(first_root),
        assignment_path=copied_artifacts["assignments"],
        exported_path=copied_artifacts["bulk_entry_csv"],
        artifacts=copied_artifacts,
        expected_hashes=outcome.hashes,
        reports=outcome.reports,
        truth_values=TRUTHS,
        blockers=blockers,
        next_action="Refresh exact official activity evidence and rerun.",
        output_dir=second_root / "output" / "review",
        package_root=second_root,
    )
    first_data = json.loads(Path(first.json_path).read_text(encoding="utf-8"))
    assert len(first_data["entries"]) == 5
    assert len(first_data["exposure"]["pairwise_overlap"]) == 10
    assert Path(first.json_path).read_bytes() == Path(second.json_path).read_bytes()
    assert Path(first.html_path).read_bytes() == Path(second.html_path).read_bytes()
    create_cowork_status_workbook(
        output_path=first_root / "output" / "NFL_DFS_Cowork_Review_five-entry.xlsx",
        run_values={"RUN_LABEL": "five-entry-render-fixture"},
        blockers=blockers,
        report_path=first_root / "output" / "cowork_run.json",
        truth_values=TRUTHS,
        readable_review=first.data,
    )


def test_reordered_flex_duplicate_is_caught_by_canonical_identity(tmp_path: Path) -> None:
    slate, template, salary_path, entry_path, outcome, _readable, blockers = _create(tmp_path)
    with Path(outcome.artifacts["assignments"]).open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        assignment_rows = list(csv.reader(handle))
    assignments = {row[0]: tuple(row[1:]) for row in assignment_rows[1:]}
    first_entry, second_entry = (entry.entry_id for entry in template.authorizations)
    first_roster = assignments[first_entry]
    assignments[second_entry] = (first_roster[0], *reversed(first_roster[1:]))
    assignment_path = Path(outcome.artifacts["assignments"])
    assignment_sha = write_assignments_csv(
        assignment_path,
        assignments,
        entry_order=(first_entry, second_entry),
    )
    export = export_review_entries(
        slate=slate,
        template=template,
        assignments=assignments,
        output_path=tmp_path / "duplicate" / "DK_REVIEW_ENTRY_reordered.csv",
    )
    assert export.file_valid
    selection_path = Path(outcome.artifacts["selection_report"])
    selection = json.loads(selection_path.read_bytes())
    selection["assignments_sha256"] = assignment_sha
    second_validation = validate_lineup(slate, assignments[second_entry])
    assert second_validation.lineup is not None
    selection["lineups"][1]["roster"] = list(assignments[second_entry])
    selection["lineups"][1]["salary"] = second_validation.lineup.salary
    selection["lineups"][1]["canonical_key"] = second_validation.lineup.canonical_key
    selection["lineups"][1]["prior_points"] = selection["lineups"][0]["prior_points"]
    selection_sha = write_run_record(selection_path, selection)
    expected_hashes = {
        **outcome.hashes,
        "assignments": assignment_sha,
        "bulk_entry_csv": export.output_sha256,
        "selection_report": selection_sha,
    }
    artifacts = {
        **outcome.artifacts,
        "bulk_entry_csv": export.output_path,
    }
    with pytest.raises(ReadableReviewError, match="CANONICAL_DUPLICATE"):
        create_readable_review(
            slate=slate,
            template=template,
            salary_path=salary_path,
            entry_path=entry_path,
            assignment_path=assignment_path,
            exported_path=export.output_path,
            artifacts=artifacts,
            expected_hashes=expected_hashes,
            reports=outcome.reports,
            truth_values=TRUTHS,
            blockers=blockers,
            next_action="Rerun.",
            output_dir=tmp_path / "duplicate",
            package_root=tmp_path,
        )
