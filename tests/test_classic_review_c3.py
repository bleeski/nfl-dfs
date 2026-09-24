"""C3 golden, mutation, rendering-safety, replay, and scale acceptance."""

from __future__ import annotations

import copy
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nfl_dfs.classic_portfolio_policy import (
    classic_portfolio_policy_template,
    validate_classic_portfolio_policy_bytes,
    write_normalized_classic_portfolio_policy,
)
from nfl_dfs.classic_review import (
    ClassicReviewError,
    _audit_template_bytes,
    create_classic_review_package,
)
from nfl_dfs.classic_scale_acceptance import run_classic_c3_scale_acceptance
from nfl_dfs.dk import parse_entries, parse_entry_bytes, parse_salaries
from nfl_dfs.hashing import sha256_file
from nfl_dfs.prior_review import run_prior_review
from nfl_dfs.projection import build_projection_package
from nfl_dfs.readable_review import _render_html
from nfl_dfs.workbook import create_cowork_status_workbook

from .test_classic_prior_review import AS_OF, _fixture


_C3_OUTPUT_KEYS = {
    "classic_export_audit",
    "bulk_entry_csv",
    "readable_review_json",
    "readable_review_html",
    "review_workbook",
}


@pytest.fixture(scope="module")
def c3_seed(tmp_path_factory: pytest.TempPathFactory):
    root = tmp_path_factory.mktemp("c3seed")
    salary, entry, package, role, status, _inactive = _fixture(
        root / "fixture", entries=3
    )
    slate = parse_salaries(salary)
    template = parse_entries(entry)
    policy_document = classic_portfolio_policy_template(
        slate,
        tuple(item.entry_id for item in template.authorizations),
        entry_sha256=template.raw_hash,
    )
    policy_path = root / "classic_policy.json"
    policy_path.write_text(json.dumps(policy_document, indent=2), encoding="utf-8")
    validated = validate_classic_portfolio_policy_bytes(
        policy_path.read_bytes(),
        slate=slate,
        entry_ids=tuple(item.entry_id for item in template.authorizations),
        entry_sha256=template.raw_hash,
    )
    assert validated.valid and validated.policy is not None
    normalized = write_normalized_classic_portfolio_policy(
        root / "classic_policy.normalized.json", validated.policy
    )
    outcome = run_prior_review(
        salary_csv=salary,
        entry_csv=entry,
        label="c3-seed",
        as_of=AS_OF,
        run_root=root / "run",
        output_root=root / "out",
        prior_package_dir=package,
        build_priors=True,
        official_status_csv=status,
        offensive_role_evidence_json=role,
        portfolio_policy=validated.policy,
        portfolio_policy_source_path=policy_path,
        portfolio_policy_source_sha256=sha256_file(policy_path),
        portfolio_policy_normalized_path=normalized,
        portfolio_policy_normalized_sha256=sha256_file(normalized),
        project=build_projection_package,
    )
    assert not outcome.blocked, outcome.blockers
    return root, salary, entry, outcome


def _direct_args(c3_seed, target: Path):
    source_root, source_salary, source_entry, outcome = c3_seed
    shutil.copytree(source_root, target)
    artifacts = {
        name: str(target / Path(path).resolve().relative_to(source_root))
        for name, path in outcome.artifacts.items()
        if name not in _C3_OUTPUT_KEYS
    }
    return {
        "salary_path": target / source_salary.resolve().relative_to(source_root),
        "entry_path": target / source_entry.resolve().relative_to(source_root),
        "artifacts": artifacts,
        "expected_hashes": outcome.hashes,
        "audit_at": AS_OF,
        "output_path": target / "fresh" / "review" / "DK_REVIEW_ENTRY_c3.csv",
        "output_dir": target / "fresh" / "review",
        "package_root": target,
    }


def _assert_no_new_output(args) -> None:
    assert not Path(args["output_path"]).exists()
    assert not (Path(args["output_dir"]) / "classic_review_export_audit.json").exists()
    assert not (Path(args["output_dir"]) / "prior_only_readable_review.json").exists()
    assert not (Path(args["output_dir"]) / "prior_only_readable_review.html").exists()


def test_golden_exact_template_reconciliation_and_readable_display(c3_seed, tmp_path: Path) -> None:
    args = _direct_args(c3_seed, tmp_path / "golden")
    result = create_classic_review_package(**args)
    source = Path(args["entry_path"]).read_bytes()
    output = Path(result.export_path).read_bytes()
    assignment = json.loads(Path(args["artifacts"]["classic_assignment"]).read_text())
    assignments = {
        item["entry_id"]: tuple(item["roster"])
        for item in assignment["entry_assignments"]
    }
    template = parse_entries(args["entry_path"])
    assert _audit_template_bytes(
        source_bytes=source,
        output_bytes=output,
        encoding=template.encoding,
        roster_start=template.roster_start_index,
        roster_width=9,
        assignments=assignments,
    ) == ()
    reparsed = parse_entry_bytes(output, source_name=result.export_path)
    assert [item.entry_id for item in reparsed.authorizations] == list(assignments)
    assert [entry["entry_id"] for entry in result.data["entries"]] == list(assignments)
    assert all(len(entry["slots"]) == 9 for entry in result.data["entries"])
    assert result.data["truths"] == {
        "FILE_VALID": True,
        "EVIDENCE_STATE": "PASS",
        "MODEL_STATUS": "PRIOR_ONLY",
        "RELEASE_DECISION": "DO_NOT_UPLOAD",
    }
    assert result.audit["recomputed"]["pairwise_person_overlap"] == result.data["exposure"]["pairwise_overlap"]
    assert not list(tmp_path.rglob("DK_UPLOAD_*.csv"))


def test_complete_package_copy_replay_is_byte_identical(c3_seed, tmp_path: Path) -> None:
    first_args = _direct_args(c3_seed, tmp_path / "source")
    second_args = _direct_args(c3_seed, tmp_path / "copy")
    first = create_classic_review_package(**first_args)
    second = create_classic_review_package(**second_args)
    for left, right in (
        (first.audit_path, second.audit_path),
        (first.export_path, second.export_path),
        (first.json_path, second.json_path),
        (first.html_path, second.html_path),
    ):
        assert Path(left).read_bytes() == Path(right).read_bytes()


@pytest.mark.parametrize(
    "artifact_name",
    (
        "portfolio_policy_source",
        "portfolio_policy_normalized",
        "classic_candidate_bank",
        "classic_assignment",
        "classic_portfolio_audit",
        "selection_report",
        "complete_slate_coverage",
        "classic_selected_scores",
        "team_source",
        "player_source",
        "identity_map",
        "team_projections",
        "player_opportunities",
        "source_ledger",
        "team_splits",
        "official_status_csv",
        "offensive_role_evidence_json",
    ),
)
def test_every_bound_artifact_mutation_is_withheld(
    c3_seed, tmp_path: Path, artifact_name: str
) -> None:
    args = _direct_args(c3_seed, tmp_path / artifact_name.replace(":", "-"))
    path = Path(args["artifacts"][artifact_name])
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(ClassicReviewError, match="INTAKE_HASH_MISMATCH"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


@pytest.mark.parametrize("input_name", ("salary_path", "entry_path"))
def test_salary_or_template_mutation_is_withheld(c3_seed, tmp_path: Path, input_name: str) -> None:
    args = _direct_args(c3_seed, tmp_path / input_name)
    path = Path(args[input_name])
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(ClassicReviewError, match="INTAKE_HASH_MISMATCH"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


@pytest.mark.parametrize(
    ("after_checkpoint", "next_checkpoint"),
    (("INTAKE", "PRE_EXPORT"), ("PRE_EXPORT", "POST_WRITE"), ("POST_WRITE", "PRE_RENDER")),
)
def test_mutation_at_every_hash_boundary_is_withheld(
    c3_seed,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    after_checkpoint: str,
    next_checkpoint: str,
) -> None:
    from nfl_dfs import classic_review

    args = _direct_args(c3_seed, tmp_path / after_checkpoint.lower())
    hash_only_source = Path(args["artifacts"]["team_source"])
    original = classic_review._hash_checkpoint
    mutated = False

    def hook(name, paths, expected):
        nonlocal mutated
        result = original(name, paths, expected)
        if name == after_checkpoint and not mutated:
            hash_only_source.write_bytes(hash_only_source.read_bytes() + b" ")
            mutated = True
        return result

    monkeypatch.setattr(classic_review, "_hash_checkpoint", hook)
    with pytest.raises(ClassicReviewError, match=next_checkpoint + "_HASH_MISMATCH"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


@pytest.mark.parametrize(
    ("label", "field", "value", "expected"),
    (
        ("CANDIDATE_BANK", "status", "CANDIDATE_BANK_TIMEOUT", "CANDIDATE_BANK_NOT_ACCEPTED"),
        ("CANDIDATE_BANK", "status", "CANDIDATE_BANK_INCOMPLETE", "CANDIDATE_BANK_NOT_ACCEPTED"),
        ("CANDIDATE_BANK", "status", "CANDIDATE_BANK_SOLVER_ERROR", "CANDIDATE_BANK_NOT_ACCEPTED"),
        ("C2_AUDIT", "status", "FAIL", "C2_AUDIT_NOT_PASS"),
    ),
)
def test_blocking_bank_solver_and_audit_states_are_withheld(
    c3_seed,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    label: str,
    field: str,
    value: str,
    expected: str,
) -> None:
    from nfl_dfs import classic_review

    args = _direct_args(c3_seed, tmp_path / (label + value).lower())
    original = classic_review._strict_json

    def hook(raw, *, label: str, canonical: bool = False):
        result = original(raw, label=label, canonical=canonical)
        if label == locals_label:
            changed = copy.deepcopy(dict(result))
            changed[field] = value
            return changed
        return result

    locals_label = label
    monkeypatch.setattr(classic_review, "_strict_json", hook)
    with pytest.raises(ClassicReviewError, match=expected):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


def test_nonoptimal_joint_selection_state_is_withheld(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import classic_review

    args = _direct_args(c3_seed, tmp_path / "nonoptimal")
    original = classic_review._strict_json

    def hook(raw, *, label: str, canonical: bool = False):
        result = original(raw, label=label, canonical=canonical)
        if label == "SELECTION":
            changed = copy.deepcopy(dict(result))
            changed["portfolio_policy"]["enforcement"]["joint_selection_status"] = "PORTFOLIO_SELECTION_TIMEOUT"
            return changed
        return result

    monkeypatch.setattr(classic_review, "_strict_json", hook)
    with pytest.raises(ClassicReviewError, match="JOINT_SELECTION_NOT_OPTIMAL"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


def _limit_stopped(c3_seed, tmp_path, monkeypatch, *, joint=None, scope=None, bank_status=None):
    """C3 on a seed whose records say a limit stopped the joint solve or the bank."""

    from nfl_dfs import classic_review

    args = _direct_args(c3_seed, tmp_path)
    original = classic_review._strict_json

    def hook(raw, *, label: str, canonical: bool = False):
        result = original(raw, label=label, canonical=canonical)
        if label == "SELECTION":
            changed = copy.deepcopy(dict(result))
            enforcement = changed["portfolio_policy"]["enforcement"]
            if joint is not None:
                enforcement["joint_selection_status"] = joint
                enforcement["optimality_scope"] = scope
            if bank_status is not None:
                enforcement["candidate_bank_status"] = bank_status
            return changed
        if label == "CANDIDATE_BANK" and bank_status is not None:
            changed = copy.deepcopy(dict(result))
            changed["status"] = bank_status
            return changed
        return result

    monkeypatch.setattr(classic_review, "_strict_json", hook)
    return args


def test_a_limit_incumbent_joint_selection_is_accepted_and_named(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Session 08: labelled, with no optimality scope, and named as a limitation."""

    args = _limit_stopped(c3_seed, tmp_path / "incumbent", monkeypatch,
                          joint="FEASIBLE_LIMIT_ACTUAL_CANDIDATE_BANK", scope=None)
    result = create_classic_review_package(**args)
    assert result.audit["joint_selection"]["status"] == "FEASIBLE_LIMIT_ACTUAL_CANDIDATE_BANK"
    assert result.audit["joint_selection"]["optimality_scope"] is None
    assert any(item.startswith("PORTFOLIO_SELECTION_LIMIT_INCUMBENT:") for item in result.audit["limitations"])
    assert result.data["limitations"] == result.audit["limitations"]
    html = Path(result.html_path).read_text(encoding="utf-8")
    assert "feasible, not proven optimal" in html
    assert Path(result.export_path).is_file()


def test_a_limit_incumbent_claiming_the_optimum_scope_is_withheld(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    args = _limit_stopped(c3_seed, tmp_path / "overclaim", monkeypatch,
                          joint="FEASIBLE_LIMIT_ACTUAL_CANDIDATE_BANK", scope="ACTUAL_CANDIDATE_BANK")
    with pytest.raises(ClassicReviewError, match="CLASSIC_C3_C2_OPTIMALITY_SCOPE_MISMATCH"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


@pytest.mark.parametrize("bank_status", ("BOUNDED_TIME_LIMIT_STOP", "BOUNDED_SEARCH_LIMIT_STOP"))
def test_a_limit_stopped_bank_is_accepted_and_named(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, bank_status: str
) -> None:
    args = _limit_stopped(c3_seed, tmp_path / bank_status.lower(), monkeypatch, bank_status=bank_status)
    result = create_classic_review_package(**args)
    assert result.audit["candidate_bank"]["status"] == bank_status
    assert any(item.startswith(f"CANDIDATE_BANK_STOPPED_AT_LIMIT:{bank_status}:")
               for item in result.audit["limitations"])
    # A proven optimum over the bank stays scoped to it.
    assert result.audit["joint_selection"]["optimality_scope"] == "ACTUAL_CANDIDATE_BANK"


@pytest.mark.parametrize(
    ("limit", "model_status"), (("nodes", "kSolutionLimit"), ("time", "kTimeLimit"))
)
def test_run_slate_delivers_a_limited_incumbent_from_a_time_stopped_bank(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, limit: str, model_status: str
) -> None:
    """The card's acceptance, end to end through run-slate and the C3 export.

    Real HiGHS stops the joint selection holding the bank's witness start:
    at `mip_max_nodes=0` (kSolutionLimit), and at a 1e-9 s time limit
    (kTimeLimit), which a clock can only reach sooner. The bank's total budget
    runs out two solves into its top-k fill; its other limits are ones a small
    bank cannot reach. The export is delivered, both limits are named `S`
    limitations, and this host's candidate rate is still read from the bank's
    own report.
    """

    import highspy

    from nfl_dfs import classic_portfolio, selection
    from nfl_dfs.classic_portfolio_policy import classic_portfolio_policy_template
    from nfl_dfs.deadline import read_candidate_rate

    from .test_classic_portfolio_c2 import UNREACHABLE_LIMITS, _stop_in_the_fill
    from .test_deadline_controller import _classic, _truth_codes

    def policy_file(attachments: Path) -> Path:
        slate = parse_salaries(attachments / "salary.csv")
        entries = parse_entries(attachments / "entries.csv")
        path = attachments.parent / "classic_policy.json"
        path.write_text(json.dumps(classic_portfolio_policy_template(
            slate, tuple(item.entry_id for item in entries.authorizations),
            entry_sha256=entries.raw_hash, limits=UNREACHABLE_LIMITS,
        )), encoding="utf-8")
        return path

    def limited():
        model = highspy.Highs()
        model.setOptionValue("presolve", "off")  # presolve alone picks one lineup of a small bank
        if limit == "nodes":
            model.setOptionValue("mip_max_nodes", 0)
        return model

    real_solve = selection.solve_classic_portfolio

    def solve(policy, bank, **kwargs):
        if limit == "time":
            kwargs["time_limit_seconds"] = 1e-9
        return real_solve(policy, bank, solver_factory=limited, **kwargs)

    monkeypatch.setattr(selection, "solve_classic_portfolio", solve)
    _stop_in_the_fill(monkeypatch, classic_portfolio, after=2)
    code, report, root = _classic(tmp_path, monkeypatch, run_id="limited", policy=policy_file,
                                  deadline="2099-01-01T00:00:00+00:00")
    assert code == 0 and report["stage"] == "PRIOR_ONLY_CLASSIC_C3_REVIEW_EXPORT", report["blockers"][:3]
    assert report["improvement"]["status"] == "DELIVERED"
    assert report["latest_deliverable"]["producer"] != "run-slate:baseline"
    assert Path(report["bulk_entry_csv"]).is_file()
    assert report["RELEASE_DECISION"] == "DO_NOT_UPLOAD"
    codes = _truth_codes(report)
    assert codes["PORTFOLIO_SELECTION_LIMIT_INCUMBENT"] == "S"
    assert codes["CANDIDATE_BANK_STOPPED_AT_LIMIT"] == "S"
    named = next(item for item in report["release_truths"]["delivery_limitations"]
                 if item["code"] == "PORTFOLIO_SELECTION_LIMIT_INCUMBENT")
    assert f"stopped at {model_status} with a validated incumbent" in named["detail"]
    audit = json.loads(Path(report["prior_review_artifacts"]["classic_export_audit"]).read_text(encoding="utf-8"))
    assert audit["joint_selection"] == {
        "status": "FEASIBLE_LIMIT_ACTUAL_CANDIDATE_BANK", "optimality_scope": None, "c2_audit_status": "PASS",
    }
    assert audit["candidate_bank"]["status"] == "BOUNDED_TIME_LIMIT_STOP"
    rate = report["deadline"]["candidate_rate"]
    assert rate["basis"] == "BANK_REPORT" and rate["seconds_per_candidate"] > 0
    assert read_candidate_rate(tmp_path / "runs" / "host_candidate_rates.json", mode="CLASSIC") is not None


def test_prefilled_stale_partial_and_unauthorized_exports_are_withheld(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import classic_review

    stale = _direct_args(c3_seed, tmp_path / "stale")
    Path(stale["output_path"]).parent.mkdir(parents=True)
    Path(stale["output_path"]).write_bytes(b"stale")
    with pytest.raises(ClassicReviewError, match="OUTPUT_EXISTS"):
        create_classic_review_package(**stale)

    prefilled = _direct_args(c3_seed, tmp_path / "prefilled")
    source = Path(prefilled["entry_path"])
    source.write_bytes(source.read_bytes().replace(b",,,,,,,,,,,", b",43700001,,,,,,,,,", 1))
    changed_hashes = dict(prefilled["expected_hashes"])
    changed_hashes["entry_csv"] = sha256_file(source)
    prefilled["expected_hashes"] = changed_hashes
    with pytest.raises(ClassicReviewError, match="BLANK_CELL_AUTHORITY_REQUIRED"):
        create_classic_review_package(**prefilled)
    _assert_no_new_output(prefilled)

    unauthorized = _direct_args(c3_seed, tmp_path / "unauthorized")
    original_upload = classic_review.write_upload_bytes
    monkeypatch.setattr(
        classic_review,
        "write_upload_bytes",
        lambda template, assignments: original_upload(template, assignments) + b"UNAUTHORIZED\r\n",
    )
    with pytest.raises(ClassicReviewError, match="PHYSICAL_LINE_COUNT_CHANGED"):
        create_classic_review_package(**unauthorized)
    _assert_no_new_output(unauthorized)
    monkeypatch.undo()

    partial = _direct_args(c3_seed, tmp_path / "partial")
    original_atomic = classic_review._atomic_write

    def partial_write(path, raw):
        if Path(path) == Path(partial["output_path"]):
            Path(path).write_bytes(raw[: max(1, len(raw) // 2)])
            return "not-the-proposed-hash"
        return original_atomic(path, raw)

    monkeypatch.setattr(classic_review, "_atomic_write", partial_write)
    with pytest.raises(ClassicReviewError, match="PARTIAL_OR_WRITE_HASH_MISMATCH"):
        create_classic_review_package(**partial)
    _assert_no_new_output(partial)


def test_selected_unavailable_and_missing_current_role_stop(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import classic_review

    unavailable = _direct_args(c3_seed, tmp_path / "unavailable")
    assignment = json.loads(Path(unavailable["artifacts"]["classic_assignment"]).read_text())
    selected_id = assignment["entry_assignments"][0]["roster"][0]
    status = Path(unavailable["artifacts"]["official_status_csv"])
    status.write_text(
        status.read_text(encoding="utf-8").replace(f",{selected_id},ACTIVE,", f",{selected_id},INACTIVE,"),
        encoding="utf-8",
        newline="",
    )
    hashes = dict(unavailable["expected_hashes"])
    hashes["official_status_csv"] = sha256_file(status)
    unavailable["expected_hashes"] = hashes
    with pytest.raises(ClassicReviewError, match="SELECTED_ACTIVITY_NOT_ACTIVE"):
        create_classic_review_package(**unavailable)
    _assert_no_new_output(unavailable)

    missing_role = _direct_args(c3_seed, tmp_path / "missing-role")
    original = classic_review._strict_json

    def hook(raw, *, label: str, canonical: bool = False):
        result = original(raw, label=label, canonical=canonical)
        if label == "OFFENSIVE_ROLE_EVIDENCE":
            changed = copy.deepcopy(dict(result))
            changed["declarations"] = []
            return changed
        return result

    monkeypatch.setattr(classic_review, "_strict_json", hook)
    with pytest.raises(ClassicReviewError, match="SELECTED_ROLE_EVIDENCE_MISSING"):
        create_classic_review_package(**missing_role)
    _assert_no_new_output(missing_role)


@pytest.mark.parametrize("where", ["coverage", "gate"])
def test_artifacts_that_disagree_about_missing_activity_still_refuse(
    c3_seed, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, where: str
) -> None:
    """R28 moved only the missing row (Session 09). A bound artifact that claims
    a selected person lacks a row when C3's own re-read finds one, or a gate
    that hides a gap, is conflicted evidence and still refuses the export."""

    from nfl_dfs import classic_review

    args = _direct_args(c3_seed, tmp_path / where)
    assignment = json.loads(Path(args["artifacts"]["classic_assignment"]).read_text())
    selected_id = assignment["entry_assignments"][0]["roster"][0]
    person = next(
        row.underlying_id for row in parse_salaries(args["salary_path"]).players if row.dk_id == selected_id
    )
    original = classic_review._strict_json

    def hook(raw, *, label: str, canonical: bool = False):
        result = original(raw, label=label, canonical=canonical)
        if label != "COVERAGE":
            return result
        changed = copy.deepcopy(dict(result))
        if where == "coverage":
            changed["official_status_coverage"]["selected_without_row"] = [person]
        else:
            changed["selected_evidence_gate"]["activity_gaps"] = [{"person": person}]
            changed["selected_evidence_gate"]["status"] = "PASS_WITH_NAMED_LIMITATIONS"
        return changed

    monkeypatch.setattr(classic_review, "_strict_json", hook)
    with pytest.raises(ClassicReviewError, match="SELECTED_ACTIVITY_COVERAGE_MISMATCH"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


def test_a_coverage_that_names_a_file_the_review_does_not_track_refuses(c3_seed, tmp_path: Path) -> None:
    args = _direct_args(c3_seed, tmp_path / "untracked-status")
    args["artifacts"] = {k: v for k, v in args["artifacts"].items() if k != "official_status_csv"}
    args["expected_hashes"] = {k: v for k, v in args["expected_hashes"].items() if k != "official_status_csv"}
    with pytest.raises(ClassicReviewError, match="CLASSIC_C3_OFFICIAL_STATUS_ARTIFACT_REQUIRED"):
        create_classic_review_package(**args)
    _assert_no_new_output(args)


def test_classic_html_and_workbook_injection_defenses(c3_seed, tmp_path: Path) -> None:
    args = _direct_args(c3_seed, tmp_path / "render-safety")
    result = create_classic_review_package(**args)
    review = copy.deepcopy(result.data)
    review["entries"][0]["contest_name"] = "<script>=unsafe</script>"
    review["entries"][0]["slots"][0]["name"] = "=1+1"
    review["artifacts"][0]["path"] = "+unsafe"
    html = _render_html(review, data_sha256="0" * 64).decode("utf-8")
    assert "<script>=unsafe</script>" not in html
    assert "&lt;script&gt;=unsafe&lt;/script&gt;" in html
    workbook_path = tmp_path / "classic-review.xlsx"
    create_cowork_status_workbook(
        output_path=workbook_path,
        run_values={"RUN_LABEL": "=unsafe"},
        blockers=("@unsafe",),
        report_path=tmp_path / "cowork_run.json",
        truth_values=result.data["truths"],
        readable_review=review,
        review_csv_path=result.export_path,
    )
    workbook = load_workbook(workbook_path, data_only=False)
    assert len(workbook.sheetnames) == 8
    hostile_cells = [
        cell
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "unsafe" in cell.value
    ]
    assert hostile_cells
    assert all(cell.data_type != "f" for cell in hostile_cells)


@pytest.mark.parametrize("entry_count", (1, 3, 20, 150))
def test_registered_full_fixture_scale_and_exact_entry_order(
    tmp_path: Path, entry_count: int
) -> None:
    repo = Path(__file__).resolve().parents[1]
    report = run_classic_c3_scale_acceptance(
        repo_root=repo,
        work_root=tmp_path / f"scale-{entry_count}",
        entry_count=entry_count,
    )
    assert report["status"] == "PASS"
    assert report["fixture"] == {
        "salary_sha256": "d8ce5c1dd643ec419a0c4bc99b05a5f619f99dd23bc8369f70b795d8ac6132fb",
        "people": 719,
        "teams": 24,
        "games": 12,
    }
    assert report["candidate_bank"]["produced"] == report["candidate_bank"]["requested"]
    assert report["candidate_bank"]["completeness"] in {
        "EXHAUSTIVE",
        "CANDIDATE_BANK_INCOMPLETE",
    }
    assert report["joint_selection"]["selected_entries"] == entry_count
    assert report["c3"]["exact_entry_order"] == [
        str(7000000001 + index) for index in range(entry_count)
    ]
    assert report["copied_package_replay"]["status"] == "PASS"
    assert report["resources"]["within_registered_limits"] is True
    assert report["resources"]["process_peak_rss_bytes"] > 0
    assert report["resources"]["traced_python_peak_bytes"] > 0
    assert report["c3"]["no_dk_upload"] is True
