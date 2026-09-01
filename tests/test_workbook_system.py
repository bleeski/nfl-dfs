from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from nfl_dfs.system import doctor
from nfl_dfs.workbook import create_operator_input_workbook


def test_operator_workbook_has_exact_five_sheet_surface(tmp_path: Path) -> None:
    path = create_operator_input_workbook(tmp_path / "operator.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Run Control", "Evidence Paste", "Portfolio", "QA", "Upload"]
    assert workbook["Run Control"]["B17"].value == "LARGE_GPP"
    assert workbook["Evidence Paste"]["J108"].value.startswith("=SUMPRODUCT")
    assert workbook["Upload"]["B4"].value == "DO_NOT_UPLOAD"
    assert workbook["Run Control"].freeze_panes == "A4"


def test_doctor_sqlite_and_runtime(tmp_path: Path) -> None:
    report = doctor(tmp_path)
    assert report.sqlite_integrity == "ok"
    assert report.sqlite_journal_mode in {"WAL", "DELETE"}
    assert report.available_memory_bytes > 0
    assert report.processors >= 1
