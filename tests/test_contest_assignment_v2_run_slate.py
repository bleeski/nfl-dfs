"""Session 76: the run path on `within_contest_diversity_v2`, read through the audit, the review and the workbook.

`tests/test_contest_assignment_run_slate.py` and `tests/test_contest_assignment_classic.py` already follow every exit's
assignment into the file, the audit and both reviews, now on v2. This file checks what only v2 adds: the objective the
step reports is the one the delivered bytes bear out, the policy audit names v2, and the review, its HTML and the
workbook read the block by its objective while a v1 block renders as it did.
"""

from __future__ import annotations

from pathlib import Path

from nfl_dfs import contest_assignment as ca
from nfl_dfs import readable_review
from nfl_dfs.dk import parse_entries

from .test_contest_assignment_run_slate import (
    INTERLEAVED_CONTESTS,
    SIX,
    _assert_truths_unchanged,
    _run,
    _selection_rosters,
)
from .test_entry_groups import SD3_CONTROLS, _cells, _readable


def _objective(total):
    return (*(total[name] for name in ca.OBJECTIVE_ORDER[:4]), float(total["spread"]))


def test_the_policy_exit_reports_v2_s_objective_and_the_bytes_bear_it_out(tmp_path, monkeypatch):
    code, report, entries, slate = _run(
        tmp_path, monkeypatch, run_id="cd-v2", contests=INTERLEAVED_CONTESTS, policy_controls=SD3_CONTROLS)
    assert code == 0, report["blockers"]
    _assert_truths_unchanged(report)
    step = report["prior_review_reports"]["contest_assignment"]
    assert step["contest_assignment_version"] == ca.CONTEST_ASSIGNMENT_VERSION_V2
    assert step["score_definition"] == ca.CONTEST_OBJECTIVE_DEFINITION
    assert step["objective_order"] == list(ca.OBJECTIVE_ORDER) and step["single_entry_rows_move"] is True

    people = ca.people_from_slate(slate.players)
    contest_of = {e.entry_id: e.contest_id for e in parse_entries(entries).authorizations}
    delivered = _cells(Path(report["latest_deliverable"]["path"]))
    after = ca.objectives_from_rosters(
        [(e, contest_of[e], tuple(r)) for e, r in delivered.items() if any(r)], mode=ca.MODE_SHOWDOWN, people=people)
    solver = ca.objectives_from_rosters(
        [(e, contest_of[e], r) for e, r in zip(SIX, _selection_rosters(report))], mode=ca.MODE_SHOWDOWN, people=people)
    assert after == step["contest_objective_after"] and ca.objective_total(after) == step["objective_after"]
    assert solver == step["contest_objective_before"]
    assert not ca.objective_improves(_objective(step["objective_before"]), _objective(step["objective_after"]))

    audit = report["prior_review_reports"]["portfolio_policy_audit"]["contest_assignment"]
    assert audit["contest_assignment_version"] == ca.CONTEST_ASSIGNMENT_VERSION_V2
    assert audit["status"] == "PASS" and audit["findings"] == []

    block = _readable(report)["contest_assignment"]
    assert block["reported_statistics_match"] is True
    assert block["objective_after"] == step["objective_after"] and block["objective_order"] == list(ca.OBJECTIVE_ORDER)
    html = Path(report["prior_review_artifacts"]["readable_review_html"]).read_text("utf-8")
    assert "ranks a repeated Captain as the costliest difference" in html
    assert "v1 score reading (before, after)" in html
    assert "contests hold one entry and are left as they were" not in html

    from openpyxl import load_workbook

    sheet = load_workbook(report["review_workbook"])["Exposure"]
    values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert "v1 score reading before, after" in values and "Same-Captain pairs before, after" in values


def _small_block(version):
    rows = [ca.EntryRow("1", "A", ("c1", "s2"), "p"), ca.EntryRow("2", "A", ("c1", "s3"), "p"),
            ca.EntryRow("3", "B", ("c4", "s5"), "p"), ca.EntryRow("4", "B", ("c6", "s7"), "p"),
            ca.EntryRow("5", "S", ("c8", "s9"), "p")]
    people = {dk_id: ca.Person(dk_id[1:]) for row in rows for dk_id in row.roster}
    out = ca.diversify(rows, mode=ca.MODE_SHOWDOWN, people=people, restarts=5, time_limit_seconds=None,
                       version=version)
    block, problems = ca.review_block(
        mode=ca.MODE_SHOWDOWN, people=people, selected_in_solver_order=[r.roster for r in rows], reported=out.report,
        rows=[(r.entry_id, r.contest_id, tuple(out.assignments[r.entry_id]), "all") for r in rows])
    assert problems == []
    return block, "".join(readable_review._contest_assignment_html(block))


def test_a_v1_block_renders_as_it_did():
    _block, html = _small_block(ca.CONTEST_ASSIGNMENT_VERSION_V1)
    assert "A contest scores its worst pair plus its mean pair" in html
    assert "1 contests hold one entry and are left as they were." in html
    assert "costliest difference" not in html and "v1 score reading" not in html


def test_a_v1_block_keeps_its_workbook_rows():
    from openpyxl import Workbook

    from nfl_dfs import workbook

    rows = {}
    for version in (ca.CONTEST_ASSIGNMENT_VERSION_V1, ca.CONTEST_ASSIGNMENT_VERSION_V2):
        block, _html = _small_block(version)
        sheet = Workbook().active
        workbook._add_contest_assignment_rows(sheet, {"contest_assignment": block}, 1)
        rows[version] = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    v1, v2 = rows[ca.CONTEST_ASSIGNMENT_VERSION_V1], rows[ca.CONTEST_ASSIGNMENT_VERSION_V2]
    assert "Total score before, after" in v1 and "Contests holding one entry (left as they were)" in v1
    assert "Score before" in v1 and not any("v1 score reading" in value for value in v1)
    assert "v1 score reading before, after" in v2 and "Contests holding one entry (their rows may move)" in v2


def test_a_v2_block_names_its_objective_and_the_one_entry_contests_it_changed():
    block, html = _small_block(ca.CONTEST_ASSIGNMENT_VERSION_V2)
    assert "ranks a repeated Captain as the costliest difference" in html
    assert "v1 score reading (before, after)" in html
    assert "left as they were" not in html
    changed = ", ".join(block["single_entry_contests_moved"]) or "none"
    assert f"1 contests hold one entry; v2 may move their rows to relieve a repeat elsewhere. Changed: {changed}." in html
