"""Session 60 (R37, P9 part 1): the injury room moves the workload.

Every number here is a test input. The assertions are about where a vacated share
goes (the position room, never another position, never another team), that a team's
pooled share is conserved exactly, that the trigger is the DraftKings status the run
binds plus a supplied official INACTIVE row and nothing else, that the redistributed
model is the one scoring receives at both call sites, and that every move is reported.

The Week 4 frozen inputs (Allen, Wilson and Ertz against 3.58, 3.04 and 0.55) are not
committed: only the salary file, the entries, the QB depth package and the earlier
scores are. The mechanism is therefore proved on fixtures shaped like the card's cases.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import replace
from datetime import timedelta
from pathlib import Path

import pytest

from nfl_dfs.participation import (
    WORKLOAD_REDISTRIBUTION_DOES_NOT_ESTABLISH,
    WORKLOAD_REDISTRIBUTION_RULE,
    WORKLOAD_REDISTRIBUTION_VERSION,
    ParticipationError,
    build_participation_contract,
    mark_declared_allocations,
    redistribute_opportunity,
    redistribute_vacated_workload,
)

from nfl_dfs.dk import parse_salaries as parse_salaries_classic

from . import test_classic_prior_review as classic
from . import test_qb_depth_roles as qb_depth
from .test_participation import _POOL, _SHARES, _model, _slate
from .test_prior_review_profile import AS_OF as SHOWDOWN_AS_OF, _prepared_run
from .test_prior_selection import _entries_bytes, _splits_bytes

FIELDS = (
    "qb_attempt_share",
    "carry_share",
    "target_share",
    "rushing_td_share",
    "receiving_td_share",
)
LEAD, BACKUP, THIRD = "SEA|RB|Sea Lead RB", "SEA|RB|Sea Backup RB", "SEA|RB|Sea Third RB"
ALPHA_WR, HURT_WR, TE = "SEA|WR|Sea Alpha WR", "SEA|WR|Sea Hurt WR", "SEA|TE|Sea TE"
SEA_QB = "SEA|QB|Sea QB"


def _pool(statuses: dict[str, str] | None = None, extra=()):
    """The shared pool with every status cleared, then `statuses` applied by name."""

    statuses = statuses or {}
    return tuple(
        (team, position, name, statuses.get(name, ""), salary)
        for team, position, name, _status, salary in _POOL
    ) + tuple(extra)


def _setup(tmp_path, statuses=None, extra=(), **contract_kwargs):
    slate = _slate(tmp_path, _pool(statuses, extra))
    model = _model(tmp_path, slate)
    contract = build_participation_contract(slate, **contract_kwargs)
    return slate, model, contract


def _by_person(model):
    return {player.underlying_id: player for player in model.players}


def _team_totals(model):
    totals: dict[tuple[str, str], float] = {}
    for player in model.players:
        for field in FIELDS:
            key = (player.team, field)
            totals[key] = totals.get(key, 0.0) + float(getattr(player, field))
    return totals


def _dk_ids(slate, name):
    return [player.dk_id for player in slate.players if player.name == name]


def _unchanged(before, after, people):
    for person in people:
        for field in FIELDS:
            assert getattr(after[person], field) == pytest.approx(getattr(before[person], field)), (
                person, field,
            )


# --------------------------------------------------------------------------- #
# Where the share goes
# --------------------------------------------------------------------------- #


def test_the_room_inherits_what_the_out_starter_vacates_in_proportion_to_prior(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    # Nobody leaves the model: the QB depth and role resolvers need every row.
    assert set(after) == set(before)
    assert out.workload_redistribution == WORKLOAD_REDISTRIBUTION_VERSION
    assert after[LEAD].carry_share == 0.0
    assert after[BACKUP].carry_share > before[BACKUP].carry_share
    assert after[THIRD].carry_share > before[THIRD].carry_share
    gain_backup = after[BACKUP].carry_share - before[BACKUP].carry_share
    gain_third = after[THIRD].carry_share - before[THIRD].carry_share
    assert gain_backup / gain_third == pytest.approx(
        before[BACKUP].carry_share / before[THIRD].carry_share
    )
    # The other room, the other team and the quarterback are untouched.
    _unchanged(before, after, [ALPHA_WR, TE, SEA_QB, "NE|RB|Lead RB", "NE|TE|Starting TE"])
    assert report["applied"] is True
    assert report["unallocated_by_team"] == {}


def test_a_wide_receiver_room_and_a_tight_end_room_move_the_same_way(tmp_path, monkeypatch):
    monkeypatch.setitem(
        _SHARES, "Sea Backup TE",
        dict(qb=0.0, carry=0.0, target=0.04, catch=0.6, ypt=6.5, rtd=0.0, ctd=0.04, cap=0.22),
    )
    extra = (("SEA", "TE", "Sea Backup TE", "", 1400),)
    slate, model, contract = _setup(tmp_path, {"Sea Hurt WR": "IR", "Sea TE": "OUT"}, extra)
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    backup_te = "SEA|TE|Sea Backup TE"
    assert after[ALPHA_WR].target_share > before[ALPHA_WR].target_share  # the WR room
    assert after[backup_te].target_share > before[backup_te].target_share * 3  # the TE room
    # Every target and receiving touchdown found a home. What the tight end held in carries and rushing
    # touchdowns, which his backup has no prior share of, stays unallocated rather than crossing a position.
    unallocated = report["unallocated_by_team"]["SEA"]
    assert set(unallocated) == {"carry_share", "rushing_td_share"}
    # The tight end's targets stayed in the tight end room, the receiver's in the receiver room.
    assert after[BACKUP].target_share == pytest.approx(before[BACKUP].target_share)
    assert after[THIRD].target_share == pytest.approx(before[THIRD].target_share)


def test_nothing_crosses_a_position_so_a_room_with_no_survivor_leaves_the_share_unallocated(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea TE": "OUT"})
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    others = [person for person in before if person.startswith("SEA|") and person != TE]
    _unchanged(before, after, others)
    # The share has nowhere to go inside the room, so it stays on the person who held it.
    assert after[TE].target_share == pytest.approx(before[TE].target_share)
    assert report["unallocated_by_team"]["SEA"]["target_share"] == pytest.approx(
        before[TE].target_share
    )
    assert report["cross_position_spill"] == "NEVER"
    # The wrapped function's own default is unchanged: it spills across the absorption set.
    spilled, _ = redistribute_opportunity(model, contract, redistribute=True)
    assert _by_person(spilled)[ALPHA_WR].target_share > before[ALPHA_WR].target_share


@pytest.mark.parametrize(
    "statuses",
    [
        {"Sea Lead RB": "OUT"},
        {"Sea Lead RB": "IR"},
        {"Sea Lead RB": "D"},
        {"Sea Lead RB": "OUT", "Sea Hurt WR": "IR"},
        {"Lead RB": "OUT", "Sea Lead RB": "OUT"},
    ],
    ids=["out", "ir", "doubtful", "two-rooms", "two-teams"],
)
def test_every_team_total_is_conserved_to_the_last_digit(tmp_path, statuses):
    slate, model, contract = _setup(tmp_path, statuses)
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _team_totals(model), _team_totals(out)
    assert set(before) == set(after)
    for key, value in before.items():
        assert after[key] == pytest.approx(value, abs=1e-12), key
    assert report["team_totals_conserved"] is True


def test_a_questionable_person_moves_nothing_and_stays_selectable(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea TE": "Q"})
    out, report = redistribute_vacated_workload(slate, model, contract)

    assert out is model  # nothing vacated, so scoring receives the model it always did
    assert out.workload_redistribution is None
    assert report["applied"] is False
    assert report["vacating_people"] == []
    assert report["moves"] == []
    assert TE in contract.selectable_people


def test_a_questionable_person_in_the_injured_room_is_still_a_survivor(tmp_path):
    """`Q` is never a trigger and never removed; like every survivor he shares the room."""

    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT", "Sea Backup RB": "Q"})
    out, report = redistribute_vacated_workload(slate, model, contract)
    after, before = _by_person(out), _by_person(model)

    assert BACKUP in contract.selectable_people
    assert [row["person"] for row in report["vacating_people"]] == [LEAD]
    assert after[BACKUP].carry_share > before[BACKUP].carry_share


# --------------------------------------------------------------------------- #
# The trigger
# --------------------------------------------------------------------------- #


def test_an_official_inactive_row_redistributes_the_same_way(tmp_path):
    slate, model, _ = _setup(tmp_path)
    official = _dk_ids(slate, "Sea Lead RB")
    # The run folds official rows into the operator exclusions; the function is told separately.
    contract = build_participation_contract(slate, operator_excluded_dk_ids=official)
    out, report = redistribute_vacated_workload(
        slate, model, contract, official_inactive_dk_ids=official
    )

    assert _by_person(out)[BACKUP].carry_share > _by_person(model)[BACKUP].carry_share
    (row,) = report["vacating_people"]
    assert row["person"] == LEAD and row["dk_status"] == ""
    assert row["triggered_by"] == ["OFFICIAL_INACTIVE_ROW"]
    assert report["official_activity_evidence"] == "NOT_IMPLIED"


def test_a_dk_status_and_an_official_row_are_both_named_when_both_hold(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    _out, report = redistribute_vacated_workload(
        slate, model, contract, official_inactive_dk_ids=_dk_ids(slate, "Sea Lead RB")
    )
    (row,) = report["vacating_people"]
    assert row["triggered_by"] == ["DK_STATUS_UNAVAILABLE", "OFFICIAL_INACTIVE_ROW"]
    assert row["dk_status"] == "OUT"


def test_an_operator_exclusion_alone_vacates_nothing(tmp_path):
    slate, model, _ = _setup(tmp_path)
    contract = build_participation_contract(
        slate, operator_excluded_dk_ids=_dk_ids(slate, "Sea Lead RB")
    )
    out, report = redistribute_vacated_workload(slate, model, contract)
    # A fade is a construction choice, not a statement that he is not playing.
    assert out is model
    assert report["applied"] is False
    assert report["operator_exclusions_that_move_nothing"] == [LEAD]


def test_an_operator_excluded_back_beside_a_real_vacancy_is_still_a_survivor(tmp_path):
    slate, model, _ = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    contract = build_participation_contract(
        slate, operator_excluded_dk_ids=_dk_ids(slate, "Sea Third RB")
    )
    out, report = redistribute_vacated_workload(slate, model, contract)
    # The model's view of reality is DK status plus official rows: the faded back still plays, so he
    # absorbs his proportional share and the selector alone decides not to pick him.
    assert _by_person(out)[THIRD].carry_share > _by_person(model)[THIRD].carry_share
    assert report["operator_exclusions_that_move_nothing"] == [THIRD]
    assert [row["person"] for row in report["vacating_people"]] == [LEAD]


def test_an_unrecognised_official_id_is_refused_in_prose(tmp_path):
    slate, model, contract = _setup(tmp_path)
    with pytest.raises(ParticipationError, match="outside the slate pool"):
        redistribute_vacated_workload(slate, model, contract, official_inactive_dk_ids=["1"])


# --------------------------------------------------------------------------- #
# Quarterbacks are the depth evidence's (R25, R36, Session 54), never this transformation's
# --------------------------------------------------------------------------- #


def _kc_pool():
    return tuple(
        (team, position, name, "OUT" if name == "KC Starter QB" else status, salary)
        for team, position, name, status, salary in qb_depth._POOL
    )


def _kc_setup(tmp_path, *, zero_weight_backup):
    slate, model, contract, _splits = qb_depth._setup(tmp_path, pool=_kc_pool())
    if zero_weight_backup:
        pinned = {"KC|QB|KC Starter QB": 1.0, "KC|QB|KC Backup QB": 0.0}
        model = replace(
            model,
            players=tuple(
                replace(p, qb_attempt_share=pinned[p.underlying_id])
                if p.underlying_id in pinned
                else p
                for p in model.players
            ),
        )
    return slate, model, contract


def _kc_qb_share(model):
    return {
        p.underlying_id: p.qb_attempt_share
        for p in model.players
        if p.team == "KC" and p.position == "QB"
    }


@pytest.mark.parametrize("with_package", [False, True], ids=["no-package", "depth-package"])
@pytest.mark.parametrize("zero_weight_backup", [False, True], ids=["weighted-backup", "zero-weight-backup"])
def test_an_out_starting_quarterback_vacates_nothing_and_the_depth_evidence_still_promotes(
    tmp_path, with_package, zero_weight_backup
):
    """Share-proportional inheritance would lift whichever backup has any prior share to the whole unit,
    played or not. Without depth evidence the attempt shares stay what they were; with it, R25 promotes."""

    from nfl_dfs.offensive_roles import resolve_offensive_roles
    from nfl_dfs.qb_depth_roles import resolve_qb_depth_roles

    slate, model, contract = _kc_setup(tmp_path, zero_weight_backup=zero_weight_backup)
    out, report = redistribute_vacated_workload(slate, model, contract)

    assert out is model
    assert report["applied"] is False and report["vacating_people"] == []
    (row,) = report["quarterbacks_left_to_the_depth_evidence"]
    assert row["person"] == "KC|QB|KC Starter QB" and row["dk_status"] == "OUT"
    assert _kc_qb_share(out) == _kc_qb_share(model)
    assert sum(_kc_qb_share(out).values()) == pytest.approx(1.0)
    if with_package:
        path = qb_depth._package(tmp_path / "pkg", slate, qb_depth._orders())
        resolution = resolve_qb_depth_roles(slate, out, contract, evidence_path=path, as_of=qb_depth.AS_OF)
        after = _kc_qb_share(resolution.model)
        assert resolution.report["starters_by_team"]["KC"] == "KC|QB|KC Backup QB"
        assert after["KC|QB|KC Backup QB"] == pytest.approx(1.0)  # R25's promoted starter gets the unit
        assert after["KC|QB|KC Starter QB"] == 0.0
        offense = resolve_offensive_roles(
            slate, resolution.model, contract, as_of=qb_depth.AS_OF,
            declared_starters=frozenset(resolution.report["starters_by_team"].values()),
            depth_evidence_sha256=resolution.evidence_sha256,
        )
        assert "KC|QB|KC Starter QB" in offense.excluded_people  # DK OUT stays out of the pool
        assert "KC|QB|KC Backup QB" not in offense.excluded_people


def test_a_quarterback_never_absorbs_a_share_either(tmp_path):
    """A running back vacates carries and rushing touchdowns; the quarterback's own are not touched."""

    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    out, _report = redistribute_vacated_workload(slate, model, contract)
    _unchanged(_by_person(model), _by_person(out), [SEA_QB, "NE|QB|Starter QB"])


def test_the_legacy_function_reports_a_share_no_survivor_can_hold_instead_of_dropping_it(tmp_path):
    """Session 60 fixed a silent drop in `redistribute_opportunity` itself: no eligible survivor for a field."""

    slate, model, contract = _setup(tmp_path, {"Sea QB": "OUT"})
    before = _by_person(model)
    reduced, report = redistribute_opportunity(model, contract, redistribute=True)
    assert SEA_QB not in _by_person(reduced)
    assert report["unallocated_by_team"]["SEA"]["qb_attempt_share"] == pytest.approx(before[SEA_QB].qb_attempt_share)
    assert report["unallocated_by_team_position"]["SEA"]["qb_attempt_share"]["QB"] == pytest.approx(
        before[SEA_QB].qb_attempt_share
    )


# --------------------------------------------------------------------------- #
# The old reading is an explicit, reported opt-out
# --------------------------------------------------------------------------- #


def test_disabled_keeps_every_survivor_at_his_prior_share_and_says_so(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    out, report = redistribute_vacated_workload(slate, model, contract, enabled=False)

    assert out is model
    assert report["applied"] is False
    assert report["rule"] == "NO_REDISTRIBUTION_SURVIVORS_KEEP_PRIOR_SHARES"
    assert report["moves"] == []
    assert report["unallocated_by_team"]["SEA"]["carry_share"] > 0.5
    assert report["transformation_version"] == WORKLOAD_REDISTRIBUTION_VERSION


# --------------------------------------------------------------------------- #
# The report
# --------------------------------------------------------------------------- #


def test_every_move_names_who_vacated_who_absorbed_and_the_shares_before_and_after(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT", "Sea Hurt WR": "IR"})
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    assert report["transformation_version"] == WORKLOAD_REDISTRIBUTION_VERSION
    assert report["does_not_establish"] == list(WORKLOAD_REDISTRIBUTION_DOES_NOT_ESTABLISH)
    assert "OFFICIAL_ACTIVE_STATUS" in report["does_not_establish"]
    assert report["trigger"] == "DK_STATUS_UNAVAILABLE_OR_SUPPLIED_OFFICIAL_INACTIVE_ROW"
    # Its own rule label: the registered V1 label includes the spill across positions this switches off.
    assert report["rule"] == WORKLOAD_REDISTRIBUTION_RULE
    assert report["rule"] != "PROPORTIONAL_TO_PRIOR_WITHIN_VACATING_POSITION_UNCAPPED_V1"
    assert report["moves_count"] == len(report["moves"]) > 0
    assert report["absorbing_people"] == sorted({BACKUP, THIRD, ALPHA_WR})
    carry = next(m for m in report["moves"] if m["field"] == "carry_share" and m["position"] == "RB")
    assert [v["person"] for v in carry["vacated"]] == [LEAD]
    assert carry["vacated"][0]["share"] == pytest.approx(before[LEAD].carry_share, abs=1e-6)
    assert carry["basis"] == "SAME_POSITION_ROOM"
    assert {a["person"] for a in carry["absorbed"]} == {BACKUP, THIRD}
    for move in report["moves"]:
        for entry in move["absorbed"]:
            field = move["field"]
            assert entry["before"] == pytest.approx(getattr(before[entry["person"]], field), abs=1e-6)
            assert entry["after"] == pytest.approx(getattr(after[entry["person"]], field), abs=1e-6)
            assert entry["gained"] == pytest.approx(entry["after"] - entry["before"], abs=2e-6)
        placed = sum(a["gained"] for a in move["absorbed"])
        assert placed + move["unallocated"] == pytest.approx(move["vacated_total"], abs=2e-6)


# --------------------------------------------------------------------------- #
# An unresolved current role never absorbs (the P1 gate must not be made moot by a model number)
# --------------------------------------------------------------------------- #


def test_a_survivor_whose_current_role_is_unresolved_takes_none_of_the_vacated_share(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    history = {THIRD: {"state": "CURRENT_ROLE_UNKNOWN", "incompatible_transfer": True, "current_team": "SEA"}}
    model = replace(model, offensive_history_by_person=history)
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    _unchanged(before, after, [THIRD])
    assert after[BACKUP].carry_share == pytest.approx(before[BACKUP].carry_share + before[LEAD].carry_share)
    assert report["not_absorbing_unresolved_current_role"] == [THIRD]
    assert report["absorbing_people"] == [BACKUP]
    assert report["team_totals_conserved"] is True


def test_a_survivor_whose_prior_row_is_not_pass_takes_none_of_it_either(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    model = replace(
        model,
        players=tuple(replace(p, evidence_state="UNKNOWN") if p.underlying_id == THIRD else p for p in model.players),
    )
    out, report = redistribute_vacated_workload(slate, model, contract)
    _unchanged(_by_person(model), _by_person(out), [THIRD])
    assert report["not_absorbing_unresolved_current_role"] == [THIRD]


def test_a_vacated_share_never_raises_an_unresolved_transfers_prior_or_flips_the_p1_gate(tmp_path):
    """Measured on the DEN@KC shape before the rule: the transfer behind a DK-OUT back went from
    8.6 to 24.2 prior points and left the material-role-change gate's exclusion."""

    from nfl_dfs.selection import select_prior_lineups

    pool = tuple(
        (team, position, name, "OUT" if name == "KC Committee RB" else status, salary)
        for team, position, name, status, salary in qb_depth._POOL
    )
    slate, model, contract, splits = qb_depth._setup(tmp_path, pool=pool, transfer=True)
    transfer = "KC|RB|KC Transfer RB"
    results = {}
    for enabled in (False, True):
        shaped, report = redistribute_vacated_workload(slate, model, contract, enabled=enabled)
        _lineups, scores, selection = select_prior_lineups(
            slate, shaped, splits, contract, count=3, as_of=qb_depth.AS_OF
        )
        results[enabled] = (scores, selection, report)

    off_scores, off_selection, _ = results[False]
    on_scores, on_selection, on_report = results[True]
    # The transfer is the only survivor in the room and may not absorb, so nothing moves at all.
    assert on_report["applied"] is False
    assert on_report["not_absorbing_unresolved_current_role"] == [transfer]
    assert on_report["unallocated_by_team"]["KC"]["carry_share"] > 0
    assert on_scores.by_person[transfer] == pytest.approx(off_scores.by_person[transfer])
    assert transfer in on_scores.offensive_role_resolution.excluded_people
    assert on_selection["offensive_roles"]["material_role_change_exclusions"] == (
        off_selection["offensive_roles"]["material_role_change_exclusions"]
    )
    assert len(on_selection["offensive_roles"]["material_role_change_exclusions"]) == 1


def test_a_move_on_a_team_with_a_declared_allocation_is_marked_never_dropped(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    _out, report = redistribute_vacated_workload(slate, model, contract)

    marked = mark_declared_allocations(report, ["SEA"])
    assert marked["superseded_by_declared_allocation"] == ["SEA"]
    assert len(marked["moves"]) == len(report["moves"]) > 0
    assert all(move["superseded_by_declared_allocation"] for move in marked["moves"])
    other = mark_declared_allocations(report, ["NE"])
    assert other["superseded_by_declared_allocation"] == []
    assert not any(move["superseded_by_declared_allocation"] for move in other["moves"])
    assert all("superseded_by_declared_allocation" not in move for move in report["moves"])  # input untouched


def test_two_vacators_at_one_position_with_nobody_left_each_keep_their_own_share(tmp_path):
    """The whole running back room is out: each keeps what he held, so the team's pool is exactly as it was."""

    slate, model, contract = _setup(
        tmp_path, {"Sea Lead RB": "OUT", "Sea Backup RB": "OUT", "Sea Third RB": "IR"}
    )
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    assert out is model and report["applied"] is False  # nobody could take anything
    for person in (LEAD, BACKUP, THIRD):
        assert after[person].carry_share == pytest.approx(before[person].carry_share)
    assert report["unallocated_by_team"]["SEA"]["carry_share"] == pytest.approx(
        sum(before[p].carry_share for p in (LEAD, BACKUP, THIRD))
    )
    assert report["team_totals_conserved"] is True


def test_two_vacators_at_one_position_split_what_the_survivor_takes_by_their_own_shares(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT", "Sea Backup RB": "IR"})
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    # The third back is the only survivor in the room: he takes both vacated carry shares.
    assert after[THIRD].carry_share == pytest.approx(
        before[THIRD].carry_share + before[LEAD].carry_share + before[BACKUP].carry_share
    )
    move = next(m for m in report["moves"] if m["field"] == "carry_share" and m["position"] == "RB")
    assert {v["person"] for v in move["vacated"]} == {LEAD, BACKUP}
    assert move["vacated_total"] == pytest.approx(before[LEAD].carry_share + before[BACKUP].carry_share)


def test_two_vacators_each_keep_their_own_share_of_a_field_nobody_in_the_room_can_take(tmp_path):
    """One field moves (so the model changes) while another has two vacators and no taker.

    Each vacator keeps his own share of the stranded field, which is what keeps the team's pool exact.
    """

    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT", "Sea Backup RB": "IR"})
    model = replace(
        model,
        players=tuple(
            replace(p, rushing_td_share=0.0) if p.underlying_id == THIRD else p for p in model.players
        ),
    )
    out, report = redistribute_vacated_workload(slate, model, contract)
    before, after = _by_person(model), _by_person(out)

    assert report["applied"] is True  # the carries moved to the third back
    assert after[THIRD].carry_share > before[THIRD].carry_share
    for person in (LEAD, BACKUP):
        assert after[person].rushing_td_share == pytest.approx(before[person].rushing_td_share)
        assert after[person].carry_share == 0.0
    assert report["unallocated_by_team"]["SEA"]["rushing_td_share"] == pytest.approx(
        before[LEAD].rushing_td_share + before[BACKUP].rushing_td_share
    )
    assert _team_totals(out)[("SEA", "rushing_td_share")] == pytest.approx(
        _team_totals(model)[("SEA", "rushing_td_share")], abs=1e-12
    )


def test_a_vacancy_that_nobody_takes_is_not_reported_as_applied(tmp_path):
    """Only an unavailable kicker: no share moves, so the model and the assumption text stay as they were."""

    slate, model, contract = _setup(tmp_path, {"Sea Kicker": "OUT"})
    out, report = redistribute_vacated_workload(slate, model, contract)
    assert out is model and out.workload_redistribution is None
    assert report["applied"] is False and report["rule"] == WORKLOAD_REDISTRIBUTION_RULE


def test_the_report_and_the_model_are_deterministic(tmp_path):
    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT", "Sea Hurt WR": "IR"})
    first = redistribute_vacated_workload(slate, model, contract)
    second = redistribute_vacated_workload(slate, model, contract)
    assert first[0] == second[0]
    assert json.dumps(first[1], sort_keys=True) == json.dumps(second[1], sort_keys=True)


# --------------------------------------------------------------------------- #
# The redistributed model is the one scoring receives, at both call sites
# --------------------------------------------------------------------------- #


def _spy_selection(monkeypatch, module):
    """Record what the redistribution received and returned and what selection was handed.

    The first `select_prior_lineups` call is the run's own. The control is the same call on the
    model as it was before the redistribution, so a score that rose did so because of the
    transformation and for no other reason.
    """

    real_select = module.select_prior_lineups
    real_redistribute = module.redistribute_vacated_workload
    seen: dict[str, object] = {"input": None, "redistributed": None, "model": None,
                               "scores": None, "control": None, "report": None}

    def redistribute(slate, model, contract, **kwargs):
        result = real_redistribute(slate, model, contract, **kwargs)
        seen.update(input=model, redistributed=result[0], report=result[1])
        return result

    def select(slate, model, splits, contract, **kwargs):
        result = real_select(slate, model, splits, contract, **kwargs)
        if seen["model"] is None:
            seen.update(model=model, scores=result[1])
            seen["control"] = real_select(slate, seen["input"], splits, contract, **kwargs)[1]
        return result

    monkeypatch.setattr(module, "redistribute_vacated_workload", redistribute)
    monkeypatch.setattr(module, "select_prior_lineups", select)
    return seen


def _assert_the_redistributed_model_reached_scoring(seen, absorber):
    assert seen["model"] is seen["redistributed"], "selection was handed a model other than the redistributed one"
    assert seen["model"].workload_redistribution == WORKLOAD_REDISTRIBUTION_VERSION
    assert seen["scores"].by_person[absorber] > seen["control"].by_person[absorber]


def _select_cli(tmp_path, capsys, name, *extra):
    from nfl_dfs.cli import main

    (tmp_path / "DKEntries.csv").write_bytes(_entries_bytes())
    (tmp_path / "stats_team_week.csv").write_bytes(_splits_bytes())
    code = main(
        [
            "select",
            "--salaries", str(tmp_path / "DKSalaries.csv"),
            "--entries", str(tmp_path / "DKEntries.csv"),
            "--team-projections", str(tmp_path / "team_projections.csv"),
            "--player-opportunities", str(tmp_path / "player_opportunities.csv"),
            "--team-splits", str(tmp_path / "stats_team_week.csv"),
            "--prior-season", "2025",
            "--as-of", "2026-09-09T12:00:00Z",
            "--output-dir", str(tmp_path / name),
            *extra,
        ]
    )
    assert code == 0
    return json.loads(capsys.readouterr().out)


def test_select_scores_the_redistributed_model_and_no_redistribute_is_the_reported_opt_out(
    tmp_path, capsys, monkeypatch
):
    from nfl_dfs import cli

    slate, _model_unused, _contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})  # writes the salary and opportunity files
    seen = _spy_selection(monkeypatch, cli)
    result = _select_cli(tmp_path, capsys, "default")

    _assert_the_redistributed_model_reached_scoring(seen, BACKUP)
    assert result["redistribution"]["applied"] is True
    assert result["redistribution"]["transformation_version"] == WORKLOAD_REDISTRIBUTION_VERSION
    assert result["status"] == "DO_NOT_UPLOAD"
    assert result["MODEL_STATUS"] == "PRIOR_ONLY"
    assert result["RELEASE_DECISION"] == "DO_NOT_UPLOAD"
    # The DK-OUT person is in no lineup.
    out_ids = set(_dk_ids(slate, "Sea Lead RB"))
    assert not out_ids & {dk for lineup in result["lineups"] for dk in lineup["roster"]}

    opted_out = _select_cli(tmp_path, capsys, "optout", "--no-redistribute")
    assert opted_out["redistribution"]["applied"] is False
    assert opted_out["redistribution"]["rule"] == "NO_REDISTRIBUTION_SURVIVORS_KEEP_PRIOR_SHARES"


def test_the_showdown_run_scores_the_redistributed_model_and_reports_every_move(tmp_path, monkeypatch):
    from nfl_dfs import prior_review
    from nfl_dfs.prior_review import run_prior_review

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=SHOWDOWN_AS_OF + timedelta(hours=6)
    )
    seen = _spy_selection(monkeypatch, prior_review)
    outcome = run_prior_review(
        salary_csv=salary_path, entry_csv=entry_path, label="ne-sea", as_of=SHOWDOWN_AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out",
        prior_package_dir=package_dir, project=project,
    )
    assert not outcome.blocked, outcome.blockers
    _assert_the_redistributed_model_reached_scoring(seen, BACKUP)
    selection = outcome.reports["selection"]
    assert selection["RELEASE_DECISION"] == "DO_NOT_UPLOAD" and selection["MODEL_STATUS"] == "PRIOR_ONLY"
    from nfl_dfs.dk import parse_salaries

    out_ids = {
        dk for name in ("Sea Lead RB", "Sea Hurt WR") for dk in _dk_ids(parse_salaries(salary_path), name)
    }
    assert not out_ids & {dk for lineup in selection["lineups"] for dk in lineup["roster"]}
    block = selection["redistribution"]
    assert block["applied"] is True and block["moves_count"] > 0
    assert {row["person"] for row in block["vacating_people"]} == {LEAD, HURT_WR}
    assert {row["dk_status"] for row in block["vacating_people"]} == {"OUT", "IR"}
    # The same block rides in the coverage the reviews read, and the stale sentence is gone.
    assert selection["pool_coverage"]["workload_redistribution"] == block
    assert "not reassigned" not in selection["pool_coverage"]["note"]


def _classic_fixture_with_out(tmp_path, monkeypatch, out_name):
    original = classic._salary_bytes

    def with_status(**kwargs):
        rows = list(csv.reader(io.StringIO(original(**kwargs).decode("utf-8"))))
        header = rows[0]
        name_at, status_at = header.index("Name"), header.index("Status")
        for row in rows[1:]:
            if row[name_at] == out_name:
                row[status_at] = "OUT"
        buffer = io.StringIO(newline="")
        csv.writer(buffer, lineterminator="\n").writerows(rows)
        return buffer.getvalue().encode("utf-8")

    monkeypatch.setattr(classic, "_salary_bytes", with_status)
    return classic._fixture(
        tmp_path, depth={"QB": 1, "RB": 2, "WR": 3, "TE": 1, "DST": 1}
    )


def test_the_classic_run_scores_the_redistributed_model_too(tmp_path, monkeypatch):
    from nfl_dfs import prior_review
    from nfl_dfs.prior_review import run_prior_review

    salary, entry, package, _role, _status, _inactive = _classic_fixture_with_out(
        tmp_path, monkeypatch, "NE RB One"
    )
    seen = _spy_selection(monkeypatch, prior_review)
    outcome = run_prior_review(
        salary_csv=salary, entry_csv=entry, label="classic-out", as_of=classic.AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out", prior_package_dir=package,
    )
    assert not outcome.blocked, outcome.blockers
    _assert_the_redistributed_model_reached_scoring(seen, "NE|RB|NE RB Two")
    block = seen["report"]
    assert [row["person"] for row in block["vacating_people"]] == ["NE|RB|NE RB One"]
    out_id = next(p.dk_id for p in parse_salaries_classic(salary).players if p.name == "NE RB One")
    rosters = json.loads(Path(outcome.artifacts["selection_report"]).read_text(encoding="utf-8"))["assignments_by_entry_id"]
    assert rosters and not any(out_id in roster for roster in rosters.values())
    coverage = json.loads(Path(outcome.artifacts["complete_slate_coverage"]).read_text(encoding="utf-8"))
    assert coverage["pool_coverage"]["workload_redistribution"]["moves_count"] == block["moves_count"]
    assert coverage["MODEL_STATUS"] == "PRIOR_ONLY" and coverage["RELEASE_DECISION"] == "DO_NOT_UPLOAD"


def test_a_declared_allocation_still_wins_for_its_team(tmp_path):
    """The evidence replaces every recipient's shares after the transformation has run."""

    from nfl_dfs.offensive_roles import FIELDS as ROLE_FIELDS, OFFENSE, resolve_offensive_roles
    from . import test_offensive_roles as roles

    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    redistributed, _report = redistribute_vacated_workload(slate, model, contract)
    declared = {}
    recipients = []
    for player in redistributed.players:
        if player.team != "SEA" or player.position not in OFFENSE:
            continue
        shares = {f: getattr(player, f) for f in ROLE_FIELDS}
        if player.underlying_id == BACKUP:
            shares["carry_share"] -= 0.10
        if player.underlying_id == THIRD:
            shares["carry_share"] += 0.10
        declared[player.underlying_id] = shares["carry_share"]
        recipients.append(
            {**roles._binding(slate, player.underlying_id), "shares": shares, "receiving_efficiency": None}
        )
    declaration = {
        "team": "SEA", "game_id": slate.games[0].game_id,
        "totals": dict.fromkeys(ROLE_FIELDS, 1.0), "unallocated": dict.fromkeys(ROLE_FIELDS, 0.0),
        "recipients": recipients,
    }
    path = roles._package(tmp_path / "roles", slate, [declaration])
    result = resolve_offensive_roles(slate, redistributed, contract, evidence_path=path, as_of=roles.AS_OF)

    after = _by_person(result.model)
    assert after[THIRD].carry_share == pytest.approx(declared[THIRD])
    assert after[THIRD].carry_share != pytest.approx(_by_person(redistributed)[THIRD].carry_share)


def test_the_offensive_role_report_stops_saying_the_vacated_volume_stays_unallocated(tmp_path):
    from nfl_dfs.offensive_roles import resolve_offensive_roles

    slate, model, contract = _setup(tmp_path, {"Sea Lead RB": "OUT"})
    redistributed, _report = redistribute_vacated_workload(slate, model, contract)
    plain = resolve_offensive_roles(slate, model, contract, as_of=SHOWDOWN_AS_OF)
    moved = resolve_offensive_roles(slate, redistributed, contract, as_of=SHOWDOWN_AS_OF)

    assert plain.report["assumptions"][0] == "HISTORY_IS_UNCONFIRMED; VACATED_VOLUME_REMAINS_UNALLOCATED"
    assert moved.report["assumptions"][0] == (
        "HISTORY_IS_UNCONFIRMED; "
        "VACATED_VOLUME_OF_UNAVAILABLE_PEOPLE_REDISTRIBUTED_WITHIN_POSITION_ROOM_BY_RULE_V1; "
        "RESIDUAL_REMAINS_UNALLOCATED"
    )
    # What the resolver reports as unallocated is now only what nobody in the room could take.
    assert plain.report["unallocated_by_team"]["SEA"]["carry_share"] == pytest.approx(0.55, abs=0.01)
    assert moved.report["unallocated_by_team"]["SEA"]["carry_share"] == pytest.approx(0.0, abs=1e-9)


def test_the_readable_review_carries_the_moves_and_names_the_version(tmp_path):
    from .test_readable_review import _create

    *_unused, readable, _blockers = _create(tmp_path)
    block = readable.data["pool_coverage"]["workload_redistribution"]
    assert readable.data["schema_version"] == "prior_only_readable_review_sd5_v3"
    assert block["applied"] is True and block["moves_count"] > 0
    assert block["transformation_version"] == WORKLOAD_REDISTRIBUTION_VERSION
    html_text = Path(readable.html_path).read_text(encoding="utf-8")
    assert "Injury-room redistribution" in html_text
    assert "Sea Lead RB" in html_text and "Sea Backup RB" in html_text
    assert "Volume left unallocated after the injury-room redistribution" in html_text
    assert "Prior-season volume left unallocated" not in html_text
    assert "OFFICIAL_ACTIVE_STATUS" in html_text  # the does-not-establish text travels with it


def test_the_classic_c3_review_carries_the_moves_and_marks_those_a_declared_allocation_replaced(
    tmp_path, monkeypatch
):
    """The third `prior_review` exit. C3 binds a declared role allocation, which wins for its team."""

    from types import SimpleNamespace

    from nfl_dfs.classic_portfolio_policy import (
        classic_portfolio_policy_template,
        validate_classic_portfolio_policy_bytes,
        write_normalized_classic_portfolio_policy,
    )
    from nfl_dfs.dk import parse_entries, parse_salaries
    from nfl_dfs.hashing import sha256_file
    from nfl_dfs.prior_review import run_prior_review
    from nfl_dfs.projection import build_projection_package

    original_roles = classic._role_evidence

    def roles_without_the_out_back(root, slate, salary_hash):
        # An unavailable person may be left out of a declared allocation; the survivors carry it all.
        visible = [p for p in slate.players if p.name != "NE RB One"]
        return original_roles(root, SimpleNamespace(players=visible, games=slate.games), salary_hash)

    monkeypatch.setattr(classic, "_role_evidence", roles_without_the_out_back)
    salary, entry, package, role, _status, _inactive = _classic_fixture_with_out(
        tmp_path / "fixture", monkeypatch, "NE RB One"
    )
    slate = parse_salaries(salary)
    template = parse_entries(entry)
    entry_ids = tuple(item.entry_id for item in template.authorizations)
    policy_path = tmp_path / "classic_policy.json"
    policy_path.write_text(
        json.dumps(
            classic_portfolio_policy_template(slate, entry_ids, entry_sha256=template.raw_hash), indent=2
        ),
        encoding="utf-8",
    )
    validated = validate_classic_portfolio_policy_bytes(
        policy_path.read_bytes(), slate=slate, entry_ids=entry_ids, entry_sha256=template.raw_hash
    )
    assert validated.valid and validated.policy is not None
    normalized = write_normalized_classic_portfolio_policy(
        tmp_path / "classic_policy.normalized.json", validated.policy
    )
    outcome = run_prior_review(
        salary_csv=salary, entry_csv=entry, label="c3-out", as_of=classic.AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out", prior_package_dir=package,
        build_priors=True, offensive_role_evidence_json=role,
        portfolio_policy=validated.policy,
        portfolio_policy_source_path=policy_path,
        portfolio_policy_source_sha256=sha256_file(policy_path),
        portfolio_policy_normalized_path=normalized,
        portfolio_policy_normalized_sha256=sha256_file(normalized),
        project=build_projection_package,
    )
    assert not outcome.blocked, outcome.blockers
    readable = json.loads(Path(outcome.artifacts["readable_review_json"]).read_text(encoding="utf-8"))
    assert readable["schema_version"] == "prior_only_readable_review_classic_c3_v4"
    block = readable["pool_coverage"]["workload_redistribution"]
    assert block["applied"] is True
    assert [row["person"] for row in block["vacating_people"]] == ["NE|RB|NE RB One"]
    # The declared allocation replaced NE's shares after the transformation ran, and the report says so.
    assert block["superseded_by_declared_allocation"] == ["NE"]
    assert all(move["superseded_by_declared_allocation"] for move in block["moves"])
    html_text = Path(outcome.artifacts["readable_review_html"]).read_text(encoding="utf-8")
    assert "Injury-room redistribution" in html_text and "NE RB One" in html_text
    assert "A declared allocation replaced this team&#x27;s shares" in html_text
    assert readable["truths"]["RELEASE_DECISION"] == "DO_NOT_UPLOAD"


def test_a_supplied_official_inactive_row_redistributes_through_the_run(tmp_path, monkeypatch):
    """The run folds official rows into the operator exclusions; they must still reach the transformation."""

    from nfl_dfs import prior_review
    from nfl_dfs.prior_review import run_prior_review

    salary, entry, package, _role, status, _inactive = classic._fixture(
        tmp_path, depth={"QB": 1, "RB": 2, "WR": 3, "TE": 1, "DST": 1}
    )
    dk_id = next(p.dk_id for p in parse_salaries_classic(salary).players if p.name == "NE RB One")
    rows = list(csv.reader(io.StringIO(status.read_text(encoding="utf-8"))))
    for row in rows[1:]:
        if row[1] == dk_id:
            row[2] = "INACTIVE"
    buffer = io.StringIO(newline="")
    csv.writer(buffer, lineterminator="\n").writerows(rows)
    status.write_text(buffer.getvalue(), encoding="utf-8", newline="")

    seen = _spy_selection(monkeypatch, prior_review)
    outcome = run_prior_review(
        salary_csv=salary, entry_csv=entry, label="classic-official", as_of=classic.AS_OF,
        run_root=tmp_path / "run", output_root=tmp_path / "out", prior_package_dir=package,
        official_status_csv=status,
    )
    assert not outcome.blocked, outcome.blockers
    (row,) = seen["report"]["vacating_people"]
    assert row["person"] == "NE|RB|NE RB One" and row["dk_status"] == ""
    assert row["triggered_by"] == ["OFFICIAL_INACTIVE_ROW"]
    _assert_the_redistributed_model_reached_scoring(seen, "NE|RB|NE RB Two")
    rosters = json.loads(Path(outcome.artifacts["selection_report"]).read_text(encoding="utf-8"))["assignments_by_entry_id"]
    assert rosters and not any(dk_id in roster for roster in rosters.values())


def test_no_review_surface_still_says_a_vacated_share_is_not_reassigned(tmp_path):
    """The JSON limitations, the HTML and the workbook all describe what the transformation did."""

    import copy

    from openpyxl import load_workbook

    from nfl_dfs.workbook import create_cowork_status_workbook

    from .test_readable_review import TRUTHS, _create

    *_unused, readable, _blockers = _create(tmp_path)
    assert readable.data["pool_coverage"]["workload_redistribution"]["applied"] is True
    assert "not reassigned" not in json.dumps(readable.data)
    assert "not reassigned" not in Path(readable.html_path).read_text(encoding="utf-8")
    observation = [
        row["observation"] for row in readable.data["evidence_observations"] if row["category"] == "unallocated_volume"
    ]
    assert all("after the injury-room redistribution" in text for text in observation)

    workbook_path = tmp_path / "review.xlsx"
    create_cowork_status_workbook(
        output_path=workbook_path, run_values={"RUN_LABEL": "x"}, blockers=(),
        report_path=tmp_path / "cowork_run.json", truth_values=TRUTHS,
        readable_review=copy.deepcopy(readable.data),
    )
    cells = [
        str(cell.value)
        for sheet in load_workbook(workbook_path, data_only=False).worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert not any("not reassigned" in text for text in cells)
    assert any("Volume left unallocated after the injury-room redistribution" in text for text in cells)
