"""Session 50c: the contest-assignment step through `run-slate`, the Classic exits.

One test per `prior_review` Classic exit (C1 sequential, C2 policy with C3's package
and export, and the C2 subset policy with a C1 fill), each showing the diversified
assignment reaching `assignments.csv`, the independent audit and the review, with
every release truth unchanged. Every total is recomputed from the delivered bytes.
The baseline is in `test_contest_assignment_baseline.py`.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import pytest

from nfl_dfs import contest_assignment as ca
from nfl_dfs.dk import parse_entries, parse_salaries
from nfl_dfs.entry_groups import prefilled_cell_id

from .test_entry_groups import (
    _cells,
    _edit,
    _named,
    _policy,
    _raw_lines,
    _readable,
    _run_slate,
    _salary_top,
    _subset_policy,
)

CONTEST_COLUMN = 2
ENTRIES = tuple(f"91000000{index}" for index in range(1, 7))
# The C1 solver order is NE, SEA, DAL, PHI, NE, NE by quarterback, and the C2 joint
# solve holds three NE lineups, so this template puts the repeats together: entries
# 1 and 5 share a contest, and 6 and 2, and 3 and 4.
THREE_CONTESTS = {
    ENTRIES[0]: "111", ENTRIES[4]: "111", ENTRIES[5]: "222", ENTRIES[1]: "222",
    ENTRIES[2]: "333", ENTRIES[3]: "333",
}
C1 = "run-slate:prior_review:CLASSIC_C1"
C2 = "run-slate:prior_review:CLASSIC"


def _classic_run(tmp_path, monkeypatch, *, run_id, contests=THREE_CONTESTS, entries=6, policy=None,
                 cells=None, wrap_step=None):
    """`run-slate` on the synthetic Classic fixture with `contests` set by Entry ID."""

    def edit(entry, slate):
        _edit(entry, cells=cells(slate) if callable(cells) else cells,
              columns={eid: {CONTEST_COLUMN: cid} for eid, cid in contests.items()})

    if wrap_step is not None:
        real_step = ca.apply_step
        monkeypatch.setattr(ca, "apply_step", lambda **kwargs: wrap_step(real_step(**kwargs), kwargs))
    return _run_slate(tmp_path, monkeypatch, run_id=run_id, entries=entries, edit=edit, policy=policy)


def _slate(entries_path: Path):
    return parse_salaries(entries_path.parent / "salary.csv")


def _assignment_rosters(report) -> dict[str, tuple[str, ...]]:
    lines = Path(report["prior_review_artifacts"]["assignments"]).read_text(encoding="utf-8").splitlines()
    return {line.split(",")[0]: tuple(line.split(",")[1:10]) for line in lines[1:] if line}


def _selection_rosters(report) -> list[tuple[str, ...]]:
    return [tuple(row["roster"]) for row in report["prior_review_reports"]["selection"]["lineups"]]


def _recompute(entries_path: Path, delivered_path: Path):
    """Per-contest statistics from the delivered bytes alone."""

    people = ca.people_from_slate(_slate(entries_path).players)
    contest_of = {e.entry_id: e.contest_id for e in parse_entries(entries_path).authorizations}
    delivered = _cells(delivered_path)
    # a row the template filled keeps its own text ("Name (ID)"), so read the IDs out of the cells
    return ca.statistics_from_rosters(
        [(eid, contest_of[eid], tuple(prefilled_cell_id(cell) for cell in roster))
         for eid, roster in delivered.items() if any(roster)],
        mode=ca.MODE_CLASSIC, people=people)


def _solver_order_total(entries_path: Path, report, order_entries) -> float:
    people = ca.people_from_slate(_slate(entries_path).players)
    contest_of = {e.entry_id: e.contest_id for e in parse_entries(entries_path).authorizations}
    stats = ca.statistics_from_rosters(
        [(eid, contest_of[eid], roster) for eid, roster in zip(order_entries, _selection_rosters(report))],
        mode=ca.MODE_CLASSIC, people=people)
    return _total(stats)


def _total(stats) -> float:
    return sum(float(row["score"]) for row in stats.values())


def _assert_truths_unchanged(report):
    truths = report["release_truths"]
    assert (truths["MODEL_STATUS"], truths["RELEASE_DECISION"]) == ("PRIOR_ONLY", "DO_NOT_UPLOAD")
    assert report["DELIVERY_STATE"] == "DELIVERABLE"


def _no_contest_repeats_a_qb(stats, quarterbacks=None):
    """No contest holds two entries with the same quarterback, unless the portfolio forces one.

    `quarterbacks` (Session 39) is the delivered rows' quarterback by row. The six lineups
    tie on prior points in this fixture, so which six the joint solve returns is a
    tie-break: when one quarterback leads more lineups than there are contests, pigeonhole
    forces `lead - contests` repeats and the step must leave exactly that many, no more.
    """

    assert stats, "the fixture holds contests of two or more entries"
    forced = 0
    if quarterbacks:
        lead = max(quarterbacks.count(person) for person in set(quarterbacks))
        forced = max(0, lead - len(stats))
    repeats = sum(reading["entries"] - reading["distinct_key_people"] for reading in stats.values())
    assert repeats == forced, (forced, stats)
    if not forced:
        for contest, reading in stats.items():
            assert reading["distinct_key_people"] == reading["entries"], (contest, reading)


# ---------------------------------------------------------------- the three exits


def test_c1_sequential_diversifies_and_the_block_lives_in_the_selection_report(tmp_path, monkeypatch):
    code, report, entries, _root = _classic_run(tmp_path, monkeypatch, run_id="cd-c1")
    assert code == 0, report["blockers"]
    assert report["latest_deliverable"]["producer"] == C1
    _assert_truths_unchanged(report)
    delivered = Path(report["latest_deliverable"]["path"])
    assert delivered.name.startswith("DK_REVIEW_ENTRY_C1_")

    step = report["prior_review_reports"]["contest_assignment"]
    assert step["contest_assignment_version"] == "within_contest_diversity_v1"
    assert step["status"] == "IMPROVED" and step["moved_rows"] > 0 and step["pools"] == {"all": 6}
    assert "PAYOUT_OR_CONTEST_WORTH" in step["does_not_establish"]

    # assignments.csv, the delivered CSV and the selection hold the same lineups, entry for entry
    assignment = _assignment_rosters(report)
    cells = _cells(delivered)
    assert {eid: tuple(cells[eid]) for eid in ENTRIES} == assignment
    assert sorted(assignment.values()) == sorted(_selection_rosters(report))

    # the delivered bytes bear the improvement out, recomputed here and not read from the step
    after = _recompute(entries, delivered)
    assert _total(after) < _solver_order_total(entries, report, ENTRIES)
    assert f"{_total(after):.6f}" == step["total_score_after"]
    _no_contest_repeats_a_qb(after)
    assert max(row["worst_pair_shared_people"] for row in after.values()) <= 6

    # C1 has no readable review: its block is the step's report in the selection report and the run result
    reports = report["prior_review_reports"]
    assert "readable_review" not in reports
    assert reports["selection"]["contest_assignment"] == step
    assert {k: v for k, v in step["contests_after"].items()} == after
    export = report["export"]["c1_export"]
    assert export["status"] == "PASS" and export["contest_assignment"]["status"] == "PASS"
    assert export["contest_assignment"]["findings"] == []
    assert export["contest_assignment"]["contests"] == after


def test_c2_policy_and_c3_diversify_and_the_review_reconciles(tmp_path, monkeypatch):
    code, report, entries, _root = _classic_run(tmp_path, monkeypatch, run_id="cd-c2", policy=_policy)
    assert code == 0, report["blockers"]
    assert report["latest_deliverable"]["producer"] == C2
    _assert_truths_unchanged(report)
    delivered = Path(report["latest_deliverable"]["path"])

    step = report["prior_review_reports"]["contest_assignment"]
    assert step["status"] == "IMPROVED" and step["pools"] == {"bound": 6}
    assignment = _assignment_rosters(report)
    cells = _cells(delivered)
    assert {eid: tuple(cells[eid]) for eid in ENTRIES} == assignment
    assert sorted(assignment.values()) == sorted(_selection_rosters(report))
    # classic_assignment.json holds the policy's rows: it follows the permutation
    artifact = json.loads(Path(report["prior_review_artifacts"]["classic_assignment"]).read_text("utf-8"))
    assert {row["entry_id"]: tuple(row["roster"]) for row in artifact["entry_assignments"]} == assignment

    after = _recompute(entries, delivered)
    assert _total(after) < _solver_order_total(entries, report, ENTRIES)
    assert f"{_total(after):.6f}" == step["total_score_after"]
    slate = _slate(entries)
    by_id = {player.dk_id: player for player in slate.players}
    _no_contest_repeats_a_qb(after, [by_id[roster[0]].underlying_id for roster in assignment.values()])

    # the C2 audit recomputed the step's claims from the exact assignments.csv bytes
    audit = report["prior_review_reports"]["classic_portfolio_audit"]
    assert audit["status"] == "PASS"
    reading = audit["contest_assignment"]
    assert reading["status"] == "PASS" and reading["findings"] == [] and reading["contests"] == after
    # the selector's overlap labels follow the lineups to their new entries
    slate = _slate(entries)
    person = {p.dk_id: p.underlying_id for p in slate.players}
    labelled = report["prior_review_reports"]["selection"]["selection"]["pairwise_person_overlap"]
    assert labelled and {(row["entry_id_a"], row["entry_id_b"]): row["people"] for row in labelled} == {
        (left, right): len({person[d] for d in cells[left]} & {person[d] for d in cells[right]})
        for left in ENTRIES for right in ENTRIES if left < right}

    # C3: JSON, HTML and the workbook carry the block, reconciled
    readable = _readable(report)
    block = readable["contest_assignment"]
    assert block["reported_statistics_match"] is True and block["status"] == "IMPROVED"
    assert block["basis"] == ca.REVIEW_BASIS and block["key_person"] == "QUARTERBACK"
    assert {row["contest_id"] for row in block["contests"]} == {"111", "222", "333"}
    assert all(row["score_after"] <= row["score_before"] for row in block["contests"])
    assert block["total_score_after"] == step["total_score_after"]
    assert readable["reconciliation"]["status"] == "PASS"
    html = Path(report["prior_review_artifacts"]["readable_review_html"]).read_text("utf-8")
    assert "Within each contest" in html and "Distinct QBs" in html

    from openpyxl import load_workbook

    sheet = load_workbook(report["review_workbook"])["Exposure"]
    values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert any("within_contest_diversity_v1" in value for value in values)
    assert {"111", "222", "333"} <= set(values)


def test_a_subset_policy_keeps_bound_and_fill_lineups_in_their_own_rows(tmp_path, monkeypatch):
    """Five rows: 2 prefilled, 3 and 5 bound by C2, 1 and 4 filled by C1 (`_classic_subset_run`'s shape)."""

    contests = {ENTRIES[0]: "111", ENTRIES[1]: "111", ENTRIES[2]: "111", ENTRIES[3]: "222", ENTRIES[4]: "222"}
    holder: dict[str, object] = {}

    def prefill(slate):
        holder["prefilled"] = _salary_top(slate, 1)[0]
        return {ENTRIES[1]: _named(slate, holder["prefilled"])}

    code, report, entries, _root = _classic_run(
        tmp_path, monkeypatch, run_id="cd-subset", contests=contests, entries=5,
        policy=_subset_policy(1, 3), cells=prefill)
    assert code == 0, report["blockers"]
    _assert_truths_unchanged(report)
    bound = (ENTRIES[2], ENTRIES[4])
    fill = (ENTRIES[0], ENTRIES[3])
    step = report["prior_review_reports"]["contest_assignment"]
    assert step["pools"] == {"bound": 2, "fill": 2} and step["fixed_entry_ids"] == [ENTRIES[1]]
    selected = _selection_rosters(report)
    policy_lineups, fill_lineups = selected[:2], selected[2:]
    assignment = _assignment_rosters(report)
    assert sorted(assignment[eid] for eid in bound) == sorted(policy_lineups)
    assert sorted(assignment[eid] for eid in fill) == sorted(fill_lineups)
    delivered = Path(report["latest_deliverable"]["path"])
    assert _raw_lines(delivered)[ENTRIES[1]] == _raw_lines(entries)[ENTRIES[1]]  # the prefilled row is byte-identical
    assert tuple(_cells(delivered)[ENTRIES[1]]) != ()
    after = _recompute(entries, delivered)
    assert after["111"]["entries"] == 3 and after["222"]["entries"] == 2  # the prefilled row counts in its contest
    assert f"{_total(after):.6f}" == step["total_score_after"]
    audit = report["prior_review_reports"]["classic_portfolio_audit"]
    assert audit["status"] == "PASS" and audit["contest_assignment"]["status"] == "PASS"
    assert audit["contest_assignment"]["contests"] == after
    readable = _readable(report)
    assert readable["contest_assignment"]["reported_statistics_match"] is True
    assert readable["reconciliation"]["status"] == "PASS"
    assert readable["unbound_rows"]["entry_ids"] == list(fill)


@pytest.mark.parametrize("policy", [None, _policy], ids=["c1", "c2-c3"])
def test_the_same_inputs_give_the_same_delivered_bytes(tmp_path, monkeypatch, policy):
    first = _classic_run(tmp_path / "a", monkeypatch, run_id="cd-det", policy=policy)[1]
    second = _classic_run(tmp_path / "b", monkeypatch, run_id="cd-det", policy=policy)[1]
    for name in ("assignments", "contest_assignment"):
        assert Path(first["prior_review_artifacts"][name]).read_bytes() == Path(
            second["prior_review_artifacts"][name]).read_bytes(), name
    assert Path(first["latest_deliverable"]["path"]).read_bytes() == Path(
        second["latest_deliverable"]["path"]).read_bytes()
    record = json.loads(Path(first["prior_review_artifacts"]["contest_assignment"]).read_text("utf-8"))
    assert record["schema_version"] == "nfl_contest_assignment_step_v1" and "seconds" not in record


FIVE_ROWS = {ENTRIES[0]: "111", ENTRIES[1]: "111", ENTRIES[2]: "222", ENTRIES[3]: "222", ENTRIES[4]: "222"}


@pytest.mark.parametrize(
    ("policy", "entries", "contests", "bound"),
    [(None, 6, THREE_CONTESTS, ()), (_policy, 6, THREE_CONTESTS, ENTRIES),
     (_subset_policy(1, 3), 5, FIVE_ROWS, (ENTRIES[1], ENTRIES[3]))],
    ids=["c1", "c2-c3", "subset"])
def test_a_failed_step_leaves_the_solver_order_and_ships_with_the_gap_named(
    tmp_path, monkeypatch, policy, entries, contests, bound
):
    def broken(*_args, **_kwargs):
        raise RuntimeError("injected")

    monkeypatch.setattr(ca, "diversify", broken)
    code, report, _entries, _root = _classic_run(
        tmp_path, monkeypatch, run_id="cd-failed", policy=policy, entries=entries, contests=contests)
    assert code == 0, report["blockers"]
    _assert_truths_unchanged(report)
    step = report["prior_review_reports"]["contest_assignment"]
    assert step["status"] == "FAILED" and step["error"] == "CONTEST_ASSIGNMENT_STEP_FAILED:RuntimeError:injected"
    # the file is the solver's order: the policy's lineups fill its rows in template order, then the fill's
    selected = _selection_rosters(report)
    rows = [eid for eid in ENTRIES[:entries] if not bound or eid in bound]
    rows += [eid for eid in ENTRIES[:entries] if bound and eid not in bound]
    assert _assignment_rosters(report) == {eid: roster for eid, roster in zip(rows, selected)}
    limitations = {item["code"]: item for item in report["release_truths"]["delivery_limitations"]}
    assert limitations["CONTEST_ASSIGNMENT_STEP_FAILED"]["class"] == "P"
    if policy is not None:
        assert _readable(report)["contest_assignment"]["status"] == "FAILED"
        assert report["prior_review_reports"]["classic_portfolio_audit"]["status"] == "PASS"


def test_a_single_contest_run_is_not_applicable(tmp_path, monkeypatch):
    contests = {eid: "111" for eid in ENTRIES}
    code, report, entries, _root = _classic_run(tmp_path, monkeypatch, run_id="cd-one", contests=contests)
    assert code == 0, report["blockers"]
    step = report["prior_review_reports"]["contest_assignment"]
    assert step["status"] in {"NOT_APPLICABLE", "UNCHANGED"} and step["moved_rows"] == 0
    assert list(_assignment_rosters(report).values()) == _selection_rosters(report)


# ---------------------------------------------------------------- the audits refuse


def _capture_c2_audit(tmp_path, monkeypatch, *, run_id, contests, entries=6, policy=_policy, cells=None):
    from nfl_dfs import classic_portfolio as cp

    seen: dict[str, object] = {}
    real_audit = cp.audit_classic_portfolio

    def capture(**kwargs):
        seen["kwargs"] = kwargs
        return real_audit(**kwargs)

    monkeypatch.setattr("nfl_dfs.prior_review.audit_classic_portfolio", capture)
    code, report, entries_path, _root = _classic_run(
        tmp_path, monkeypatch, run_id=run_id, contests=contests, entries=entries, policy=policy, cells=cells)
    assert code == 0, report["blockers"]
    return real_audit, seen["kwargs"], report, entries_path


def _retext(kwargs, lines_edit):
    """The kwargs with `assignments.csv` rewritten by `lines_edit(lines)` and rehashed."""

    lines = kwargs["assignment_csv_bytes"].decode("utf-8").splitlines(keepends=True)
    raw = "".join(lines_edit(lines)).encode("utf-8")
    return {**kwargs, "assignment_csv_bytes": raw, "expected_assignment_csv_sha256": hashlib.sha256(raw).hexdigest()}


def test_the_c2_audit_refuses_a_lineup_the_selection_never_held(tmp_path, monkeypatch):
    real_audit, kwargs, report, entries = _capture_c2_audit(
        tmp_path, monkeypatch, run_id="cd-audit", contests=THREE_CONTESTS)
    assert real_audit(**kwargs).passed
    slate = _slate(entries)
    held = {tuple(roster) for roster in _assignment_rosters(report).values()}
    outsider = next(roster for roster in _salary_top(slate, 40) if tuple(roster) not in held)

    def swap(lines):
        cells = lines[1].rstrip("\r\n").split(",")
        return [lines[0], ",".join([cells[0], *outsider]) + "\n", *lines[2:]]

    audit = real_audit(**_retext(kwargs, swap))
    assert not audit.passed
    assert any(item.startswith("CONTEST_ASSIGNMENT_MULTISET_CHANGED:bound") for item in audit.problems), audit.problems
    assert any(item.startswith("CLASSIC_AUDIT_ASSIGNMENT_ARTIFACT_MISMATCH") for item in audit.problems)


def test_the_c2_audit_holds_the_csv_to_the_assignment_json_row_for_row(tmp_path, monkeypatch):
    """Two rows exchanged in the CSV alone keep the multiset and change the file: only a row check sees it."""

    real_audit, kwargs, _report, _entries = _capture_c2_audit(
        tmp_path, monkeypatch, run_id="cd-rows", contests=THREE_CONTESTS)

    def exchange(lines):
        first, second = lines[1].rstrip("\r\n").split(","), lines[2].rstrip("\r\n").split(",")
        assert first[1:] != second[1:]
        return [lines[0], ",".join([first[0], *second[1:]]) + "\n", ",".join([second[0], *first[1:]]) + "\n", *lines[3:]]

    audit = real_audit(**_retext(kwargs, exchange))
    assert not audit.passed
    assert any(item.startswith("CLASSIC_AUDIT_ASSIGNMENT_ARTIFACT_MISMATCH") for item in audit.problems)
    assert not any(item.startswith("CONTEST_ASSIGNMENT_MULTISET_CHANGED") for item in audit.problems)


def test_the_c2_audit_refuses_a_csv_that_is_not_the_hashed_one(tmp_path, monkeypatch):
    real_audit, kwargs, _report, _entries = _capture_c2_audit(
        tmp_path, monkeypatch, run_id="cd-hash", contests=THREE_CONTESTS)
    raw = kwargs["assignment_csv_bytes"] + b"\n"
    audit = real_audit(**{**kwargs, "assignment_csv_bytes": raw})
    assert not audit.passed
    assert any(item.startswith("CLASSIC_AUDIT_ASSIGNMENT_CSV_SHA256_MISMATCH") for item in audit.problems)


def test_the_c2_audit_names_a_lying_report_p_and_still_passes(tmp_path, monkeypatch):
    real_audit, kwargs, _report, _entries = _capture_c2_audit(
        tmp_path, monkeypatch, run_id="cd-lie", contests=THREE_CONTESTS)
    claim = kwargs["contest_assignment"]
    lie = json.loads(json.dumps(claim.reported_after))
    first = sorted(lie)[0]
    lie[first]["worst_pair_shared_people"] += 5
    audit = real_audit(**{**kwargs, "contest_assignment": replace(claim, reported_after=lie)})
    assert audit.passed, audit.problems
    reading = audit.contest_assignment
    assert reading["status"] == "STATS_MISMATCH" and reading["findings"] == ["CONTEST_ASSIGNMENT_STATS_MISMATCH"]
    assert reading["contests"] == real_audit(**kwargs).contest_assignment["contests"]


def test_the_c2_audit_refuses_a_moved_filled_row_read_from_the_template_bytes(tmp_path, monkeypatch):
    contests = {ENTRIES[0]: "111", ENTRIES[1]: "111", ENTRIES[2]: "111", ENTRIES[3]: "222", ENTRIES[4]: "222"}
    holder: dict[str, object] = {}

    def prefill(slate):
        holder["prefilled"] = _salary_top(slate, 1)[0]
        return {ENTRIES[1]: _named(slate, holder["prefilled"])}

    real_audit, kwargs, _report, _entries = _capture_c2_audit(
        tmp_path, monkeypatch, run_id="cd-fixed", contests=contests, entries=5,
        policy=_subset_policy(1, 3), cells=prefill)
    claim = kwargs["contest_assignment"]
    assert real_audit(**kwargs).contest_assignment["status"] == "PASS"
    rows = tuple(replace(row, roster=tuple(reversed(row.roster))) if row.entry_id == ENTRIES[1] else row
                 for row in claim.rows)
    audit = real_audit(**{**kwargs, "contest_assignment": replace(claim, rows=rows)})
    assert not audit.passed
    assert any(item.startswith(f"CONTEST_ASSIGNMENT_FIXED_ROW_MOVED:{ENTRIES[1]}") for item in audit.problems)
    # a subset policy leaves the fill pool to C1: both pools' multisets are audited
    swapped_fill = _retext(kwargs, lambda lines: [
        lines[0], *[",".join([line.split(",")[0], *holder["prefilled"]]) + "\n"
                    if line.startswith(ENTRIES[0]) else line for line in lines[1:]]])
    audit = real_audit(**swapped_fill)
    assert any(item.startswith("CONTEST_ASSIGNMENT_MULTISET_CHANGED:fill") for item in audit.problems), audit.problems


# ---------------------------------------------------------------- full-run refusals


def _is_baseline(kwargs) -> bool:
    """The baseline calls the step with no restarts; these tests wrap only the model path."""

    return kwargs.get("restarts") == 0


def _falsify_report(step, kwargs):
    if _is_baseline(kwargs):
        return step
    report = json.loads(json.dumps(step.report))
    first = sorted(report["contests_after"])[0]
    report["contests_after"][first]["worst_pair_shared_people"] += 5
    return ca.StepOutcome(step.assignments, report, step.claim, step.failure)


def _falsify_claim(step, kwargs):
    if _is_baseline(kwargs):
        return step
    lie = json.loads(json.dumps(step.claim.reported_after))
    lie[sorted(lie)[0]]["worst_pair_shared_people"] += 5
    return ca.StepOutcome(step.assignments, step.report, replace(step.claim, reported_after=lie), step.failure)


def _outsider_swap(tmp_path):
    """A wrapper putting a legal, distinct lineup the selection never held in row one."""

    def swap(step, kwargs):
        if _is_baseline(kwargs):
            return step
        slate = parse_salaries(tmp_path / "fixture" / "DKSalaries.csv")
        held = set(step.assignments.values())
        outsider = next(tuple(roster) for roster in _salary_top(slate, 40) if tuple(roster) not in held)
        return ca.StepOutcome({**step.assignments, ENTRIES[0]: outsider}, step.report, step.claim, step.failure)

    return swap


def _limitation_classes(report) -> dict[str, str]:
    return {item["code"]: item["class"] for item in report["release_truths"]["delivery_limitations"]}


@pytest.mark.parametrize("policy", [None, _policy], ids=["c1", "c2-c3"])
def test_a_lineup_outside_the_selection_is_refused_and_the_baseline_stays(tmp_path, monkeypatch, policy):
    """C1's export audit and C2's independent audit each hold the file to the step's multiset (`V`)."""

    code, report, _entries, _root = _classic_run(
        tmp_path, monkeypatch, run_id="cd-outsider", policy=policy, wrap_step=_outsider_swap(tmp_path))
    assert code != 0
    text = ";".join(report["blockers"])
    assert "CONTEST_ASSIGNMENT_MULTISET_CHANGED" in text, text
    assert report["latest_deliverable"]["producer"] == "run-slate:baseline"
    assert report["release_truths"]["RELEASE_DECISION"] == "DO_NOT_UPLOAD"
    assert not list(Path(report["latest_deliverable"]["path"]).parent.parent.glob("review/DK_REVIEW_ENTRY_*.csv"))


def test_c1_names_a_lying_report_p_and_still_ships(tmp_path, monkeypatch):
    code, report, _entries, _root = _classic_run(tmp_path, monkeypatch, run_id="cd-c1-lie", wrap_step=_falsify_claim)
    assert code == 0, report["blockers"]
    assert report["latest_deliverable"]["producer"] == C1
    assert _limitation_classes(report)["CONTEST_ASSIGNMENT_STATS_MISMATCH"] == "P"
    assert report["export"]["c1_export"]["contest_assignment"]["status"] == "STATS_MISMATCH"
    _assert_truths_unchanged(report)


def test_c3_names_a_lying_report_p_and_keeps_the_csv(tmp_path, monkeypatch):
    """DISPLAY_RECONCILIATION covers the block: reported numbers the bytes do not bear out are `P`."""

    code, report, _entries, _root = _classic_run(
        tmp_path, monkeypatch, run_id="cd-c3-lie", policy=_policy, wrap_step=_falsify_report)
    assert code == 2 or report["blockers"]
    assert "CONTEST_ASSIGNMENT_STATS_MISMATCH" in ";".join(report["blockers"])
    assert _limitation_classes(report)["CONTEST_ASSIGNMENT_STATS_MISMATCH"] == "P"
    assert report["latest_deliverable"]["producer"] == C2  # a `P` discrepancy never withholds the file
    _assert_truths_unchanged(report)


# ---------------------------------------------------------------- C3 replayed from a real run


def _capture_c3(tmp_path, monkeypatch, *, run_id):
    from nfl_dfs import prior_review
    from nfl_dfs.classic_review import create_classic_review_package

    captured: dict[str, object] = {}

    def record(**kwargs):
        captured.update({key: dict(value) if isinstance(value, dict) else value for key, value in kwargs.items()})
        return create_classic_review_package(**kwargs)

    monkeypatch.setattr(prior_review, "create_classic_review_package", record)
    code, report, entries, _root = _classic_run(tmp_path, monkeypatch, run_id=run_id, policy=_policy)
    assert code == 0 and captured, report["blockers"]
    return captured, report, entries


def _replay(tmp_path, captured, name, *, artifact=None, raw=None, keep_expected=False):
    """C3 again from the run's own call, one bound artifact replaced by `raw`."""

    from nfl_dfs.classic_review import create_classic_review_package
    from nfl_dfs.hashing import sha256_file

    folder = tmp_path / name
    folder.mkdir()
    artifacts = dict(captured["artifacts"])
    hashes = dict(captured["expected_hashes"])
    if artifact is not None:
        path = folder / Path(artifacts[artifact]).name
        path.write_bytes(raw)
        artifacts[artifact] = str(path)
        if not keep_expected:
            hashes[artifact] = sha256_file(path)
    values = dict(captured, artifacts=artifacts, expected_hashes=hashes,
                  output_path=folder / "review" / Path(captured["output_path"]).name,
                  output_dir=folder / "review")
    return values, create_classic_review_package


def _canonical(payload) -> bytes:
    return (json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str) + "\n").encode()


def test_c3_replayed_from_the_same_bytes_writes_the_same_block(tmp_path, monkeypatch):
    captured, report, _entries = _capture_c3(tmp_path / "run", monkeypatch, run_id="cd-replay")
    values, create = _replay(tmp_path, captured, "clean")
    again = create(**values)
    assert again.data["contest_assignment"] == _readable(report)["contest_assignment"]
    original = Path(captured["output_path"]).parent
    # the readable JSON and HTML name their artifacts relative to the folder, so the block is compared above
    for name in (Path(captured["output_path"]).name, "classic_review_export_audit.json"):
        assert (Path(values["output_dir"]) / name).read_bytes() == (original / name).read_bytes(), name


def test_c3_refuses_a_changed_step_record_and_an_assignment_csv_that_disagrees(tmp_path, monkeypatch):
    from nfl_dfs.classic_review import ClassicReviewError

    captured, _report, _entries = _capture_c3(tmp_path / "run", monkeypatch, run_id="cd-c3-mutate")
    step_path = Path(captured["artifacts"]["contest_assignment"])

    # a changed byte in the step's record, its hash unchanged: the intake checkpoint refuses it
    values, create = _replay(tmp_path, captured, "step-byte", artifact="contest_assignment",
                             raw=step_path.read_bytes().replace(b"IMPROVED", b"UNCHANGED"), keep_expected=True)
    with pytest.raises(ClassicReviewError):
        create(**values)
    assert not list((tmp_path / "step-byte" / "review").glob("DK_REVIEW_ENTRY_*.csv"))

    # two rows exchanged in assignments.csv, rehashed: the file and the selection record disagree
    csv_path = Path(captured["artifacts"]["assignments"])
    lines = csv_path.read_text(encoding="utf-8").splitlines(keepends=True)
    first, second = lines[1].rstrip("\n").split(","), lines[2].rstrip("\n").split(",")
    exchanged = "".join([lines[0], ",".join([first[0], *second[1:]]) + "\n",
                         ",".join([second[0], *first[1:]]) + "\n", *lines[3:]])
    values, create = _replay(tmp_path, captured, "csv-rows", artifact="assignments", raw=exchanged.encode())
    with pytest.raises(ClassicReviewError, match="CLASSIC_C3_ASSIGNMENT_CSV_DISAGREEMENT"):
        create(**values)
    assert not list((tmp_path / "csv-rows" / "review").glob("DK_REVIEW_ENTRY_*.csv"))


def test_c3_keeps_its_export_and_names_p_when_the_step_record_lies(tmp_path, monkeypatch):
    from nfl_dfs.classic_review import ClassicReviewPresentationError

    captured, _report, _entries = _capture_c3(tmp_path / "run", monkeypatch, run_id="cd-c3-record")
    record = json.loads(Path(captured["artifacts"]["contest_assignment"]).read_text(encoding="utf-8"))
    first = sorted(record["contests_after"])[0]
    record["contests_after"][first]["mean_shared_people"] += 1
    values, create = _replay(tmp_path, captured, "lying-record", artifact="contest_assignment", raw=_canonical(record))
    with pytest.raises(ClassicReviewPresentationError, match="CONTEST_ASSIGNMENT_STATS_MISMATCH") as caught:
        create(**values)
    assert Path(caught.value.export_path).is_file()  # the audited export stays; only the JSON and HTML are removed
    assert not list((tmp_path / "lying-record" / "review").glob("prior_only_readable_review.*"))


def test_c3_refuses_a_lineup_swapped_in_the_selection_record(tmp_path, monkeypatch):
    from nfl_dfs.classic_review import ClassicReviewError

    captured, _report, entries = _capture_c3(tmp_path / "run", monkeypatch, run_id="cd-c3-swap")
    selection = json.loads(Path(captured["artifacts"]["selection_report"]).read_text(encoding="utf-8"),
                           parse_float=Decimal)
    held = {tuple(roster) for roster in selection["assignments_by_entry_id"].values()}
    outsider = next(tuple(r) for r in _salary_top(_slate(entries), 40) if tuple(r) not in held)
    selection["assignments_by_entry_id"][ENTRIES[0]] = list(outsider)
    values, create = _replay(tmp_path, captured, "swap", artifact="selection_report", raw=_canonical(selection))
    with pytest.raises(ClassicReviewError, match="CLASSIC_C3_SELECTION_ASSIGNMENT_MAP_DISAGREEMENT"):
        create(**values)
    assert not list((tmp_path / "swap" / "review").glob("DK_REVIEW_ENTRY_*.csv"))
