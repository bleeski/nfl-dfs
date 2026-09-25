"""SD3: exact, deterministic Showdown portfolio-policy contract coverage."""

from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nfl_dfs.cowork import CoworkInputError, CoworkRunRequest, required_next_inputs
from nfl_dfs.dk import parse_entries, parse_salaries
from nfl_dfs.hashing import sha256_bytes
from nfl_dfs.portfolio_policy import (
    FRACTION_UNIT,
    POLICY_SCHEMA_VERSION,
    canonical_showdown_lineup_identity,
    portfolio_policy_template,
    salary_person_bindings,
    summarize_showdown_portfolio,
    validate_portfolio_policy_bytes,
    validate_portfolio_policy_file,
)

from .test_participation import _salary_bytes, _slate


def _reference(slate, index: int = 0) -> dict[str, str]:
    return salary_person_bindings(slate)[index].as_mapping()


def _document(slate, entry_ids=("900000001", "900000002"), **controls):
    return portfolio_policy_template(slate, entry_ids, controls=controls)


def _validate(slate, document, entry_ids=("900000001", "900000002"), **kwargs):
    raw = json.dumps(document, ensure_ascii=False).encode("utf-8")
    return validate_portfolio_policy_bytes(
        raw, slate=slate, entry_ids=entry_ids, **kwargs
    )


def _codes(validation) -> set[str]:
    return {problem.code for problem in validation.problems}


@pytest.mark.parametrize(
    ("entries", "fraction", "expected"),
    [
        (("1", "2"), 0.49, 0),
        (("1", "2"), 0.50, 1),
        (("1", "2"), 1.00, 2),
        (("1", "2", "3"), 0.66, 1),
        (("1", "2", "3"), 0.67, 2),
    ],
)
def test_exact_decimal_floor_boundaries(entries, fraction, expected, tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    reference = _reference(slate)
    document = _document(
        slate,
        entries,
        max_combined_person_exposure={
            "default_fraction": None,
            "overrides": [{**reference, "fraction": fraction}],
        },
    )

    validation = _validate(slate, document, entries)

    assert validation.valid, validation.blockers()
    limit = next(
        item
        for item in validation.policy.effective_limits
        if item.person.underlying_id == reference["underlying_id"]
    )
    assert limit.combined_max_entries == expected
    if expected == 0:
        assert limit.captain_max_entries == 0


def test_omitted_default_zero_full_and_override_precedence(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    target = _reference(slate, 0)
    zero_captain = _reference(slate, 1)
    document = _document(
        slate,
        max_combined_person_exposure={
            "default_fraction": 0.50,
            "overrides": [{**target, "fraction": 1.00}],
        },
        max_captain_exposure={
            "default_fraction": 0.50,
            "overrides": [{**zero_captain, "fraction": 0}],
        },
    )

    validation = _validate(slate, document)

    assert validation.valid, validation.blockers()
    limits = {item.person.underlying_id: item for item in validation.policy.effective_limits}
    assert limits[target["underlying_id"]].combined_source == "PERSON_OVERRIDE"
    assert limits[target["underlying_id"]].combined_max_entries == 2
    assert limits[zero_captain["underlying_id"]].combined_max_entries == 1
    assert limits[zero_captain["underlying_id"]].captain_max_entries == 0
    omitted = _validate(slate, _document(slate))
    assert omitted.valid
    assert all(item.combined_max_entries == 2 for item in omitted.policy.effective_limits)
    assert all(item.captain_max_entries == 2 for item in omitted.policy.effective_limits)


def test_captain_subset_contradiction_is_tightened_and_reported(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    target = _reference(slate)
    document = _document(
        slate,
        max_combined_person_exposure={
            "default_fraction": None,
            "overrides": [{**target, "fraction": 0.50}],
        },
        max_captain_exposure={
            "default_fraction": None,
            "overrides": [{**target, "fraction": 1.00}],
        },
    )

    validation = _validate(slate, document)

    assert validation.valid
    limit = next(
        item
        for item in validation.policy.effective_limits
        if item.person.underlying_id == target["underlying_id"]
    )
    assert limit.declared_captain_max_entries == 2
    assert limit.captain_max_entries == limit.combined_max_entries == 1
    assert {finding.code for finding in validation.findings} == {
        "PORTFOLIO_POLICY_CAPTAIN_TIGHTENED_BY_COMBINED"
    }


@pytest.mark.parametrize(
    ("value", "code"),
    [
        (True, "PORTFOLIO_POLICY_FRACTION_TYPE_INVALID"),
        ("0.5", "PORTFOLIO_POLICY_FRACTION_TYPE_INVALID"),
        (-0.01, "PORTFOLIO_POLICY_FRACTION_OUT_OF_RANGE"),
        (1.01, "PORTFOLIO_POLICY_FRACTION_OUT_OF_RANGE"),
        (50, "PORTFOLIO_POLICY_FRACTION_OUT_OF_RANGE"),
    ],
)
def test_malformed_fraction_types_and_ranges_are_rejected(
    value, code, tmp_path: Path
) -> None:
    slate = _slate(tmp_path)
    document = _document(
        slate,
        max_combined_person_exposure={
            "default_fraction": value,
            "overrides": [],
        },
    )
    validation = _validate(slate, document)
    assert code in _codes(validation)


def test_nonfinite_number_and_ambiguous_unit_are_rejected(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    document = _document(slate)
    raw = json.dumps(document).replace(
        '"default_fraction": null', '"default_fraction": NaN', 1
    ).encode("utf-8")
    nonfinite = validate_portfolio_policy_bytes(
        raw, slate=slate, entry_ids=("900000001", "900000002")
    )
    assert "PORTFOLIO_POLICY_FRACTION_NONFINITE" in _codes(nonfinite)

    document["controls"]["fraction_unit"] = "PERCENT_0_TO_100"
    unit = _validate(slate, document)
    assert "PORTFOLIO_POLICY_FRACTION_UNIT_INVALID" in _codes(unit)


def test_unknown_ids_conflicting_roles_and_duplicate_overrides_are_rejected(
    tmp_path: Path,
) -> None:
    slate = _slate(tmp_path)
    first = _reference(slate, 0)
    second = _reference(slate, 1)
    unknown = _document(
        slate,
        max_captain_exposure={
            "default_fraction": None,
            "overrides": [{**first, "cpt_dk_id": "99999999", "fraction": 0.5}],
        },
    )
    assert "PORTFOLIO_POLICY_UNKNOWN_DK_ID" in _codes(_validate(slate, unknown))

    conflicting = _document(
        slate,
        max_captain_exposure={
            "default_fraction": None,
            "overrides": [
                {**first, "flex_dk_id": second["flex_dk_id"], "fraction": 0.5}
            ],
        },
    )
    assert "PORTFOLIO_POLICY_CONFLICTING_ROLE_IDENTITY" in _codes(
        _validate(slate, conflicting)
    )

    duplicate = _document(
        slate,
        max_captain_exposure={
            "default_fraction": None,
            "overrides": [
                {**first, "fraction": 0.5},
                {**first, "fraction": 0.6},
            ],
        },
    )
    assert "PORTFOLIO_POLICY_DUPLICATE_OVERRIDE" in _codes(
        _validate(slate, duplicate)
    )


def test_duplicate_reordered_and_unknown_entry_bindings_are_rejected(tmp_path: Path) -> None:
    """Session 11b: a subset in template order now validates; everything else stays refused."""

    slate = _slate(tmp_path)
    duplicate = _document(slate, ("900000001", "900000001"))
    validation = _validate(slate, duplicate)
    assert "PORTFOLIO_POLICY_DUPLICATE_ENTRY_ID" in _codes(validation)
    assert "PORTFOLIO_POLICY_ENTRY_ID_BINDING_MISMATCH" in _codes(validation)

    for bound in (("900000002", "900000001"), ("900000009",), ()):
        assert "PORTFOLIO_POLICY_ENTRY_ID_BINDING_MISMATCH" in _codes(
            _validate(slate, _document(slate, bound))
        ), bound


def test_complete_person_identity_binding_detects_mutated_identity(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    document = _document(slate)
    document["bindings"]["person_identities"][0]["underlying_id"] = "NE|QB|Wrong"
    validation = _validate(slate, document)
    assert "PORTFOLIO_POLICY_UNKNOWN_PERSON_ID" in _codes(validation)
    assert "PORTFOLIO_POLICY_PERSON_IDENTITY_COVERAGE_MISMATCH" in _codes(validation)


def test_policy_hash_is_stable_under_nonsemantic_order_and_tracks_source_bytes(
    tmp_path: Path,
) -> None:
    slate = _slate(tmp_path)
    first = _reference(slate, 0)
    second = _reference(slate, 1)
    controls = {
        "max_combined_person_exposure": {
            "default_fraction": 1,
            "overrides": [
                {**first, "fraction": 0.5},
                {**second, "fraction": 0.75},
            ],
        }
    }
    document = _document(slate, **controls)
    raw_a = json.dumps(document, sort_keys=True).encode("utf-8")
    reordered = json.loads(json.dumps(document))
    reordered["bindings"]["person_identities"].reverse()
    reordered["controls"]["max_combined_person_exposure"]["overrides"].reverse()
    raw_b = (json.dumps(reordered, indent=2) + "\n").encode("utf-8")

    first_validation = validate_portfolio_policy_bytes(
        raw_a, slate=slate, entry_ids=("900000001", "900000002")
    )
    second_validation = validate_portfolio_policy_bytes(
        raw_b, slate=slate, entry_ids=("900000001", "900000002")
    )

    assert first_validation.valid and second_validation.valid
    assert first_validation.source_sha256 != second_validation.source_sha256
    assert first_validation.normalized_sha256 == second_validation.normalized_sha256
    assert first_validation.policy.canonical_bytes() == second_validation.policy.canonical_bytes()


def test_expected_hash_detects_policy_input_mutation(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    path = tmp_path / "policy.json"
    path.write_text(json.dumps(_document(slate)), encoding="utf-8")
    original = sha256_bytes(path.read_bytes())
    path.write_text(path.read_text(encoding="utf-8") + "\n", encoding="utf-8")

    validation = validate_portfolio_policy_file(
        path,
        slate=slate,
        entry_ids=("900000001", "900000002"),
        expected_sha256=original,
    )

    assert "PORTFOLIO_POLICY_INPUT_MUTATED" in _codes(validation)


def test_policy_and_source_exclusions_are_exact_and_stricter_than_caps(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    policy_excluded = _reference(slate, 0)
    source_excluded = _reference(slate, 1)
    document = _document(
        slate,
        excluded_people=[policy_excluded],
        max_combined_person_exposure={"default_fraction": 1, "overrides": []},
        max_captain_exposure={"default_fraction": 1, "overrides": []},
    )

    validation = _validate(
        slate,
        document,
        externally_excluded_people=(source_excluded["underlying_id"],),
    )

    assert validation.valid, validation.blockers()
    limits = {item.person.underlying_id: item for item in validation.policy.effective_limits}
    assert limits[policy_excluded["underlying_id"]].combined_max_entries == 0
    assert limits[source_excluded["underlying_id"]].captain_max_entries == 0
    assert limits[source_excluded["underlying_id"]].exclusion_source == (
        "SOURCE_OR_PARTICIPATION_PRECEDENCE"
    )
    assert sum(
        finding.code == "PORTFOLIO_POLICY_EXCLUSION_TAKES_PRECEDENCE"
        for finding in validation.findings
    ) == 2


def test_necessary_capacity_checks_are_named_without_solver_claims(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    zero = _document(
        slate,
        max_combined_person_exposure={"default_fraction": 0, "overrides": []},
        max_captain_exposure={"default_fraction": 0, "overrides": []},
    )
    zero_validation = _validate(slate, zero)
    assert {
        "PORTFOLIO_POLICY_PERSON_CAPACITY_INSUFFICIENT",
        "PORTFOLIO_POLICY_COMBINED_EXPOSURE_CAPACITY_INSUFFICIENT",
        "PORTFOLIO_POLICY_CAPTAIN_CAPACITY_INSUFFICIENT",
    } <= _codes(zero_validation)
    assert all("solver" not in problem.message.casefold() for problem in zero_validation.problems)

    people = salary_person_bindings(slate)
    keep = {item.underlying_id for item in people[:3]} | {
        item.underlying_id for item in people[-4:]
    }
    excluded = [item.as_mapping() for item in people if item.underlying_id not in keep]
    overlap = _document(
        slate,
        excluded_people=excluded,
        max_pairwise_person_overlap=4,
    )
    overlap_validation = _validate(slate, overlap)
    assert "PORTFOLIO_POLICY_PAIRWISE_OVERLAP_CAPACITY_INSUFFICIENT" in _codes(
        overlap_validation
    )


def _legal_roster(slate):
    by_person = {
        binding.underlying_id: binding for binding in salary_person_bindings(slate)
    }
    flex_rows = sorted(
        (row for row in slate.players if row.role == "FLEX"),
        key=lambda row: (row.salary, row.dk_id),
    )
    for captain in sorted(
        (row for row in slate.players if row.role == "CPT"),
        key=lambda row: (row.salary, row.dk_id),
    ):
        flex = [row for row in flex_rows if row.underlying_id != captain.underlying_id]
        for start in range(max(1, len(flex) - 4)):
            selected = flex[start : start + 5]
            if len(selected) != 5:
                continue
            roster = (captain.dk_id, *(row.dk_id for row in selected))
            identity, _people, _captain = canonical_showdown_lineup_identity(slate, roster)
            assert identity
            return roster, by_person
    raise AssertionError("fixture did not contain a legal lineup")


def test_canonical_uniqueness_overlap_and_role_agnostic_person_counts(tmp_path: Path) -> None:
    slate = _slate(tmp_path)
    roster, by_person = _legal_roster(slate)
    reversed_flex = (roster[0], *reversed(roster[1:]))
    summary = summarize_showdown_portfolio(
        slate, {"1": roster, "2": reversed_flex}
    )
    assert len(summary.duplicate_canonical_lineups) == 1
    assert summary.pairwise_person_overlap == (("1", "2", 6),)
    assert set(dict(summary.combined_person_counts).values()) == {2}
    assert sum(dict(summary.captain_counts).values()) == 2

    old_captain = next(row for row in slate.players if row.dk_id == roster[0])
    new_captain_flex = next(row for row in slate.players if row.dk_id == roster[1])
    swapped = (
        by_person[new_captain_flex.underlying_id].cpt_dk_id,
        by_person[old_captain.underlying_id].flex_dk_id,
        *roster[2:],
    )
    distinct = summarize_showdown_portfolio(slate, {"1": roster, "2": swapped})
    assert distinct.duplicate_canonical_lineups == ()
    assert distinct.pairwise_person_overlap == (("1", "2", 6),)


def test_request_round_trip_and_path_confinement_include_policy(tmp_path: Path) -> None:
    root = tmp_path / "allowed"
    root.mkdir()
    policy = root / "policy.json"
    policy.write_text("{}", encoding="utf-8")
    request = CoworkRunRequest.from_mapping(
        {"profile": "prior_review", "portfolio_policy_json": str(policy)},
        allowed_roots=(root,),
    )
    assert CoworkRunRequest.from_mapping(
        request.to_dict(), allowed_roots=(root,)
    ) == request
    assert required_next_inputs(request) == required_next_inputs(CoworkRunRequest())

    outside = tmp_path / "outside.json"
    outside.write_text("{}", encoding="utf-8")
    with pytest.raises(CoworkInputError, match="portfolio_policy_json is outside"):
        CoworkRunRequest.from_mapping(
            {"portfolio_policy_json": str(outside)}, allowed_roots=(root,)
        )


def test_appg_never_changes_policy_integer_math(tmp_path: Path) -> None:
    first_path = tmp_path / "first.csv"
    second_path = tmp_path / "second.csv"
    first_path.write_bytes(_salary_bytes())
    second_path.write_bytes(_salary_bytes().replace(b",0,", b",98765.4321,"))
    first_slate = parse_salaries(first_path)
    second_slate = parse_salaries(second_path)
    controls = {
        "max_combined_person_exposure": {
            "default_fraction": 0.67,
            "overrides": [],
        }
    }
    first = _validate(first_slate, _document(first_slate, ("1", "2", "3"), **controls), ("1", "2", "3"))
    second = _validate(second_slate, _document(second_slate, ("1", "2", "3"), **controls), ("1", "2", "3"))
    assert first.valid and second.valid
    assert [item.combined_max_entries for item in first.policy.effective_limits] == [
        item.combined_max_entries for item in second.policy.effective_limits
    ]


def test_cowork_policy_is_enforced_audited_snapshotted_and_replays_identically(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from nfl_dfs import prior_review as prior_review_module
    from .test_prior_review_profile import (
        _attachments,
        _cowork_args,
        _prepared_run,
    )

    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=datetime.now(timezone.utc) + timedelta(hours=6)
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    slate = parse_salaries(salary_path)
    entries = parse_entries(entry_path)
    policy_path = tmp_path / "policy" / "portfolio.json"
    policy_path.parent.mkdir()
    policy_path.write_text(
        json.dumps(
            portfolio_policy_template(
                slate,
                [entry.entry_id for entry in entries.authorizations],
                controls={"max_pairwise_person_overlap": 4},
            )
        ),
        encoding="utf-8",
    )
    runs = tmp_path / "runs"
    outputs = tmp_path / "outputs"
    retained = outputs / "earlier" / "DK_REVIEW_ENTRY_keep.csv"
    retained.parent.mkdir(parents=True)
    retained.write_bytes(b"preserve me")
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", runs)
    real = prior_review_module.run_prior_review
    monkeypatch.setattr(
        cli, "run_prior_review", lambda **kwargs: real(**kwargs, project=project)
    )

    first_code = cli.command_cowork_run(
        _cowork_args(
            tmp_path,
            attachments,
            run_id="sd4-first",
            output_dir=str(outputs),
            portfolio_policy_json=str(policy_path),
            prior_package_dir=str(package_dir),
        )
    )
    assert first_code == 0
    first_report = json.loads(
        (outputs / "sd4-first" / "cowork_run.json").read_text(encoding="utf-8")
    )
    assert first_report["FILE_VALID"] is True
    assert first_report["EVIDENCE_STATE"] == "UNKNOWN"
    assert first_report["MODEL_STATUS"] == "PRIOR_ONLY"
    assert first_report["RELEASE_DECISION"] == "DO_NOT_UPLOAD"
    assert Path(first_report["bulk_entry_csv"]).is_file()
    readable = first_report["prior_review_reports"]["readable_review"]
    assert readable["reconciliation"]["status"] == "PASS"
    assert readable["truths"] == {
        "FILE_VALID": True,
        "EVIDENCE_STATE": "UNKNOWN",
        "MODEL_STATUS": "PRIOR_ONLY",
        "RELEASE_DECISION": "DO_NOT_UPLOAD",
    }
    assert [entry["entry_id"] for entry in readable["entries"]] == [
        entry.entry_id for entry in entries.authorizations
    ]
    assert all(len(entry["slots"]) == 6 for entry in readable["entries"])
    assert all(
        entry["salary_total"] + entry["salary_remaining"] == 50_000
        for entry in readable["entries"]
    )
    assert readable["exposure"]["canonical_uniqueness"] == "PASS"
    assert readable["exposure"]["effective_pairwise_person_overlap"] == 4
    assert readable["exposure"]["pairwise_overlap"][0]["actual_people"] <= 4
    assert "official_activity" in {
        observation["category"] for observation in readable["evidence_observations"]
    }
    readable_json = Path(first_report["prior_review_artifacts"]["readable_review_json"])
    readable_html = Path(first_report["prior_review_artifacts"]["readable_review_html"])
    assert sha256_bytes(readable_json.read_bytes()) == first_report["prior_review_hashes"]["readable_review_json"]
    assert sha256_bytes(readable_html.read_bytes()) == first_report["prior_review_hashes"]["readable_review_html"]
    workbook = load_workbook(first_report["review_workbook"], data_only=False)
    assert workbook.sheetnames == [
        "Run Control", "Evidence Paste", "Portfolio", "QA", "Upload",
        "Exposure", "Review Evidence", "Artifacts",
    ]
    assert workbook["Upload"]["B17"].value == "PASS"
    assert workbook["Upload"]["B5"].value == "TRUE"
    assert workbook["Upload"]["B7"].value == "PRIOR_ONLY"
    assert workbook["Upload"]["B8"].value == "DO_NOT_UPLOAD"
    # 22 rows since P1b added the QB_DEPTH_ROLE_EVIDENCE_JSON control. The
    # number is asserted rather than derived on purpose: the print area is what
    # an operator actually sees on paper, so a row appearing or vanishing should
    # fail here and be looked at, not pass silently.
    assert str(workbook["Run Control"].print_area) == (
        "'Run Control'!$A$1:$D$22"
    )
    assert str(workbook["Evidence Paste"].print_area) == (
        "'Evidence Paste'!$A$1:$E$105,'Evidence Paste'!$G$4:$J$108,"
        "'Evidence Paste'!$L$4:$O$105,'Evidence Paste'!$Q$4:$W$37"
    )
    assert str(workbook["QA"].print_area) == "'QA'!$A$1:$G$44"
    assert str(workbook["Upload"].print_area) == "'Upload'!$A$1:$B$18"
    for sheet_name in (
        "Portfolio", "Exposure", "Review Evidence", "Artifacts",
    ):
        sheet = workbook[sheet_name]
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.print_title_rows == "$1:$4"
        assert len(sheet.sheet_view.selection) == 1
        assert sheet.sheet_view.selection[0].pane == "bottomLeft"
    assert first_report["portfolio_policy"]["valid"] is True
    assert first_report["portfolio_policy"]["enforcement_status"] == (
        "ENFORCED_AND_INDEPENDENTLY_AUDITED"
    )
    assert first_report["portfolio_policy"]["selector"]["enforcement_status"] == "PASS"
    assert first_report["portfolio_policy"]["independent_audit"]["status"] == "PASS"
    assert sha256_bytes(
        Path(first_report["portfolio_policy"]["normalized_policy"]).read_bytes()
    ) == first_report["portfolio_policy"]["normalized_policy_sha256"]
    assert retained.read_bytes() == b"preserve me"
    first_bytes = Path(first_report["bulk_entry_csv"]).read_bytes()
    snapshotted_request = runs / "sd4-first" / "run_request.json"
    request = json.loads(snapshotted_request.read_text(encoding="utf-8"))
    snapshotted_policy = Path(request["portfolio_policy_json"])
    assert snapshotted_policy.parent == runs / "sd4-first" / "inputs"
    assert snapshotted_policy.read_bytes() == policy_path.read_bytes()

    policy_path.write_text(policy_path.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    second_code = cli.command_cowork_run(
        _cowork_args(
            tmp_path,
            attachments,
            request=str(snapshotted_request),
            input_dir=None,
            profile=None,
            run_id="sd4-replay",
            output_dir=str(outputs),
            portfolio_policy_json=None,
            prior_package_dir=str(package_dir),
            build_priors=False,
        )
    )
    assert second_code == 0
    second_report = json.loads(
        (outputs / "sd4-replay" / "cowork_run.json").read_text(encoding="utf-8")
    )
    assert second_report["portfolio_policy"]["source_policy_sha256"] == (
        first_report["portfolio_policy"]["source_policy_sha256"]
    )
    assert second_report["portfolio_policy"]["normalized_policy_sha256"] == (
        first_report["portfolio_policy"]["normalized_policy_sha256"]
    )
    assert second_report["portfolio_policy"]["independent_audit"]["status"] == "PASS"
    assert Path(second_report["bulk_entry_csv"]).read_bytes() == first_bytes


def test_policy_lineup_count_mismatch_fails_before_generation_and_preserves_outputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from .test_prior_review_profile import _attachments, _cowork_args

    prepared = tmp_path / "prepared"
    prepared.mkdir()
    slate = _slate(prepared)
    salary_path = prepared / "DKSalaries.csv"
    entry_path = prepared / "DKEntries.csv"
    entry_path.write_bytes(
        b"Entry ID,Contest Name,Contest ID,Entry Fee,CPT,FLEX,FLEX,FLEX,FLEX,FLEX\n"
        b"900000001,Contest,1,$20,,,,,,\n"
        b"900000002,Contest,1,$20,,,,,,\n"
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    policy_path = tmp_path / "policy.json"
    policy_path.write_text(
        json.dumps(portfolio_policy_template(slate, ("900000001", "900000002"))),
        encoding="utf-8",
    )
    outputs = tmp_path / "outputs"
    retained = outputs / "earlier" / "DK_REVIEW_ENTRY_keep.csv"
    retained.parent.mkdir(parents=True)
    retained.write_bytes(b"keep")
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")

    code = cli.command_cowork_run(
        _cowork_args(
            tmp_path,
            attachments,
            run_id="lineup-count-mismatch",
            output_dir=str(outputs),
            portfolio_policy_json=str(policy_path),
            lineup_count=1,
            build_priors=True,
        )
    )
    assert code == 2
    report = json.loads(
        (outputs / "lineup-count-mismatch" / "cowork_run.json").read_text(
            encoding="utf-8"
        )
    )
    assert any(
        blocker.startswith("PORTFOLIO_POLICY_LINEUP_COUNT_MUST_MATCH_ENTRIES:")
        for blocker in report["blockers"]
    )
    assert report["bulk_entry_csv"] is None
    assert retained.read_bytes() == b"keep"
    assert list((outputs / "lineup-count-mismatch").rglob("DK_REVIEW_ENTRY_*.csv")) == []


def test_policy_is_not_silently_ignored_by_the_diagnostic_profile(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import cli
    from .test_prior_review_profile import _attachments, _cowork_args

    prepared = tmp_path / "prepared"
    prepared.mkdir()
    slate = _slate(prepared)
    salary_path = prepared / "DKSalaries.csv"
    entry_path = prepared / "DKEntries.csv"
    entry_path.write_bytes(
        b"Entry ID,Contest Name,Contest ID,Entry Fee,CPT,FLEX,FLEX,FLEX,FLEX,FLEX\n"
        b"900000001,Contest,1,$20,,,,,,\n"
        b"900000002,Contest,1,$20,,,,,,\n"
    )
    attachments = _attachments(tmp_path, salary_path, entry_path)
    policy_path = tmp_path / "policy.json"
    policy_path.write_text(
        json.dumps(portfolio_policy_template(slate, ("900000001", "900000002"))),
        encoding="utf-8",
    )
    monkeypatch.setattr(cli, "DEFAULT_RUNS_DIR", tmp_path / "runs")
    code = cli.command_cowork_run(
        _cowork_args(
            tmp_path,
            attachments,
            profile="diagnostic",
            run_id="wrong-profile",
            output_dir=str(tmp_path / "outputs"),
            portfolio_policy_json=str(policy_path),
        )
    )
    assert code == 2
    report = json.loads(
        (tmp_path / "outputs" / "wrong-profile" / "cowork_run.json").read_text(
            encoding="utf-8"
        )
    )
    assert any(
        blocker.startswith("PORTFOLIO_POLICY_PROFILE_UNSUPPORTED_SD4:")
        for blocker in report["blockers"]
    )
    assert report["bulk_entry_csv"] is None


def test_policy_assignment_artifact_mutation_fails_independent_audit_and_writes_no_review(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from nfl_dfs import prior_review as prior_review_module
    from nfl_dfs.portfolio_policy import write_normalized_portfolio_policy
    from .test_prior_review_profile import _prepared_run

    now = datetime.now(timezone.utc)
    salary_path, entry_path, package_dir, project = _prepared_run(
        tmp_path, expires_at=now + timedelta(hours=6)
    )
    slate = parse_salaries(salary_path)
    entries = parse_entries(entry_path)
    source_path = tmp_path / "policy.json"
    source_path.write_text(
        json.dumps(
            portfolio_policy_template(
                slate,
                [entry.entry_id for entry in entries.authorizations],
                controls={"max_pairwise_person_overlap": 4},
            )
        ),
        encoding="utf-8",
    )
    validation = validate_portfolio_policy_file(
        source_path,
        slate=slate,
        entry_ids=tuple(entry.entry_id for entry in entries.authorizations),
    )
    assert validation.valid and validation.policy is not None
    normalized_path = write_normalized_portfolio_policy(
        tmp_path / "policy.normalized.json", validation.policy
    )
    real_write = prior_review_module.write_assignments_csv

    def tampering_write(path, assignments, **kwargs):
        digest = real_write(path, assignments, **kwargs)
        Path(path).write_bytes(Path(path).read_bytes() + b"tampered\n")
        return digest

    monkeypatch.setattr(prior_review_module, "write_assignments_csv", tampering_write)
    output_root = tmp_path / "out"
    retained = output_root / "earlier" / "DK_REVIEW_ENTRY_keep.csv"
    retained.parent.mkdir(parents=True)
    retained.write_bytes(b"keep")
    outcome = prior_review_module.run_prior_review(
        salary_csv=salary_path,
        entry_csv=entry_path,
        label="audit-mutation",
        as_of=now,
        run_root=tmp_path / "run",
        output_root=output_root,
        prior_package_dir=package_dir,
        portfolio_policy=validation.policy,
        portfolio_policy_source_path=source_path,
        portfolio_policy_source_sha256=validation.source_sha256,
        portfolio_policy_normalized_path=normalized_path,
        portfolio_policy_normalized_sha256=validation.policy.normalized_sha256,
        project=project,
    )
    assert outcome.blocked
    assert outcome.stage == "EXPORT"
    assert "PORTFOLIO_POLICY_INDEPENDENT_AUDIT_FAILED" in outcome.blockers[0]
    assert outcome.reports["portfolio_policy_audit"]["status"] == "FAIL"
    assert retained.read_bytes() == b"keep"
    assert list(output_root.rglob("DK_REVIEW_ENTRY_audit-mutation.csv")) == []


# ----------------------------------------------------------------- Session 11b: subset binding


def test_integer_caps_and_denominators_follow_the_bound_rows_not_the_fillable_rows(tmp_path: Path) -> None:
    """A 0.5 fraction over 4 bound rows of 10 fillable rows allows 2, not 5."""

    slate = _slate(tmp_path)
    fillable = tuple(str(900000001 + index) for index in range(10))
    bound = (fillable[1], fillable[3], fillable[6], fillable[8])
    document = _document(
        slate, bound,
        max_combined_person_exposure={"default_fraction": 0.5, "overrides": []},
        max_captain_exposure={"default_fraction": 0.5, "overrides": []},
    )
    validation = _validate(slate, document, fillable)
    assert validation.valid, validation.blockers()
    policy = validation.policy
    assert policy.entry_ids == bound and policy.entry_count == 4
    assert {item.combined_max_entries for item in policy.effective_limits} == {2}
    assert {item.captain_max_entries for item in policy.effective_limits} == {2}
    assert policy.as_mapping()["effective"]["entry_count_denominator"] == 4
    assert policy.as_mapping()["bindings"]["entry_ids"] == list(bound)
    full = _validate(slate, _document(
        slate, fillable, max_combined_person_exposure={"default_fraction": 0.5, "overrides": []}), fillable)
    assert {item.combined_max_entries for item in full.policy.effective_limits} == {5}


def test_classic_integer_domains_follow_the_bound_rows(tmp_path: Path) -> None:
    from nfl_dfs.classic_portfolio_policy import (
        classic_portfolio_policy_template,
        default_search_limits,
        validate_classic_portfolio_policy_bytes,
    )

    supplied = Path(__file__).resolve().parent / "fixtures" / "supplied"
    slate = parse_salaries(supplied / "DKSalaries Salary CSV Classic.csv")
    template = parse_entries(supplied / "DKEntries CSV 20 entries.csv")
    fillable = tuple(entry.entry_id for entry in template.authorizations)[:10]
    bound = fillable[2:6]

    def validate(document):
        return validate_classic_portfolio_policy_bytes(
            json.dumps(document).encode("utf-8"), slate=slate, entry_ids=fillable,
            entry_sha256=template.raw_hash)

    document = classic_portfolio_policy_template(slate, bound, entry_sha256=template.raw_hash)
    assert document["selection"]["limits"] == default_search_limits(4).as_mapping()
    assert {rule["maximum_entries"] for rule in document["controls"]["stack_rules"]} == {4}
    validation = validate(document)
    assert validation.valid, validation.blockers()
    assert validation.policy.entry_count == 4 and validation.policy.entry_ids == bound
    assert {bound_.maximum_entries for bound_ in validation.policy.player_bounds} == {4}
    assert validation.policy.as_mapping()["effective"]["entry_count_denominator"] == 4
    person = document["bindings"]["people"][0]
    reference = {key: person[key] for key in ("underlying_id", "dk_id")}
    document["controls"]["player_exposure_bounds"] = [
        {**reference, "minimum_entries": 0, "maximum_entries": 5}]
    codes = {issue.code for issue in validate(document).problems}
    assert codes and "CLASSIC_POLICY_ENTRY_ID_BINDING_MISMATCH" not in codes  # 5 exceeds the 4 bound rows


def _fill(tmp_path, *, fill_count, policy_controls=None, max_person_overlap=4, forbidden_rosters=()):
    from nfl_dfs.selection import select_prior_lineups

    from .test_prior_selection import _prepared

    tmp_path.mkdir(exist_ok=True)
    slate, model, contract, splits = _prepared(tmp_path)
    bound = ("1", "2")
    fillable = ("1", "2", *(str(3 + index) for index in range(fill_count)))
    raw = json.dumps(portfolio_policy_template(slate, bound, controls=policy_controls or {})).encode("utf-8")
    validation = validate_portfolio_policy_bytes(raw, slate=slate, entry_ids=fillable)
    assert validation.valid, validation.blockers()
    return slate, select_prior_lineups(
        slate, model, splits, contract, count=2, fill_count=fill_count, portfolio_policy=validation.policy,
        max_person_overlap=max_person_overlap, forbidden_rosters=forbidden_rosters)


def test_sd3_fills_the_unbound_rows_after_its_joint_solve_under_the_runs_exclusions_only(tmp_path: Path) -> None:
    from nfl_dfs.lineups import roster_canonical_key
    from nfl_dfs.selection import select_prior_lineups

    from .test_prior_selection import _prepared

    (tmp_path / "top").mkdir()
    slate, model, contract, splits = _prepared(tmp_path / "top")
    (best,), _scores, _report = select_prior_lineups(slate, model, splits, contract, count=1)
    by_id = {row.dk_id: row for row in slate.players}
    faded = by_id[best.captain_dk_id].underlying_id
    identity = next(item for item in salary_person_bindings(slate) if item.underlying_id == faded)
    (second,), _scores, _report = select_prior_lineups(
        slate, model, splits, contract, count=1, forbidden_rosters=(best.roster,))

    slate, (lineups, _scores, report) = _fill(
        tmp_path / "fill", fill_count=2, forbidden_rosters=(second.roster,),
        policy_controls={"excluded_people": [identity.as_mapping()]})

    assert len(lineups) == 4 and [lineup.index for lineup in lineups] == [1, 2, 3, 4]
    people = [{by_id[dk_id].underlying_id for dk_id in lineup.roster} for lineup in lineups]
    assert all(faded not in group for group in people[:2])  # the policy's exclusion binds its rows
    assert lineups[2].roster == best.roster  # and not the fill's: its first lineup is the run's best
    keys = [roster_canonical_key(slate, lineup.roster) for lineup in lineups]
    assert len(set(keys)) == 4 and roster_canonical_key(slate, second.roster) not in keys
    fill = report["unbound_fill"]
    assert fill["source"] == "SHOWDOWN_SEQUENTIAL" and fill["lineups"] == 2
    assert fill["lineup_indexes"] == [3, 4]
    assert fill["no_good_rosters"] == {"policy_lineups": 2, "prefilled_rosters": 1}
    assert report["selected_lineup_count"] == 2 and len(report["pairwise_person_overlap"]) == 1


def test_an_unbound_fill_that_runs_out_of_distinct_lineups_delivers_nothing(tmp_path: Path) -> None:
    from nfl_dfs.selection import SelectionError

    with pytest.raises(SelectionError, match="SOLVER_RETURNED_NO_LINEUP:.*stage=UNBOUND_FILL") as caught:
        _fill(tmp_path, fill_count=4, max_person_overlap=0)
    assert caught.value.status == "SOLVER_RETURNED_NO_LINEUP"
    assert caught.value.facts["stage"] == "UNBOUND_FILL"


def test_the_sd3_audit_checks_the_bound_rows_and_holds_the_artifact_to_the_unbound_ones(tmp_path: Path) -> None:
    from nfl_dfs.portfolio_enforcement import audit_policy_assignments, build_policy_candidate_bank

    slate = _slate(tmp_path)
    fillable = ("1", "2", "3")
    raw = json.dumps(portfolio_policy_template(slate, ("1", "3"), controls={})).encode("utf-8")
    policy = validate_portfolio_policy_bytes(raw, slate=slate, entry_ids=fillable).policy
    objective = {row.dk_id: float(index) for index, row in enumerate(slate.players)}
    bank = build_policy_candidate_bank(slate, objective, candidate_limit=3, total_time_limit_seconds=5,
                                       per_solve_time_limit_seconds=1)
    rosters = {entry: bank.candidates[index].roster for index, entry in enumerate(fillable)}

    def audit(artifact_rows, unbound):
        artifact = ("\n".join(["Entry ID,CPT,FLEX,FLEX,FLEX,FLEX,FLEX",
                               *(",".join((entry, *rosters[entry])) for entry in artifact_rows)]) + "\n").encode()
        normalized = policy.canonical_bytes()
        return audit_policy_assignments(
            slate=slate, policy=policy, assignments=[(entry, rosters[entry]) for entry in ("1", "3")],
            salary_bytes=(tmp_path / "DKSalaries.csv").read_bytes(), entry_bytes=b"entries",
            expected_entry_sha256=sha256_bytes(b"entries"), source_policy_bytes=raw,
            expected_source_policy_sha256=sha256_bytes(raw), normalized_policy_bytes=normalized,
            expected_normalized_policy_sha256=sha256_bytes(normalized), assignment_artifact_bytes=artifact,
            expected_assignment_artifact_sha256=sha256_bytes(artifact), unbound_entry_ids=unbound)

    passed = audit(("1", "2", "3"), ("2",))
    assert passed.passed, passed.problems
    assert passed.entry_ids == ("1", "3") and len(passed.pairwise_person_overlap) == 1
    for artifact_rows, unbound in ((("1", "2", "3"), ()), (("1", "3"), ("2",)), (("1", "2", "2", "3"), ("2",))):
        failed = audit(artifact_rows, unbound)
        assert any(problem.startswith("PORTFOLIO_AUDIT_ASSIGNMENT_ARTIFACT_MISMATCH")
                   for problem in failed.problems), (artifact_rows, unbound)


def test_the_runs_own_exclusions_bind_the_unbound_fill_too(tmp_path: Path) -> None:
    from nfl_dfs.participation import build_participation_contract
    from nfl_dfs.selection import select_prior_lineups

    from .test_prior_selection import _prepared

    slate, model, _contract, splits = _prepared(tmp_path)
    (best,), _scores, _report = select_prior_lineups(slate, model, splits, _contract, count=1)
    by_id = {row.dk_id: row for row in slate.players}
    faded = by_id[best.captain_dk_id].underlying_id
    ids = [row.dk_id for row in slate.players if row.underlying_id == faded]
    contract = build_participation_contract(slate, operator_excluded_dk_ids=tuple(ids))
    raw = json.dumps(portfolio_policy_template(slate, ("1", "2"), controls={})).encode("utf-8")
    policy = validate_portfolio_policy_bytes(
        raw, slate=slate, entry_ids=("1", "2", "3", "4"), externally_excluded_people=(faded,)).policy
    lineups, _scores, report = select_prior_lineups(
        slate, model, splits, contract, count=2, fill_count=2, portfolio_policy=policy)
    assert len(lineups) == 4 and report["unbound_fill"]["lineups"] == 2
    assert all(faded not in {by_id[dk_id].underlying_id for dk_id in lineup.roster} for lineup in lineups)


def test_each_fill_solve_gets_what_the_bank_and_joint_solve_leave_of_the_window() -> None:
    from nfl_dfs.prior_review import _fill_solve_seconds

    class Window:
        def __init__(self, seconds, passed=False):
            self.seconds, self.passed_at_start = seconds, passed

        def improvement_remaining(self):
            return self.seconds

    assert _fill_solve_seconds(None, 3) is None and _fill_solve_seconds(Window(100.0), 0) is None
    assert _fill_solve_seconds(Window(100.0), 4) == pytest.approx(2.0)  # 10% of 100 s over 5
    assert _fill_solve_seconds(Window(10_000.0), 4) == 10.0  # never above the sequential default
    assert _fill_solve_seconds(Window(1.0), 4) == 0.5  # never under the solver's minimum
    assert _fill_solve_seconds(Window(100.0, passed=True), 4) == 0.5
    # Session 11c: a C2 policy's bank and joint solve take its declared limits.
    assert _fill_solve_seconds(Window(100.0), 4, declared_search_seconds=60.0) == pytest.approx(8.0)
    assert _fill_solve_seconds(Window(100.0), 4, declared_search_seconds=120.0) == 0.5
    assert _fill_solve_seconds(Window(1_000.0), 4, declared_search_seconds=60.0) == 10.0
